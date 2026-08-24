import copy
import html
import json
import math
import mimetypes
import os
import re
import socket
import subprocess
import sys
import threading
import time
import traceback
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
try:
    from importlib.metadata import PackageNotFoundError, version as get_package_version
except Exception:
    class PackageNotFoundError(Exception):
        pass

    def get_package_version(_package_name):
        raise PackageNotFoundError

from itertools import product
from pathlib import Path
from urllib.parse import parse_qs, quote, unquote, urlparse

DEPENDENCY_IMPORT_ERRORS = {}
DEFAULT_UI_HOST = "127.0.0.1"
DEFAULT_UI_PORT = 8765
DEFAULT_UI_PORT_SCAN_LIMIT = 10
UI_LAUNCH_HINT_FILENAME = "llm_expert_bench_ui_url.txt"
LOCAL_EXPERT_BATTLE_SUITE_ID = "local-expert-battle-48"
LOCAL_EXPERT_BATTLE_FILENAME = "local_expert_battle.json"
LOCAL_EXPERT_WIKI_DIR_ENV = "DIY_LLM_WIKI_DIR"
LLAMA_CPP_LAUNCHER_ROOT_ENV = "DIY_LLAMACPP_ROOT"
CURRENT_LLAMA_CPP_MODEL_FALLBACK = "current-llama.cpp-model"

try:
    import matplotlib.pyplot as plt
except Exception as exc:
    plt = None
    DEPENDENCY_IMPORT_ERRORS["matplotlib"] = exc

try:
    import pandas as pd
except Exception as exc:
    pd = None
    DEPENDENCY_IMPORT_ERRORS["pandas"] = exc

try:
    import questionary
    from questionary import Choice
except Exception as exc:
    questionary = None
    Choice = None
    DEPENDENCY_IMPORT_ERRORS["questionary"] = exc

try:
    import requests
except Exception as exc:
    requests = None
    DEPENDENCY_IMPORT_ERRORS["requests"] = exc

try:
    from openai import OpenAI
except Exception as exc:
    OpenAI = None
    DEPENDENCY_IMPORT_ERRORS["openai"] = exc

try:
    import msvcrt
except ImportError:
    msvcrt = None


NVIDIA_SMI_QUERY = [
    "nvidia-smi",
    "--query-gpu=index,name,memory.used,memory.total",
    "--format=csv,noheader,nounits",
]


PARAM_INFO = {
    "temperature": {
        "label": "溫度 (Temperature)",
        "range": "0.0 - 2.0",
        "desc": "調高更有創意，調低更穩定；Qwen 類模型常用 0.1 或 1.0。",
        "default": "0.1, 0.8",
        "backends": ["ollama", "llama.cpp"],
        "backend_keys": {"ollama": "temperature", "llama.cpp": "temperature"},
    },
    "num_ctx": {
        "label": "上下文長度 (Num_Ctx)",
        "range": "128 - 262144",
        "desc": "可測長上下文與 TTFT 影響；此項為 Ollama 請求級參數。",
        "default": "4096, 8192",
        "backends": ["ollama"],
        "backend_keys": {"ollama": "num_ctx"},
    },
    "num_predict": {
        "label": "最大生成 Token",
        "range": "-1, 1 - 4096",
        "desc": "控制回覆長度；llama.cpp 會映射為 n_predict。",
        "default": "256, 512",
        "backends": ["ollama", "llama.cpp"],
        "backend_keys": {"ollama": "num_predict", "llama.cpp": "n_predict"},
    },
    "top_p": {
        "label": "核心採樣 (Top_P)",
        "range": "0.0 - 1.0",
        "desc": "限制候選詞機率總和，數值越低越保守。",
        "default": "0.8, 0.95",
        "backends": ["ollama", "llama.cpp"],
        "backend_keys": {"ollama": "top_p", "llama.cpp": "top_p"},
    },
    "min_p": {
        "label": "最小概率 (Min_P)",
        "range": "0.0 - 1.0",
        "desc": "過濾低機率雜訊詞；常見平衡點在 0.05 左右。",
        "default": "0.02, 0.05",
        "backends": ["ollama", "llama.cpp"],
        "backend_keys": {"ollama": "min_p", "llama.cpp": "min_p"},
    },
    "repeat_penalty": {
        "label": "重複懲罰 (Repeat Penalty)",
        "range": "1.0 - 2.0",
        "desc": "降低重複句與繞圈輸出。",
        "default": "1.05, 1.15",
        "backends": ["ollama", "llama.cpp"],
        "backend_keys": {"ollama": "repeat_penalty", "llama.cpp": "repeat_penalty"},
    },
    "num_gpu": {
        "label": "GPU 層數 / 顯存卸載",
        "range": "0 - 100",
        "desc": "對 Ollama TPS 影響很大；llama.cpp 通常在 server 啟動時設定。",
        "default": "25, 50",
        "backends": ["ollama"],
        "backend_keys": {"ollama": "num_gpu"},
    },
    "enable_thinking": {
        "label": "Thinking / Reasoning 開關",
        "range": "enable | disable",
        "desc": "測試是否啟用模型的 thinking / reasoning 模式。Ollama 會送出 `think`，llama.cpp 會送入 `chat_template_kwargs.enable_thinking`。",
        "default": "disable, enable",
        "value_type": "boolean",
        "backends": ["ollama", "llama.cpp"],
        "backend_keys": {"ollama": "think", "llama.cpp": "enable_thinking"},
        "request_targets": {"ollama": "body", "llama.cpp": "chat_template_kwargs"},
    },
}

PARAM_GROUPS = {
    "🔥 生成核心": ["temperature", "num_ctx", "num_predict"],
    "⚖️ 採樣與懲罰": ["top_p", "min_p", "repeat_penalty"],
    "🧠 Thinking / Reasoning": ["enable_thinking"],
    "🖥️ 硬體與部署": ["num_gpu"],
}


CAPABILITY_OPTIONS = {
    "chat": {
        "label": "聊天能力",
        "description": "一般對話輸出，沿用目前的文字串流 benchmark。",
        "default_prompt": (
            "解釋一下 3D 列印使用 PETG 時，長期受力下的潛變（creep）風險，"
            "以及有哪些實際的改善方式。"
        ),
    },
    "tools": {
        "label": "Tools 調用能力",
        "description": "要求模型先呼叫工具，檢查是否真的輸出 tool_calls。",
        "default_prompt": (
            "請查詢台北今天的天氣。若你支援 tools 或 function calling，"
            "請先呼叫 `lookup_weather` 工具，不要直接回答。"
        ),
    },
}


SUITE_QUESTION_REQUIRED_FIELDS = (
    "id",
    "category",
    "title",
    "prompt",
    "expected_output",
    "evaluation_guide",
)

SUITE_SMOKE_7 = {
    "id": "suite-smoke-7",
    "version": "1.0.0",
    "title": "Seven-skill smoke suite / 七項能力冒煙測試",
    "description": (
        "One fixed question each for math, logic, reasoning, reading, translation, "
        "writing, and coding. / 數學、邏輯、推理、閱讀、翻譯、寫作與程式各一題。"
    ),
    "question_schema": {
        "id": "Stable question identifier / 穩定題目識別碼",
        "category": "Machine-readable skill category / 機器可讀能力分類",
        "title": "Bilingual short title / 中英雙語短標題",
        "prompt": "Complete user prompt sent to the model / 實際送給模型的完整提示",
        "expected_output": "Reference answer or expected response shape / 參考答案或預期輸出形式",
        "evaluation_guide": "Manual or future automatic scoring guidance / 人工或未來自動評分準則",
    },
    "questions": [
        {
            "id": "smoke7-math-01",
            "category": "math",
            "title": "Discount and tax / 折扣與稅額",
            "prompt": (
                "一件商品原價 800 元，先打 85 折，再針對折後價格加收 5% 稅金。"
                "請列出計算式，並以兩位小數給出最後應付金額。"
            ),
            "expected_output": "800 × 0.85 × 1.05 = 714.00 元。",
            "evaluation_guide": "The calculation and final amount 714.00 must both be correct.",
        },
        {
            "id": "smoke7-logic-01",
            "category": "logic",
            "title": "Truth-teller puzzle / 誠實者邏輯題",
            "prompt": (
                "A 說：「B 在說謊。」B 說：「我們兩個都在說謊。」已知每個人不是永遠說真話，"
                "就是永遠說假話。請判斷 A、B 各是哪一種人，並用兩句話說明理由。"
            ),
            "expected_output": "A 說真話，B 說假話。",
            "evaluation_guide": "The conclusion must be A truthful and B lying, with a consistent explanation.",
        },
        {
            "id": "smoke7-reasoning-01",
            "category": "reasoning",
            "title": "Access-chain reasoning / 權限鏈推理",
            "prompt": (
                "所有金屬鑰匙都放在紅盒中；紅盒放在上鎖的櫃子裡。小美可以進入放置櫃子的房間，"
                "但經理不在場時不能打開任何上鎖物件。今天經理不在。小美今天能拿到金屬鑰匙嗎？"
                "請依條件逐步回答，不要加入題目沒有提供的假設。"
            ),
            "expected_output": "不能；她雖能進房間，但無權打開上鎖的櫃子，因此無法取得紅盒內鑰匙。",
            "evaluation_guide": "The answer must be no and connect room access, the locked cabinet, and manager absence.",
        },
        {
            "id": "smoke7-reading-01",
            "category": "reading",
            "title": "Short-passage comprehension / 短文理解",
            "prompt": (
                "閱讀短文：『工廠把夜間冷卻水泵改為依溫度自動調速後，用電量下降 18%。"
                "不過在第一週，兩次溫度感測器誤報讓泵浦全速運轉。工程團隊因此加入雙感測器交叉驗證，"
                "之後四週沒有再發生誤報。』請回答：(1) 用電量下降的直接原因是什麼？"
                "(2) 團隊為何加入雙感測器交叉驗證？每題各用一句話。"
            ),
            "expected_output": "(1) 水泵改為依溫度自動調速。(2) 為避免單一感測器誤報使泵浦全速運轉。",
            "evaluation_guide": "Both answers must be grounded only in the passage and preserve the causal relationship.",
        },
        {
            "id": "smoke7-translation-01",
            "category": "translation",
            "title": "Technical translation / 技術翻譯",
            "prompt": (
                "請把下列繁體中文翻譯成自然、精確的英文，只輸出譯文；保留 GPU、24 GB 與 128k context "
                "三個技術標記不變：『這張 GPU 有 24 GB 記憶體，但啟用 128k context 時仍須留意 KV cache 的成長。』"
            ),
            "expected_output": (
                "This GPU has 24 GB of memory, but KV cache growth still needs to be monitored "
                "when 128k context is enabled."
            ),
            "evaluation_guide": "Meaning must be accurate, English natural, and all three required markers preserved.",
        },
        {
            "id": "smoke7-writing-01",
            "category": "writing",
            "title": "Concise professional writing / 精簡商務寫作",
            "prompt": (
                "請用繁體中文寫一封簡短專業郵件，通知團隊明天下午 3 點的模型部署延後到後天下午 2 點。"
                "內容必須包含主旨、延後原因是『驗證尚未完成』、向收件者致歉，以及請大家回覆是否能配合新時間。"
                "全文控制在 120 個中文字以內。"
            ),
            "expected_output": "A concise Traditional Chinese email containing all four required elements and the new time.",
            "evaluation_guide": "Check subject, both times, stated reason, apology, reply request, tone, and length constraint.",
        },
        {
            "id": "smoke7-coding-01",
            "category": "coding",
            "title": "Order-preserving deduplication / 保序去重",
            "prompt": (
                "請用 Python 實作 `dedupe_keep_order(items)`：移除重複項目但保留第一次出現的順序，"
                "時間複雜度需為 O(n)。請提供型別標註、函式本體與一個輸入輸出範例，不要使用外部套件。"
            ),
            "expected_output": "A valid O(n) Python implementation using a seen set plus an ordered result list.",
            "evaluation_guide": "Code must be valid, preserve first occurrence order, use type hints, and include one example.",
        },
    ],
}

BUILTIN_SUITES = {SUITE_SMOKE_7["id"]: SUITE_SMOKE_7}


def get_local_expert_battle_definition():
    """Load the repository-managed battle set without baking private wiki text into git."""
    suite_path = Path(__file__).with_name(LOCAL_EXPERT_BATTLE_FILENAME)
    try:
        with suite_path.open("r", encoding="utf-8") as file:
            suite = json.load(file)
    except FileNotFoundError as exc:
        raise ValueError(f"Local Expert Battle test set is missing: {suite_path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"Invalid Local Expert Battle test set: {exc}") from exc

    if suite.get("id") != LOCAL_EXPERT_BATTLE_SUITE_ID:
        raise ValueError(f"Unexpected Local Expert Battle suite id: {suite.get('id')}")
    return suite


def get_local_expert_wiki_root():
    configured_path = os.environ.get(LOCAL_EXPERT_WIKI_DIR_ENV, "").strip()
    return Path(configured_path).expanduser() if configured_path else Path.home() / "wiki"


def load_local_expert_wiki_excerpt(question):
    source_name = str(question.get("wiki_source") or "").strip()
    if not source_name:
        return "", "", 0

    wiki_root = get_local_expert_wiki_root()
    source_path = (wiki_root / source_name).resolve()
    try:
        source_path.relative_to(wiki_root.resolve())
    except ValueError as exc:
        raise ValueError(f"Wiki source must remain inside {wiki_root}: {source_name}") from exc
    if not source_path.is_file():
        raise ValueError(
            f"Wiki source is unavailable: {source_path}. Set {LOCAL_EXPERT_WIKI_DIR_ENV} to your wiki directory."
        )

    content = source_path.read_text(encoding="utf-8", errors="replace").replace("\r\n", "\n")
    start = max(0, int(question.get("wiki_excerpt_start") or 0))
    excerpt = content[start : start + 2400].strip()
    if len(excerpt) < 2000:
        # A shorter final segment is still useful, but make the evidence limitation explicit to the reviewer.
        excerpt = content[max(0, len(content) - 2400) :].strip()
    if len(excerpt) < 1200:
        raise ValueError(f"Wiki source is too short for a long-summary test: {source_path}")
    return excerpt, source_path.as_posix(), len(excerpt)


def hydrate_local_expert_battle_suite(suite):
    hydrated = copy.deepcopy(suite)
    for question in hydrated.get("questions") or []:
        template = question.pop("prompt_template", "")
        if template:
            excerpt, source_path, excerpt_chars = load_local_expert_wiki_excerpt(question)
            question["prompt"] = str(template).format(excerpt=excerpt)
            question["wiki_source_path"] = source_path
            question["wiki_excerpt_chars"] = excerpt_chars
    return hydrated


def get_suite_definition(suite_id):
    suite = (
        hydrate_local_expert_battle_suite(get_local_expert_battle_definition())
        if suite_id == LOCAL_EXPERT_BATTLE_SUITE_ID
        else BUILTIN_SUITES.get(suite_id)
    )
    if suite is None:
        raise ValueError(f"Unknown benchmark suite: {suite_id}")

    questions = suite.get("questions") or []
    if not questions:
        raise ValueError(f"Benchmark suite has no questions: {suite_id}")

    question_ids = set()
    for question in questions:
        missing_fields = [field for field in SUITE_QUESTION_REQUIRED_FIELDS if not question.get(field)]
        if missing_fields:
            raise ValueError(
                f"Question in {suite_id} is missing required fields: {', '.join(missing_fields)}"
            )
        if question["id"] in question_ids:
            raise ValueError(f"Duplicate question id in {suite_id}: {question['id']}")
        question_ids.add(question["id"])
    return copy.deepcopy(suite)


def resolve_benchmark_questions(config):
    capability = config.get("capability", "chat")
    if capability in BUILTIN_SUITES or capability == LOCAL_EXPERT_BATTLE_SUITE_ID:
        suite = get_suite_definition(capability)
        return [
            {
                "suite_id": suite["id"],
                "suite_version": suite["version"],
                **question,
            }
            for question in suite["questions"]
        ]

    return [
        {
            "suite_id": "",
            "suite_version": "",
            "id": "",
            "category": "",
            "title": "",
            "prompt": str(config.get("prompt", "")),
            "expected_output": "",
            "evaluation_guide": "",
        }
    ]

TOOL_BENCHMARK_TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "lookup_weather",
            "description": "Look up the current weather for a city.",
            "parameters": {
                "type": "object",
                "properties": {
                    "city": {
                        "type": "string",
                        "description": "City name in Chinese or English.",
                    },
                    "unit": {
                        "type": "string",
                        "enum": ["celsius", "fahrenheit"],
                        "description": "Preferred temperature unit.",
                    },
                },
                "required": ["city"],
            },
        },
    }
]


def get_ollama_models():
    try:
        response = requests.get("http://localhost:11434/api/tags", timeout=2)
        response.raise_for_status()
        models = response.json().get("models", [])
        return [model["name"] for model in models if model.get("name")]
    except requests.RequestException:
        return []


def get_openai_compatible_models(base_url, timeout_seconds=3):
    """Return model IDs currently served by an OpenAI-compatible endpoint."""
    models_url = str(base_url).rstrip("/") + "/models"
    try:
        response = requests.get(models_url, timeout=timeout_seconds)
        response.raise_for_status()
        payload = response.json()
    except (requests.RequestException, ValueError, TypeError):
        return []

    seen = set()
    model_ids = []
    for item in payload.get("data") or []:
        if not isinstance(item, dict):
            continue
        model_id = str(item.get("id") or "").strip()
        normalized_id = model_id.casefold()
        if not model_id or normalized_id in seen:
            continue
        seen.add(normalized_id)
        model_ids.append(model_id)
    return model_ids


def get_llama_cpp_launcher_root():
    configured_path = os.environ.get(LLAMA_CPP_LAUNCHER_ROOT_ENV, "").strip()
    if configured_path:
        return Path(configured_path).expanduser()
    return Path(__file__).resolve().parent.parent / "easy_llamacpp"


def get_llama_cpp_models():
    """Read the catalog refreshed by easy_llamacpp's configured GGUF scanner."""
    launcher_root = get_llama_cpp_launcher_root()
    index_path = launcher_root / "json" / "model-index.json"
    try:
        with index_path.open("r", encoding="utf-8-sig") as file:
            payload = json.load(file)
    except (OSError, json.JSONDecodeError):
        return []

    default_model_id = str(payload.get("default_model_id") or "")
    catalog = []
    seen_names = set()
    for item in payload.get("models") or []:
        if not isinstance(item, dict):
            continue
        model_path = str(item.get("path") or "").strip()
        model_name = str(item.get("name") or Path(model_path).stem).strip()
        normalized_name = model_name.casefold()
        path_name = Path(model_path).name.casefold()
        if (
            not model_name
            or not model_path.casefold().endswith(".gguf")
            or normalized_name in seen_names
            or "mmproj" in normalized_name
            or "mmproj" in path_name
        ):
            continue
        seen_names.add(normalized_name)
        catalog.append(
            {
                "name": model_name,
                "path": model_path,
                "id": str(item.get("id") or ""),
                "available": Path(model_path).is_file(),
                "is_default": str(item.get("id") or "") == default_model_id,
            }
        )

    return sorted(catalog, key=lambda item: (not item["is_default"], item["name"].casefold()))


def resolve_llama_cpp_auto_switch_models(model_names):
    catalog_by_name = {item["name"].casefold(): item for item in get_llama_cpp_models()}
    resolved_models = []
    missing_models = []
    unavailable_models = []
    for model_name in model_names:
        model_entry = catalog_by_name.get(str(model_name).casefold())
        if model_entry is None:
            missing_models.append(str(model_name))
            continue
        if not model_entry.get("available"):
            unavailable_models.append(str(model_name))
            continue
        resolved_models.append(model_entry)
    if missing_models:
        raise ValueError(
            "Selected llama.cpp model(s) are not in easy_llamacpp model-index.json: "
            + ", ".join(missing_models)
        )
    if unavailable_models:
        raise ValueError(
            "Selected GGUF file(s) are missing on disk: " + ", ".join(unavailable_models)
        )
    return resolved_models


def get_llama_cpp_switch_port(base_url):
    parsed = urlparse(str(base_url))
    if parsed.hostname not in {"127.0.0.1", "localhost", "::1"}:
        raise ValueError("Auto-switch only supports a local llama.cpp URL (localhost or 127.0.0.1).")
    return parsed.port or 80


def start_llama_cpp_model(model_entry, base_url, ready_timeout_seconds=240):
    launcher_root = get_llama_cpp_launcher_root()
    launcher_script = launcher_root / "PS1" / "Start_LCPP.ps1"
    if not launcher_script.is_file():
        raise RuntimeError(f"easy_llamacpp launcher is missing: {launcher_script}")
    model_path = Path(str(model_entry.get("path") or ""))
    if not model_path.is_file():
        raise RuntimeError(f"Selected GGUF is missing: {model_path}")

    port = get_llama_cpp_switch_port(base_url)
    completed = subprocess.run(
        [
            "powershell.exe",
            "-NoProfile",
            "-ExecutionPolicy",
            "Bypass",
            "-File",
            str(launcher_script),
            "-BypassMenu",
            "-Background",
            "-NoBrowser",
            "-NoPause",
            "-ReturnNonZeroOnError",
            "-Port",
            str(port),
            "-ModelPath",
            str(model_path),
        ],
        cwd=str(launcher_root),
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if completed.returncode != 0:
        detail = (completed.stderr or completed.stdout).strip()
        raise RuntimeError(detail or f"easy_llamacpp failed to start {model_entry['name']}")

    expected_ids = {str(model_entry.get("name") or "").casefold(), model_path.stem.casefold()}
    models_url = str(base_url).rstrip("/") + "/models"
    deadline = time.monotonic() + max(1, int(ready_timeout_seconds))
    last_error = ""
    while time.monotonic() < deadline:
        try:
            response = requests.get(models_url, timeout=3)
            response.raise_for_status()
            payload = response.json()
            served_ids = {
                str(item.get("id") or "").casefold()
                for item in payload.get("data", [])
                if isinstance(item, dict)
            }
            if expected_ids & served_ids:
                return
            last_error = "llama-server is responding with a different model"
        except (requests.RequestException, ValueError) as exc:
            last_error = str(exc)
        time.sleep(1)
    raise RuntimeError(
        f"Timed out waiting for {model_entry['name']} at {models_url}. "
        f"Last status: {last_error or 'no response'}"
    )


BOOLEAN_TRUE_ALIASES = {
    "1",
    "true",
    "on",
    "enable",
    "enabled",
    "yes",
    "y",
    "think",
    "thinking",
    "啟用",
    "開啟",
    "是",
}

BOOLEAN_FALSE_ALIASES = {
    "0",
    "false",
    "off",
    "disable",
    "disabled",
    "no",
    "n",
    "nothink",
    "關閉",
    "停用",
    "否",
}


def get_param_value_type(param_key):
    return PARAM_INFO.get(param_key, {}).get("value_type", "number")


def parse_boolean_value(raw_value):
    normalized = str(raw_value).strip().lower()
    if normalized in BOOLEAN_TRUE_ALIASES:
        return True
    if normalized in BOOLEAN_FALSE_ALIASES:
        return False
    raise ValueError(
        f"無法解析布林值：{raw_value}。請使用 enable/disable、true/false、on/off 或 1/0。"
    )


def format_param_value_for_display(param_key, value):
    if get_param_value_type(param_key) == "boolean":
        normalized_value = value
        if isinstance(value, str):
            normalized_value = parse_boolean_value(value)
        return "enable" if bool(normalized_value) else "disable"
    return str(value)


def format_param_values_for_display(param_key, values):
    return ", ".join(format_param_value_for_display(param_key, value) for value in values)


def resolve_thinking_mode(value):
    if value in (None, "", "N/A"):
        return "default"
    try:
        return "enable" if parse_boolean_value(value) else "disable"
    except ValueError:
        return str(value)


def get_thinking_mode_for_run(params):
    if not isinstance(params, dict) or "enable_thinking" not in params:
        return "default"
    return resolve_thinking_mode(params.get("enable_thinking"))


def get_param_request_target(param_key, backend):
    info = PARAM_INFO.get(param_key, {})
    request_targets = info.get("request_targets", {})
    if backend in request_targets:
        return request_targets[backend]
    return "options" if backend == "ollama" else "body"


def parse_csv_values(raw_text, param_key=None):
    value_type = get_param_value_type(param_key)
    values = []
    for item in raw_text.split(","):
        item = item.strip()
        if not item:
            continue
        if value_type == "boolean":
            values.append(parse_boolean_value(item))
            continue
        try:
            if "." in item or "e" in item.lower():
                value = float(item)
                values.append(int(value) if value.is_integer() else value)
            else:
                values.append(int(item))
        except ValueError as exc:
            raise ValueError(f"無法解析數值：{item}") from exc

    if not values:
        raise ValueError("至少需要一個測試值。")

    return values


def ask_param_values(param_key):
    info = PARAM_INFO[param_key]
    while True:
        raw_value = questionary.text(
            f"輸入 {info['label']} 測試值 (逗號隔開):",
            default=info["default"],
        ).ask()
        if raw_value is None:
            return None
        try:
            return parse_csv_values(raw_value, param_key=param_key)
        except ValueError as exc:
            print(f"⚠️ {exc} 請重新輸入。")


def format_param_dict(params):
    if not params:
        return "預設參數"
    joined = ", ".join(
        f"{key}={format_param_value_for_display(key, value)}" for key, value in params.items()
    )
    return "{" + joined + "}"


def build_backend_options(backend, params):
    backend_options = {}
    for key, value in params.items():
        backend_key = PARAM_INFO[key]["backend_keys"].get(backend)
        if backend_key:
            backend_options[backend_key] = value
    return backend_options


def build_backend_extra_body(backend, params):
    extra_body = {}
    options = {}
    chat_template_kwargs = {}

    for key, value in params.items():
        backend_key = PARAM_INFO[key]["backend_keys"].get(backend)
        if not backend_key:
            continue

        request_target = get_param_request_target(key, backend)
        if request_target == "options":
            options[backend_key] = value
        elif request_target == "chat_template_kwargs":
            chat_template_kwargs[backend_key] = value
        else:
            extra_body[backend_key] = value

    if backend == "ollama":
        if options:
            extra_body["options"] = options
        return extra_body

    if options:
        extra_body.update(options)
    if chat_template_kwargs:
        extra_body["chat_template_kwargs"] = chat_template_kwargs
    return extra_body


def build_ollama_modelfile_params(params):
    modelfile_params = {}
    for key, value in params.items():
        backend_key = PARAM_INFO[key]["backend_keys"].get("ollama")
        if not backend_key:
            continue
        if get_param_request_target(key, "ollama") != "options":
            continue
        modelfile_params[backend_key] = value
    return modelfile_params


SYSTEM_PROMPT_BLOCK_SEPARATOR = "---"
BACK_ACTION = "__back__"


def build_system_prompt_variants(system_prompts):
    prompts = [prompt.strip() for prompt in (system_prompts or []) if (prompt or "").strip()]
    if not prompts:
        return [{"label": "N/A", "text": ""}]
    return [{"label": f"SP{index}", "text": prompt} for index, prompt in enumerate(prompts, start=1)]


def build_benchmark_messages(capability, prompt, system_prompt_text=""):
    messages = []
    if (system_prompt_text or "").strip():
        messages.append({"role": "system", "content": system_prompt_text.strip()})
    if capability == "tools":
        messages.append(
            {
                "role": "system",
                "content": (
                    "You are being benchmarked for tool calling. "
                    "If a suitable tool is provided, call the tool before answering."
                ),
            }
        )
    messages.append({"role": "user", "content": prompt})
    return messages


def build_chat_request_payload(config, model, request_kwargs, system_prompt_text="", prompt=None):
    capability = config.get("capability", "chat")
    request_prompt = config["prompt"] if prompt is None else prompt
    payload = {
        "model": model,
        "messages": build_benchmark_messages(capability, request_prompt, system_prompt_text),
        "stream": True,
        **request_kwargs,
    }
    if config.get("backend") == "ollama":
        payload["stream_options"] = {"include_usage": True}
    if capability == "tools":
        payload["tools"] = TOOL_BENCHMARK_TOOLS
        payload["tool_choice"] = "auto"
    return payload


def parse_non_content_types(value):
    if value is None:
        return set()
    if isinstance(value, str):
        items = [item.strip() for item in value.split(",")]
    else:
        items = [str(item).strip() for item in value]
    return {item for item in items if item and item != "none"}


def adjust_classification_for_capability(classification, capability):
    if capability != "tools":
        return classification

    adjusted = classification.copy()
    non_content_types = parse_non_content_types(adjusted.get("Non_Content_Types"))
    if "tool_calls" in non_content_types:
        adjusted["Status"] = "ok"
        adjusted["Output_Category"] = "tool_call"
        finish_reason = adjusted.get("Finish_Reason") or "unknown"
        adjusted["Diagnosis"] = (
            "Received tool_calls payload during the tool benchmark "
            f"(finish_reason={finish_reason})."
        )
        return adjusted

    if adjusted["Status"] == "ok" and adjusted["Output_Category"] == "normal_content":
        adjusted["Status"] = "warning"
        adjusted["Output_Category"] = "text_reply_without_tool"
        adjusted["Diagnosis"] = (
            "Received textual content, but no tool_calls payload was emitted during the "
            "tool benchmark."
        )
    elif adjusted["Output_Category"] == "empty_reply":
        adjusted["Diagnosis"] = (
            "The tool benchmark finished without textual content or tool_calls payload."
        )
    elif adjusted["Output_Category"] == "non_content_stream":
        adjusted["Diagnosis"] = (
            "The tool benchmark returned non-content payloads, but none were tool_calls."
        )

    return adjusted


def query_nvidia_vram_snapshot():
    try:
        result = subprocess.run(
            NVIDIA_SMI_QUERY,
            capture_output=True,
            text=True,
            check=True,
            timeout=3,
        )
    except (FileNotFoundError, subprocess.SubprocessError):
        return None

    snapshot = []
    for line in result.stdout.splitlines():
        parts = [part.strip() for part in line.split(",", 3)]
        if len(parts) != 4:
            continue
        try:
            snapshot.append(
                {
                    "index": int(parts[0]),
                    "name": parts[1],
                    "memory_used_mib": int(parts[2]),
                    "memory_total_mib": int(parts[3]),
                }
            )
        except ValueError:
            continue

    return snapshot or None


def empty_vram_metrics():
    return {
        "VRAM_Base_MiB": None,
        "VRAM_Peak_MiB": None,
        "VRAM_Delta_MiB": None,
        "VRAM_Detail": "N/A",
    }


def format_mib_value(value):
    return "N/A" if value is None or pd.isna(value) else f"{int(value)} MiB"


def format_numeric_value(value, digits):
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value):.{digits}f}"


def format_probability_value(value, digits=1):
    if value is None or pd.isna(value):
        return "N/A"
    return f"{float(value) * 100:.{digits}f}%"


def format_text_value(value, default="N/A"):
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return default
    text = str(value).strip()
    return text or default


def format_token_count(value, default="N/A"):
    if value is None:
        return default
    try:
        if pd.isna(value):
            return default
        return int(value)
    except (TypeError, ValueError):
        return default


def calculate_efficiency_score(tps, vram_peak_mib):
    if tps is None or pd.isna(tps):
        return None
    if vram_peak_mib is None or pd.isna(vram_peak_mib) or vram_peak_mib <= 0:
        return None
    score = tps / (vram_peak_mib / 1024)
    return round(score, 3)


TOKEN_ESTIMATE_PATTERN = re.compile(
    r"[\u3400-\u4dbf\u4e00-\u9fff\u3040-\u30ff\uac00-\ud7af]"
    r"|[A-Za-z0-9]+(?:['._:/-][A-Za-z0-9]+)*"
    r"|[^\s]"
)


def estimate_token_count(text):
    normalized_text = normalize_text_content(text)
    if not normalized_text:
        return 0
    return len(TOKEN_ESTIMATE_PATTERN.findall(normalized_text))


def parse_optional_int(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass

    if isinstance(value, bool):
        return int(value)
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value)

    text = str(value).strip()
    if not text:
        return None
    try:
        return int(text)
    except ValueError:
        try:
            return int(float(text))
        except ValueError:
            return None


def convert_ns_to_seconds(value):
    numeric_value = parse_optional_int(value)
    if numeric_value is None:
        return None
    if numeric_value <= 0:
        return 0.0
    return round(numeric_value / 1_000_000_000, 6)


def calculate_duration_tps(unit_count, duration_seconds):
    if unit_count is None or pd.isna(unit_count) or unit_count <= 0:
        return None
    if duration_seconds is None or pd.isna(duration_seconds):
        return None
    if duration_seconds <= 0:
        return 0.0
    return round(unit_count / duration_seconds, 2)


def calculate_text_tps(unit_count, first_text_time, end_time):
    if unit_count is None or pd.isna(unit_count) or unit_count <= 0:
        return None
    if first_text_time is None or end_time is None:
        return None

    generation_time = end_time - first_text_time
    if generation_time <= 0:
        return 0.0
    return round(unit_count / generation_time, 2)


def calculate_text_duration(unit_count, first_text_time, end_time):
    if unit_count is None or pd.isna(unit_count) or unit_count <= 0:
        return None
    if first_text_time is None or end_time is None:
        return None

    generation_time = end_time - first_text_time
    if generation_time <= 0:
        return 0.0
    return round(generation_time, 3)


def calculate_phase_time(start_time, end_time):
    if start_time is None or end_time is None:
        return None
    duration = end_time - start_time
    if duration <= 0:
        return 0.0
    return round(duration, 3)


def calculate_output_thinking_ratio(output_chars, thinking_chars):
    if thinking_chars is None or pd.isna(thinking_chars) or thinking_chars <= 0:
        return None
    if output_chars is None or pd.isna(output_chars):
        return None
    return round(output_chars / thinking_chars, 3)


def normalize_text_content(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        parts = []
        for item in value:
            if isinstance(item, dict):
                text = item.get("text")
                if text:
                    parts.append(str(text))
            elif item is not None:
                parts.append(str(item))
        return "".join(parts)
    return str(value)


def extract_object_payload(raw_object):
    if raw_object is None:
        return {}

    if isinstance(raw_object, dict):
        raw_payload = raw_object
    elif hasattr(raw_object, "model_dump"):
        try:
            raw_payload = raw_object.model_dump(exclude_none=True)
        except TypeError:
            raw_payload = raw_object.model_dump()
    elif hasattr(raw_object, "dict"):
        try:
            raw_payload = raw_object.dict(exclude_none=True)
        except TypeError:
            raw_payload = raw_object.dict()
    else:
        try:
            raw_payload = vars(raw_object)
        except TypeError:
            return {}

    payload = {}
    for key, value in raw_payload.items():
        if value is None:
            continue
        if isinstance(value, str) and value == "":
            continue
        if isinstance(value, (list, dict)) and not value:
            continue
        payload[key] = value

    return payload


def extract_delta_payload(delta):
    return extract_object_payload(delta)


def extract_stream_usage_metrics(chunk):
    payload = extract_object_payload(chunk)
    usage_payload = extract_object_payload(payload.get("usage") or getattr(chunk, "usage", None))

    prompt_tokens = parse_optional_int(usage_payload.get("prompt_tokens"))
    completion_tokens = parse_optional_int(usage_payload.get("completion_tokens"))
    total_tokens = parse_optional_int(usage_payload.get("total_tokens"))
    prompt_eval_count = parse_optional_int(payload.get("prompt_eval_count"))
    eval_count = parse_optional_int(payload.get("eval_count"))

    prompt_token_value = prompt_tokens if prompt_tokens is not None else prompt_eval_count
    completion_token_value = completion_tokens if completion_tokens is not None else eval_count
    total_token_value = total_tokens
    if total_token_value is None and prompt_token_value is not None and completion_token_value is not None:
        total_token_value = prompt_token_value + completion_token_value

    return {
        "prompt_tokens": prompt_token_value,
        "completion_tokens": completion_token_value,
        "total_tokens": total_token_value,
        "prompt_eval_count": prompt_eval_count if prompt_eval_count is not None else prompt_token_value,
        "eval_count": eval_count if eval_count is not None else completion_token_value,
        "prompt_eval_duration_s": convert_ns_to_seconds(payload.get("prompt_eval_duration")),
        "eval_duration_s": convert_ns_to_seconds(payload.get("eval_duration")),
        "total_duration_s": convert_ns_to_seconds(payload.get("total_duration")),
        "load_duration_s": convert_ns_to_seconds(payload.get("load_duration")),
    }


def normalize_non_content_type(field_name):
    field_map = {
        "role": "role",
        "tool_calls": "tool_calls",
        "reasoning": "reasoning",
        "refusal": "refusal",
        "audio": "audio",
        "function_call": "tool_calls",
    }
    return field_map.get(field_name, "other")


def inspect_stream_chunk(chunk):
    chunk_info = {
        "content": "",
        "non_content_types": [],
        "finish_reason": None,
    }

    choices = getattr(chunk, "choices", None) or []
    if not choices:
        return chunk_info

    choice = choices[0]
    delta_payload = extract_delta_payload(getattr(choice, "delta", None))
    chunk_info["content"] = normalize_text_content(delta_payload.pop("content", None))
    chunk_info["non_content_types"] = sorted(
        {normalize_non_content_type(field_name) for field_name in delta_payload}
    )
    chunk_info["finish_reason"] = getattr(choice, "finish_reason", None) or None
    return chunk_info


def classify_stream_result(
    chunk_records,
    start_time,
    end_time,
    first_event_time=None,
    first_content_time=None,
    first_thinking_time=None,
    error_message=None,
):
    total_chunks = len(chunk_records)
    content_chunks = 0
    thinking_chunks = 0
    non_content_chunks = 0
    non_content_types = set()
    finish_reason = None
    output_chars = 0
    thinking_chars = 0
    output_tokens = 0
    thinking_tokens = 0
    usage_metrics = {}

    for record in chunk_records:
        if record["content"]:
            content_chunks += 1
            output_chars += len(record["content"])
            output_tokens += record.get("content_tokens", estimate_token_count(record["content"]))
        if record.get("thinking"):
            thinking_chunks += 1
            thinking_chars += len(record["thinking"])
            thinking_tokens += record.get("thinking_tokens", estimate_token_count(record["thinking"]))
        if record["non_content_types"]:
            non_content_chunks += 1
            non_content_types.update(record["non_content_types"])
        if record["finish_reason"]:
            finish_reason = record["finish_reason"]
        for key, value in (record.get("usage_metrics") or {}).items():
            if value is not None:
                usage_metrics[key] = value

    first_event_seconds = round(first_event_time - start_time, 3) if first_event_time is not None else None
    stream_duration_seconds = round(end_time - start_time, 3)

    if content_chunks > 0 and first_content_time is not None:
        ttft = round(first_content_time - start_time, 3)
        generation_time = end_time - first_content_time
        tps = round(content_chunks / generation_time, 2) if generation_time > 0 else 0.0
    else:
        ttft = None
        tps = None

    substantive_non_content_types = sorted(
        non_content_type for non_content_type in non_content_types if non_content_type != "role"
    )

    if content_chunks > 0:
        output_category = "normal_content"
        if error_message:
            status = "error"
            diagnosis = "Stream interrupted after textual content was received."
        elif not finish_reason:
            status = "warning"
            diagnosis = "Textual content was received, but the stream ended without a terminal finish_reason."
        else:
            status = "ok"
            diagnosis = f"Received textual content and completed with finish_reason={finish_reason}."
    elif error_message:
        status = "error"
        output_category = "early_stop"
        diagnosis = "Stream interrupted before any textual content was received."
    elif finish_reason:
        status = "warning"
        if substantive_non_content_types:
            output_category = "non_content_stream"
            diagnosis = (
                f"Completed with finish_reason={finish_reason} but only non-content payloads "
                f"were received: {', '.join(substantive_non_content_types)}."
            )
        else:
            output_category = "empty_reply"
            diagnosis = f"Completed with finish_reason={finish_reason} but no textual content was received."
    else:
        status = "warning"
        output_category = "early_stop"
        diagnosis = "Stream ended before any textual content or terminal finish_reason was received."

    prompt_tokens = usage_metrics.get("prompt_tokens")
    completion_tokens = usage_metrics.get("completion_tokens")
    total_tokens = usage_metrics.get("total_tokens")
    prefill_time_seconds = usage_metrics.get("prompt_eval_duration_s")
    prefill_tps = calculate_duration_tps(prompt_tokens, prefill_time_seconds)
    if prefill_tps is None and prompt_tokens is not None and ttft is not None:
        prefill_time_seconds = ttft
        prefill_tps = calculate_duration_tps(prompt_tokens, ttft)

    return {
        "Status": status,
        "Output_Category": output_category,
        "Diagnosis": diagnosis,
        "Finish_Reason": finish_reason,
        "Total_Chunks": total_chunks,
        "Content_Chunks": content_chunks,
        "Thinking_Chunks": thinking_chunks,
        "Non_Content_Chunks": non_content_chunks,
        "Non_Content_Types": ", ".join(sorted(non_content_types)) if non_content_types else "none",
        "First_Event_s": first_event_seconds,
        "Stream_Duration_s": stream_duration_seconds,
        "TTFT": ttft,
        "TPS": tps,
        "Prompt_Tokens": prompt_tokens,
        "Completion_Tokens": completion_tokens,
        "Total_Tokens": total_tokens,
        "Prefill_Time_s": prefill_time_seconds,
        "Prefill_TPS": prefill_tps,
        "Thinking_Chars": thinking_chars,
        "Thinking_Tokens": thinking_tokens,
        "Output_Chars": output_chars,
        "Output_Tokens": output_tokens,
        "Thinking_Time_s": calculate_phase_time(start_time, first_content_time),
        "Answer_Time_s": calculate_phase_time(first_content_time, end_time),
        "Output_Time_s": calculate_text_duration(output_chars, first_content_time, end_time),
        "Thinking_TPS": calculate_text_tps(thinking_tokens, first_thinking_time, end_time),
        "Output_TPS": calculate_text_tps(output_tokens, first_content_time, end_time),
        "Output_Thinking_Ratio": calculate_output_thinking_ratio(output_chars, thinking_chars),
    }


def build_result_row(
    run_id,
    config,
    model,
    param_set,
    applied_params,
    display_params,
    classification,
    vram_metrics,
    output_text,
    error_message,
):
    efficiency_score = calculate_efficiency_score(classification["TPS"], vram_metrics["VRAM_Peak_MiB"])
    return {
        "Run_ID": run_id,
        "Status": classification["Status"],
        "Capability": config.get("capability", "chat"),
        "Output_Category": classification["Output_Category"],
        "Diagnosis": classification["Diagnosis"],
        "Finish_Reason": classification["Finish_Reason"],
        "Backend": config["backend"],
        "Model": model,
        "Params": param_set.copy(),
        "Applied_Params": applied_params.copy(),
        "Config_Str": display_params,
        "TPS": classification["TPS"],
        "TTFT": classification["TTFT"],
        "First_Event_s": classification["First_Event_s"],
        "Stream_Duration_s": classification["Stream_Duration_s"],
        "Prompt_Tokens": classification.get("Prompt_Tokens"),
        "Prefill_Time_s": classification.get("Prefill_Time_s"),
        "Prefill_TPS": classification.get("Prefill_TPS"),
        "Total_Chunks": classification["Total_Chunks"],
        "Content_Chunks": classification["Content_Chunks"],
        "Non_Content_Chunks": classification["Non_Content_Chunks"],
        "Non_Content_Types": classification["Non_Content_Types"],
        "VRAM_Base_MiB": vram_metrics["VRAM_Base_MiB"],
        "VRAM_Peak_MiB": vram_metrics["VRAM_Peak_MiB"],
        "VRAM_Delta_MiB": vram_metrics["VRAM_Delta_MiB"],
        "VRAM_Detail": vram_metrics["VRAM_Detail"],
        "Efficiency_Score": efficiency_score,
        "Output_Chars": len(output_text),
        "Output_Text": output_text,
        "Error": error_message or "",
    }


def filter_eligible_results(df, capability="chat"):
    success_categories = {"tool_call"} if capability == "tools" else {"normal_content"}
    return df[(df["Status"] == "ok") & (df["Output_Category"].isin(success_categories))].copy()


def build_outcome_summary_dataframe(df):
    outcome_df = (
        df["Output_Category"]
        .fillna("unknown")
        .value_counts(dropna=False)
        .rename_axis("Output Category")
        .reset_index(name="Count")
    )
    return outcome_df


def build_tool_call_success_summary_dataframe(df):
    if df.empty or "Model" not in df.columns:
        return pd.DataFrame(
            columns=[
                "Model",
                "Total Runs",
                "Tool Call Success Count",
                "Tool Call Success Probability",
            ]
        )

    summary_rows = []
    for model, group in df.groupby("Model", dropna=False, sort=True):
        total_runs = int(len(group))
        success_count = int(
            ((group["Status"] == "ok") & (group["Output_Category"] == "tool_call")).sum()
        )
        success_probability = success_count / total_runs if total_runs else None
        summary_rows.append(
            {
                "Model": model,
                "Total Runs": total_runs,
                "Tool Call Success Count": success_count,
                "Tool Call Success Probability": format_probability_value(success_probability),
            }
        )

    return pd.DataFrame(summary_rows)


def wrap_markdown_table_headers(df):
    header_map = {
        "System Prompt": "System Prompt<br>Variant",
        "Thinking Mode": "Thinking Mode<br>State",
        "Output Category": "Output<br>Category",
        "Finish Reason": "Finish<br>Reason",
        "Prompt Tokens": "Prompt<br>Tokens",
        "Prefill TPS (tok/s)": "Prefill TPS<br>(tok/s)",
        "Total Output (chars)": "Total Output<br>(chars)",
        "Total Output Time (s)": "Total Output Time<br>(s)",
        "TPS (chunk/s)": "TPS<br>(chunk/s)",
        "Thinking TPS (tok/s)": "Thinking TPS<br>(tok/s)",
        "Output TPS (tok/s)": "Output TPS<br>(tok/s)",
        "Output/Thinking Ratio": "Output/Thinking<br>Ratio",
        "TTFT (s)": "TTFT<br>(s)",
        "First Event (s)": "First Event<br>(s)",
        "VRAM Peak (MiB)": "VRAM Peak<br>(MiB)",
        "Efficiency Score (TPS/GiB Peak)": "Efficiency Score<br>(TPS/GiB Peak)",
        "Chunks (content/total)": "Chunks<br>(content/total)",
        "Total Runs": "Total<br>Runs",
        "Tool Call Success Count": "Tool Call Success<br>Count",
        "Tool Call Success Probability": "Tool Call Success<br>Probability",
    }
    return df.rename(columns=header_map)


def summarize_vram_samples(samples):
    if not samples:
        return empty_vram_metrics()

    baseline_snapshot = samples[0]
    total_base_mib = sum(item["memory_used_mib"] for item in baseline_snapshot)
    total_peak_mib = max(sum(item["memory_used_mib"] for item in snapshot) for snapshot in samples)

    per_gpu = {}
    for item in baseline_snapshot:
        per_gpu[item["index"]] = {
            "index": item["index"],
            "name": item["name"],
            "memory_total_mib": item["memory_total_mib"],
            "base_mib": item["memory_used_mib"],
            "peak_mib": item["memory_used_mib"],
        }

    for snapshot in samples:
        for item in snapshot:
            state = per_gpu.setdefault(
                item["index"],
                {
                    "index": item["index"],
                    "name": item["name"],
                    "memory_total_mib": item["memory_total_mib"],
                    "base_mib": item["memory_used_mib"],
                    "peak_mib": item["memory_used_mib"],
                },
            )
            state["peak_mib"] = max(state["peak_mib"], item["memory_used_mib"])
            state["memory_total_mib"] = item["memory_total_mib"]
            state["name"] = item["name"]

    detail_parts = []
    for gpu_index in sorted(per_gpu):
        gpu = per_gpu[gpu_index]
        delta_mib = gpu["peak_mib"] - gpu["base_mib"]
        detail_parts.append(
            (
                f"GPU {gpu['index']} {gpu['name']}: "
                f"{gpu['base_mib']} -> {gpu['peak_mib']} / {gpu['memory_total_mib']} MiB "
                f"(+{delta_mib} MiB)"
            )
        )

    return {
        "VRAM_Base_MiB": total_base_mib,
        "VRAM_Peak_MiB": total_peak_mib,
        "VRAM_Delta_MiB": total_peak_mib - total_base_mib,
        "VRAM_Detail": " | ".join(detail_parts) if detail_parts else "N/A",
    }


class NvidiaVRAMMonitor:
    def __init__(self, interval_seconds=0.2):
        self.interval_seconds = interval_seconds
        self.samples = []
        self._stop_event = threading.Event()
        self._thread = None
        self.enabled = False

    def start(self):
        baseline_snapshot = query_nvidia_vram_snapshot()
        if not baseline_snapshot:
            return False

        self.samples = [baseline_snapshot]
        self.enabled = True
        self._thread = threading.Thread(target=self._poll_loop, daemon=True)
        self._thread.start()
        return True

    def _poll_loop(self):
        while not self._stop_event.wait(self.interval_seconds):
            snapshot = query_nvidia_vram_snapshot()
            if snapshot:
                self.samples.append(snapshot)

    def stop(self):
        if not self.enabled:
            return empty_vram_metrics()

        self._stop_event.set()
        if self._thread is not None:
            self._thread.join(timeout=1)

        final_snapshot = query_nvidia_vram_snapshot()
        if final_snapshot:
            self.samples.append(final_snapshot)

        return summarize_vram_samples(self.samples)


def build_summary_dataframe(df):
    summary_source = df.copy()
    for column_name in (
        "Prompt_Tokens",
        "Thinking_Tokens",
        "Answer_Tokens",
        "Completion_Tokens",
        "Total_Tokens",
        "Token_Count_Source",
        "Prefill_TPS",
        "Output_Chars",
        "Output_Time_s",
        "Thinking_Time_s",
        "Answer_Time_s",
        "Thinking_TPS",
        "Output_TPS",
        "Output_Thinking_Ratio",
    ):
        if column_name not in summary_source.columns:
            summary_source[column_name] = None

    summary_columns = [
        "Run_ID",
        "Status",
        "Output_Category",
        "Model",
        "Finish_Reason",
        "Config_Str",
        "Prompt_Tokens",
        "Thinking_Tokens",
        "Answer_Tokens",
        "Completion_Tokens",
        "Total_Tokens",
        "Token_Count_Source",
        "Prefill_TPS",
        "Output_Chars",
        "Output_Time_s",
        "Thinking_Time_s",
        "Answer_Time_s",
        "TPS",
        "Thinking_TPS",
        "Output_TPS",
        "Output_Thinking_Ratio",
        "TTFT",
        "First_Event_s",
        "Content_Chunks",
        "Total_Chunks",
        "VRAM_Peak_MiB",
        "Efficiency_Score",
    ]
    if "Capability" in df.columns:
        summary_columns.insert(2, "Capability")
    has_question_metadata = (
        "Question_ID" in df.columns
        and df["Question_ID"].fillna("").astype(str).str.strip().ne("").any()
    )
    if has_question_metadata:
        question_columns = [
            "Suite_ID",
            "Question_ID",
            "Question_Category",
            "Question_Title",
        ]
        for column_name in question_columns:
            if column_name not in summary_source.columns:
                summary_source[column_name] = ""
        insert_at = summary_columns.index("Capability") + 1 if "Capability" in summary_columns else 2
        for column_name in reversed(question_columns):
            summary_columns.insert(insert_at, column_name)
    if "System_Prompt_Label" in df.columns:
        summary_columns.insert(summary_columns.index("Model") + 1, "System_Prompt_Label")
    if "Thinking_Mode" in df.columns:
        thinking_mode_insert_at = summary_columns.index("Model") + 1
        if "System_Prompt_Label" in summary_columns:
            thinking_mode_insert_at = summary_columns.index("System_Prompt_Label") + 1
        summary_columns.insert(thinking_mode_insert_at, "Thinking_Mode")

    summary_df = summary_source[summary_columns].copy()
    summary_df["Finish_Reason"] = summary_df["Finish_Reason"].apply(format_text_value)
    if "Thinking_Mode" in summary_df.columns:
        summary_df["Thinking_Mode"] = summary_df["Thinking_Mode"].apply(resolve_thinking_mode)
    summary_df["Prompt_Tokens"] = summary_df["Prompt_Tokens"].apply(format_token_count)
    for token_column in (
        "Thinking_Tokens",
        "Answer_Tokens",
        "Completion_Tokens",
        "Total_Tokens",
    ):
        summary_df[token_column] = summary_df[token_column].apply(format_token_count)
    summary_df["Token_Count_Source"] = summary_df["Token_Count_Source"].apply(format_text_value)
    summary_df["Prefill_TPS"] = summary_df["Prefill_TPS"].apply(
        lambda value: format_numeric_value(value, 2)
    )
    summary_df["TPS"] = summary_df["TPS"].apply(lambda value: format_numeric_value(value, 2))
    summary_df["Thinking_TPS"] = summary_df["Thinking_TPS"].apply(
        lambda value: format_numeric_value(value, 2)
    )
    summary_df["Output_TPS"] = summary_df["Output_TPS"].apply(
        lambda value: format_numeric_value(value, 2)
    )
    summary_df["Output_Thinking_Ratio"] = summary_df["Output_Thinking_Ratio"].apply(
        lambda value: format_numeric_value(value, 3)
    )
    summary_df["Output_Chars"] = summary_df["Output_Chars"].apply(
        lambda value: "N/A" if pd.isna(value) else int(value)
    )
    summary_df["Output_Time_s"] = summary_df["Output_Time_s"].apply(
        lambda value: format_numeric_value(value, 3)
    )
    summary_df["Thinking_Time_s"] = summary_df["Thinking_Time_s"].apply(
        lambda value: format_numeric_value(value, 3)
    )
    summary_df["Answer_Time_s"] = summary_df["Answer_Time_s"].apply(
        lambda value: format_numeric_value(value, 3)
    )
    summary_df["TTFT"] = summary_df["TTFT"].apply(lambda value: format_numeric_value(value, 3))
    summary_df["First_Event_s"] = summary_df["First_Event_s"].apply(
        lambda value: format_numeric_value(value, 3)
    )
    summary_df["VRAM_Peak_MiB"] = summary_df["VRAM_Peak_MiB"].apply(
        lambda value: "N/A" if pd.isna(value) else int(value)
    )
    summary_df["Efficiency_Score"] = summary_df["Efficiency_Score"].apply(
        lambda value: format_numeric_value(value, 3)
    )
    summary_df["Chunks (content/total)"] = summary_df.apply(
        lambda row: f"{int(row['Content_Chunks'])}/{int(row['Total_Chunks'])}",
        axis=1,
    )
    summary_df = summary_df.drop(columns=["Content_Chunks", "Total_Chunks"])
    return summary_df.rename(
        columns={
            "Run_ID": "Run",
            "Capability": "Capability",
            "Suite_ID": "Suite ID",
            "Question_ID": "Question ID",
            "Question_Category": "Question Category",
            "Question_Title": "Question Title",
            "Output_Category": "Output Category",
            "System_Prompt_Label": "System Prompt",
            "Thinking_Mode": "Thinking Mode",
            "Finish_Reason": "Finish Reason",
            "Config_Str": "Config",
            "Prompt_Tokens": "Prompt Tokens",
            "Thinking_Tokens": "Thinking Tokens",
            "Answer_Tokens": "Answer Tokens",
            "Completion_Tokens": "Completion Tokens",
            "Total_Tokens": "Total Tokens",
            "Token_Count_Source": "Token Count Source",
            "Prefill_TPS": "Prefill TPS (tok/s)",
            "Output_Chars": "Total Output (chars)",
            "Output_Time_s": "Total Output Time (s)",
            "Thinking_Time_s": "Thinking Time (s)",
            "Answer_Time_s": "Answer Time (s)",
            "TPS": "TPS (chunk/s)",
            "Thinking_TPS": "Thinking TPS (tok/s)",
            "Output_TPS": "Output TPS (tok/s)",
            "Output_Thinking_Ratio": "Output/Thinking Ratio",
            "TTFT": "TTFT (s)",
            "First_Event_s": "First Event (s)",
            "VRAM_Peak_MiB": "VRAM Peak (MiB)",
            "Efficiency_Score": "Efficiency Score (TPS/GiB Peak)",
        }
    )


def dataframe_to_text_table(df):
    try:
        return df.to_markdown(index=False)
    except Exception:
        return df.to_string(index=False)


def dataframe_to_report_table(df):
    try:
        if any("<br>" in str(column_name) for column_name in df.columns):
            return df.to_html(index=False, escape=False, border=0)
        return df.to_markdown(index=False)
    except Exception:
        return df.to_string(index=False)


def html_escape_text(value):
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except TypeError:
        pass
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False)
    return html.escape(str(value))


def build_model_metric_charts_html(df):
    """Render portable, dependency-free model comparison bars for HTML reports."""
    if df is None or df.empty or "Model" not in df.columns:
        return '<p class="empty-note">No model metrics are available for comparison. / 沒有可比較的模型指標。</p>'

    source = df.copy()
    if "Status" in source.columns:
        ok_source = source[source["Status"] == "ok"]
        if not ok_source.empty:
            source = ok_source

    metric_specs = (
        ("Output_TPS", "Output TPS / 回覆速率", "tok/s", False),
        ("TTFT", "TTFT / 首字延遲", "s", True),
        ("VRAM_Peak_MiB", "VRAM Peak / 顯存峰值", "MiB", True),
        ("Efficiency_Score", "Efficiency Score / 效率分數", "TPS/GiB Peak", False),
    )
    palette = ("#bd5d38", "#2d6f73", "#9a7b2f", "#6d5a86", "#55733f", "#a04d61")
    chart_cards = []

    for column_name, title, unit, lower_is_better in metric_specs:
        if column_name not in source.columns:
            continue
        values = []
        for model, group in source.groupby("Model", sort=True):
            numeric_values = pd.to_numeric(group[column_name], errors="coerce").dropna()
            numeric_values = numeric_values[numeric_values.map(lambda value: math.isfinite(float(value)))]
            if not numeric_values.empty:
                values.append((str(model), float(numeric_values.mean())))
        if not values:
            continue

        scale_max = max(value for _, value in values) or 1.0
        rows = []
        for index, (model, value) in enumerate(values):
            width = max(1.5, min(100.0, value / scale_max * 100.0)) if value > 0 else 0.0
            rows.append(
                '<div class="metric-bar-row">'
                f'<div class="metric-bar-label" title="{html.escape(model)}">{html.escape(model)}</div>'
                '<div class="metric-bar-track">'
                f'<div class="metric-bar-fill" style="width:{width:.2f}%;background:{palette[index % len(palette)]}"></div>'
                '</div>'
                f'<div class="metric-bar-value">{value:.2f}</div>'
                '</div>'
            )
        direction = "Lower is better / 越低越好" if lower_is_better else "Higher is better / 越高越好"
        chart_cards.append(
            '<article class="metric-chart">'
            f'<h3>{html.escape(title)}</h3>'
            f'<p>{html.escape(unit)} · {html.escape(direction)}</p>'
            f'{"".join(rows)}'
            '</article>'
        )

    if not chart_cards:
        return '<p class="empty-note">No model metrics are available for comparison. / 沒有可比較的模型指標。</p>'
    return '<div class="metric-chart-grid">' + "".join(chart_cards) + "</div>"


def html_escape_header(value):
    return html.escape(str(value)).replace("&lt;br&gt;", "<br>")


def dataframe_to_html_table(df, table_class="report-table", empty_message="No data"):
    columns = [str(column_name) for column_name in getattr(df, "columns", [])]
    parts = [f'<table class="{table_class}">', "<thead><tr>"]

    if columns:
        for column_name in columns:
            parts.append(f"<th>{html_escape_header(column_name)}</th>")
    else:
        parts.append("<th>Value</th>")

    parts.append("</tr></thead><tbody>")

    if df is None or df.empty:
        parts.append(
            f'<tr><td class="empty-cell" colspan="{max(len(columns), 1)}">{html_escape_text(empty_message)}</td></tr>'
        )
    else:
        for _, row in df.iterrows():
            parts.append("<tr>")
            for column_name in columns:
                parts.append(f"<td>{html_escape_text(row[column_name])}</td>")
            parts.append("</tr>")

    parts.append("</tbody></table>")
    return "".join(parts)


def make_excel_friendly_dataframe(df):
    excel_df = df.copy()
    excel_df.columns = [re.sub(r"<br\s*/?>", "\n", str(column)) for column in excel_df.columns]
    return excel_df


def style_excel_worksheet(worksheet):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    worksheet.freeze_panes = "A2"
    header_fill = PatternFill(fill_type="solid", fgColor="F6EFE2")

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = header_fill

    for column_cells in worksheet.columns:
        lengths = []
        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            lines = value.splitlines() or [value]
            lengths.append(max(len(line) for line in lines))
            if cell.row > 1:
                cell.alignment = Alignment(vertical="top")
        column_width = min(max(lengths) + 2 if lengths else 12, 42)
        worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = max(column_width, 12)


def save_summary_excel_workbook(df, config, report_stem):
    capability = config.get("capability", "chat")
    workbook_path = Path(f"{report_stem}_summary.xlsx")
    summary_df = make_excel_friendly_dataframe(localize_report_dataframe(build_summary_dataframe(df)))
    outcome_summary_df = make_excel_friendly_dataframe(
        localize_report_dataframe(build_outcome_summary_dataframe(df))
    )
    suite_questions_df = make_excel_friendly_dataframe(
        localize_report_dataframe(build_suite_questions_dataframe(config))
    )
    question_statistics_df = make_excel_friendly_dataframe(
        localize_report_dataframe(build_question_statistics_dataframe(df))
    )
    tool_call_success_summary_df = (
        make_excel_friendly_dataframe(
            localize_report_dataframe(build_tool_call_success_summary_dataframe(df))
        )
        if capability == "tools"
        else pd.DataFrame()
    )

    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        summary_df.to_excel(writer, sheet_name="Summary", index=False)
        outcome_summary_df.to_excel(writer, sheet_name="Outcome Summary", index=False)
        if not suite_questions_df.empty:
            suite_questions_df.to_excel(writer, sheet_name="Suite Questions", index=False)
        if not question_statistics_df.empty:
            question_statistics_df.to_excel(writer, sheet_name="Question Stats", index=False)
        if capability == "tools" and not tool_call_success_summary_df.empty:
            tool_call_success_summary_df.to_excel(writer, sheet_name="Tool Call Success", index=False)

        for worksheet in writer.book.worksheets:
            style_excel_worksheet(worksheet)

    return workbook_path


def build_download_button(href, label):
    if not href:
        return ""
    return (
        f'<a class="download-button" href="{html_escape_text(href)}" download>'
        f"{html_escape_text(label)}</a>"
    )


def key_value_rows_to_html_table(rows, table_class="kv-table"):
    frame = pd.DataFrame(rows, columns=["Field", "Value"])
    return dataframe_to_html_table(frame, table_class=table_class)


def bullet_list_to_html(items, list_class="note-list"):
    parts = [f'<ul class="{list_class}">']
    for item in items:
        parts.append(f"<li>{html_escape_text(item)}</li>")
    parts.append("</ul>")
    return "".join(parts)


def bilingual_text(english, chinese):
    return f"{english} / {chinese}"


STATUS_BILINGUAL_MAP = {
    "ok": bilingual_text("ok", "正常"),
    "warning": bilingual_text("warning", "警告"),
    "error": bilingual_text("error", "錯誤"),
}

CAPABILITY_BILINGUAL_MAP = {
    "chat": bilingual_text("chat", "對話"),
    "tools": bilingual_text("tools", "工具呼叫"),
    "suite-smoke-7": bilingual_text("suite-smoke-7", "七項能力冒煙套裝"),
}

QUESTION_CATEGORY_BILINGUAL_MAP = {
    "math": bilingual_text("math", "數學"),
    "logic": bilingual_text("logic", "邏輯"),
    "reasoning": bilingual_text("reasoning", "推理"),
    "reading": bilingual_text("reading", "閱讀"),
    "translation": bilingual_text("translation", "翻譯"),
    "writing": bilingual_text("writing", "寫作"),
    "coding": bilingual_text("coding", "程式"),
}

OUTPUT_CATEGORY_BILINGUAL_MAP = {
    "normal_content": bilingual_text("normal_content", "正常文字輸出"),
    "tool_call": bilingual_text("tool_call", "工具呼叫"),
    "text_reply_without_tool": bilingual_text("text_reply_without_tool", "有文字回覆但未呼叫工具"),
    "empty_reply": bilingual_text("empty_reply", "空回覆"),
    "non_content_stream": bilingual_text("non_content_stream", "僅非文字串流"),
    "early_stop": bilingual_text("early_stop", "提早中斷"),
}

FINISH_REASON_BILINGUAL_MAP = {
    "stop": bilingual_text("stop", "正常結束"),
    "tool_calls": bilingual_text("tool_calls", "工具呼叫結束"),
    "length": bilingual_text("length", "長度上限"),
}

THINKING_MODE_BILINGUAL_MAP = {
    "enable": bilingual_text("enable", "啟用"),
    "disable": bilingual_text("disable", "停用"),
    "default": bilingual_text("default", "依後端預設"),
}

REPORT_HEADER_BILINGUAL_MAP = {
    "Run": "Run<br>執行編號",
    "Status": "Status<br>狀態",
    "Capability": "Capability<br>能力模式",
    "Suite ID": "Suite ID<br>套裝識別碼",
    "Suite Version": "Suite Version<br>套裝版本",
    "Question ID": "Question ID<br>題目識別碼",
    "Question Category": "Question Category<br>題目分類",
    "Question Title": "Question Title<br>題目名稱",
    "Prompt": "Prompt<br>題目內容",
    "Expected Output": "Expected Output<br>預期輸出",
    "Evaluation Guide": "Evaluation Guide<br>評估提示",
    "Output Category": "Output Category<br>輸出分類",
    "Model": "Model<br>模型",
    "System Prompt": "System Prompt<br>系統提示",
    "Thinking Mode": "Thinking Mode<br>思考模式",
    "Finish Reason": "Finish Reason<br>結束原因",
    "Config": "Config<br>測試設定",
    "Prompt Tokens": "Prompt Tokens<br>提示詞 Token 數",
    "Thinking Tokens": "Thinking Tokens<br>思考 Token 數",
    "Answer Tokens": "Answer Tokens<br>回答 Token 數",
    "Completion Tokens": "Completion Tokens<br>完成 Token 數",
    "Total Tokens": "Total Tokens<br>總 Token 數",
    "Token Count Source": "Token Count Source<br>Token 計數來源",
    "Prefill TPS (tok/s)": "Prefill TPS<br>(tok/s)<br>預填充速率",
    "Total Output (chars)": "Total Output<br>(chars)<br>總輸出字數",
    "Total Output Time (s)": "Total Output Time<br>(s)<br>總輸出時間",
    "Thinking Time (s)": "Thinking Time<br>(s)<br>思考時間",
    "Answer Time (s)": "Answer Time<br>(s)<br>回答時間",
    "Avg Thinking Time (s)": "Avg Thinking Time<br>(s)<br>平均思考時間",
    "Avg Answer Time (s)": "Avg Answer Time<br>(s)<br>平均回答時間",
    "Runs": "Runs<br>執行次數",
    "TPS (chunk/s)": "TPS<br>(chunk/s)<br>輸出速率",
    "Thinking TPS (tok/s)": "Thinking TPS<br>(tok/s)<br>思考速率",
    "Output TPS (tok/s)": "Output TPS<br>(tok/s)<br>回覆速率",
    "Output/Thinking Ratio": "Output/Thinking<br>Ratio<br>輸出思考比",
    "TTFT (s)": "TTFT<br>(s)<br>首字延遲",
    "First Event (s)": "First Event<br>(s)<br>首事件延遲",
    "VRAM Peak (MiB)": "VRAM Peak<br>(MiB)<br>顯存峰值",
    "Efficiency Score (TPS/GiB Peak)": "Efficiency Score<br>(TPS/GiB Peak)<br>效率分數",
    "Chunks (content/total)": "Chunks<br>(content/total)<br>片段數",
    "Count": "Count<br>次數",
    "Total Runs": "Total Runs<br>總測試次數",
    "Tool Call Success Count": "Tool Call Success<br>Count<br>工具成功次數",
    "Tool Call Success Probability": "Tool Call Success<br>Probability<br>工具成功率",
}


def localize_status_value(value):
    return STATUS_BILINGUAL_MAP.get(value, value)


def localize_capability_value(value):
    return CAPABILITY_BILINGUAL_MAP.get(value, value)


def localize_output_category_value(value):
    return OUTPUT_CATEGORY_BILINGUAL_MAP.get(value, value)


def localize_finish_reason_value(value):
    if value in (None, "", "N/A"):
        return bilingual_text("N/A", "無")
    return FINISH_REASON_BILINGUAL_MAP.get(value, value)


def localize_system_prompt_label(value):
    if value in (None, "", "N/A"):
        return bilingual_text("N/A", "未使用額外 system prompt")
    return bilingual_text(str(value), "系統提示變體")


def localize_thinking_mode_value(value):
    normalized = resolve_thinking_mode(value)
    return THINKING_MODE_BILINGUAL_MAP.get(normalized, value)


def localize_report_dataframe(df):
    localized_df = df.copy()
    if "Status" in localized_df.columns:
        localized_df["Status"] = localized_df["Status"].apply(localize_status_value)
    if "Capability" in localized_df.columns:
        localized_df["Capability"] = localized_df["Capability"].apply(localize_capability_value)
    if "Question Category" in localized_df.columns:
        localized_df["Question Category"] = localized_df["Question Category"].apply(
            lambda value: QUESTION_CATEGORY_BILINGUAL_MAP.get(value, value)
        )
    if "Output Category" in localized_df.columns:
        localized_df["Output Category"] = localized_df["Output Category"].apply(
            localize_output_category_value
        )
    if "Finish Reason" in localized_df.columns:
        localized_df["Finish Reason"] = localized_df["Finish Reason"].apply(
            localize_finish_reason_value
        )
    if "System Prompt" in localized_df.columns:
        localized_df["System Prompt"] = localized_df["System Prompt"].apply(
            localize_system_prompt_label
        )
    if "Thinking Mode" in localized_df.columns:
        localized_df["Thinking Mode"] = localized_df["Thinking Mode"].apply(
            localize_thinking_mode_value
        )
    return localized_df.rename(columns=REPORT_HEADER_BILINGUAL_MAP)


def serialize_result_value(value):
    if isinstance(value, (dict, list)):
        return value
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except TypeError:
        pass
    if hasattr(value, "item") and callable(value.item):
        try:
            return value.item()
        except (TypeError, ValueError):
            pass
    return value


def save_raw_outputs(df, report_stem):
    output_path = Path(f"{report_stem}_outputs.jsonl")
    with output_path.open("w", encoding="utf-8") as file:
        for _, row in df.iterrows():
            payload = {column: serialize_result_value(value) for column, value in row.items()}
            file.write(json.dumps(payload, ensure_ascii=False) + "\n")
    return output_path


def ensure_report_output_dir(base_dir="."):
    report_dir = Path(base_dir) / "Report"
    report_dir.mkdir(parents=True, exist_ok=True)
    return report_dir


def select_plot_dataframe(df, capability="chat"):
    if capability == "tools":
        return df.copy(), False

    eligible_df = filter_eligible_results(df, capability=capability)
    if eligible_df.empty:
        return df.copy(), True
    return eligible_df.copy(), False


def normalize_output_text(text):
    normalized = (text or "").replace("\r\n", "\n").strip()
    return normalized or "[No text returned]"


def pause_before_exit():
    if not sys.stdin.isatty():
        return

    try:
        if msvcrt is not None:
            print("\n按任意鍵結束...", end="", flush=True)
            msvcrt.getch()
            print()
        else:
            input("\n按 Enter 結束...")
    except (EOFError, KeyboardInterrupt):
        pass


def get_installed_version(package_name):
    try:
        return get_package_version(package_name)
    except PackageNotFoundError:
        return None


def show_windows_message_dialog(title, message, style=0):
    if sys.platform != "win32":
        return False

    try:
        import ctypes

        ctypes.windll.user32.MessageBoxW(None, message, title, style)
        return True
    except Exception:
        return False


def show_windows_error_dialog(title, message):
    return show_windows_message_dialog(title, message, style=0x10)


def show_windows_info_dialog(title, message):
    return show_windows_message_dialog(title, message, style=0x40)


def ensure_runtime_ready(require_questionary=False):
    dependency_errors = dict(DEPENDENCY_IMPORT_ERRORS)
    if not require_questionary:
        dependency_errors.pop("questionary", None)

    if not dependency_errors:
        return

    guidance = {
        "openai": "請執行 `pip install -U openai`，並確認版本為 1.x 以上。",
        "questionary": "請執行 `pip install -U questionary prompt_toolkit`。",
        "matplotlib": "請執行 `pip install -U matplotlib`。",
        "pandas": "請執行 `pip install -U pandas`。",
        "requests": "請執行 `pip install -U requests`。",
    }
    lines = [
        "程式啟動前檢查失敗。",
        "這通常代表另一台電腦雖然有安裝套件，但版本不相容或安裝不完整。",
        "",
    ]
    for package_name, exc in dependency_errors.items():
        installed_version = get_installed_version(package_name)
        version_label = f"已安裝 {installed_version}" if installed_version else "未偵測到安裝版本"
        hint = guidance.get(package_name, f"請重新安裝 `{package_name}`。")
        lines.append(f"- {package_name} ({version_label}): {exc}. {hint}")

    raise RuntimeError("\n".join(lines))


def persist_crash_log(error_text):
    crash_log_path = Path("llm_expert_bench_crash.log")
    crash_log_path.write_text(error_text, encoding="utf-8")
    return crash_log_path


def persist_ui_launch_hint(url, report_root):
    launch_hint_path = Path(UI_LAUNCH_HINT_FILENAME)
    launch_hint_path.write_text(
        "\n".join(
            [
                "DIY LLM Benchmark UI",
                f"URL: {url}",
                f"Report directory: {report_root}",
                f"Started at: {time.strftime('%Y-%m-%d %H:%M:%S')}",
                "",
                "If the browser did not open automatically, copy the URL above into your browser.",
            ]
        ),
        encoding="utf-8",
    )
    return launch_hint_path


def build_suite_questions_dataframe(config):
    capability = config.get("capability", "chat")
    if capability not in BUILTIN_SUITES:
        return pd.DataFrame()

    questions = resolve_benchmark_questions(config)
    return pd.DataFrame(
        [
            {
                "Suite ID": question["suite_id"],
                "Suite Version": question["suite_version"],
                "Question ID": question["id"],
                "Question Category": question["category"],
                "Question Title": question["title"],
                "Prompt": question["prompt"],
                "Expected Output": question["expected_output"],
                "Evaluation Guide": question["evaluation_guide"],
            }
            for question in questions
        ]
    )


def build_question_statistics_dataframe(df):
    if "Question_ID" not in df.columns:
        return pd.DataFrame()

    question_df = df[df["Question_ID"].fillna("").astype(str).str.strip().ne("")].copy()
    if question_df.empty:
        return pd.DataFrame()

    group_columns = [
        "Suite_ID",
        "Question_ID",
        "Question_Category",
        "Question_Title",
        "Model",
    ]
    for column_name in group_columns:
        if column_name not in question_df.columns:
            question_df[column_name] = ""
    for column_name in (
        "Prompt_Tokens",
        "Thinking_Tokens",
        "Answer_Tokens",
        "Completion_Tokens",
        "Total_Tokens",
        "Thinking_Time_s",
        "Answer_Time_s",
    ):
        if column_name not in question_df.columns:
            question_df[column_name] = None

    records = []
    for group_values, group in question_df.groupby(group_columns, dropna=False, sort=False):
        suite_id, question_id, category, title, model = group_values

        def sum_tokens(column_name):
            values = pd.to_numeric(group[column_name], errors="coerce")
            return int(values.fillna(0).sum())

        def average_seconds(column_name):
            values = pd.to_numeric(group[column_name], errors="coerce").dropna()
            return round(float(values.mean()), 3) if not values.empty else None

        token_sources = sorted(
            {
                str(value).strip()
                for value in group.get("Token_Count_Source", pd.Series(dtype=str)).dropna()
                if str(value).strip()
            }
        )
        records.append(
            {
                "Suite ID": suite_id,
                "Question ID": question_id,
                "Question Category": category,
                "Question Title": title,
                "Model": model,
                "Runs": len(group),
                "Prompt Tokens": sum_tokens("Prompt_Tokens"),
                "Thinking Tokens": sum_tokens("Thinking_Tokens"),
                "Answer Tokens": sum_tokens("Answer_Tokens"),
                "Completion Tokens": sum_tokens("Completion_Tokens"),
                "Total Tokens": sum_tokens("Total_Tokens"),
                "Avg Thinking Time (s)": average_seconds("Thinking_Time_s"),
                "Avg Answer Time (s)": average_seconds("Answer_Time_s"),
                "Token Count Source": ", ".join(token_sources) or "N/A",
            }
        )
    return pd.DataFrame(records)


def build_ui_launch_message(url, report_root, launch_hint_path, browser_opened):
    lines = [
        "DIY LLM Benchmark local UI is ready.",
        url,
        "",
        f"Report directory: {report_root}",
        f"Launch hint file: {launch_hint_path.resolve()}",
        "",
    ]
    if browser_opened:
        lines.append("A browser tab was opened automatically if Windows allowed it.")
    else:
        lines.append("Browser auto-open did not succeed. Please open the URL manually.")
    return "\n".join(lines)


def try_open_browser(url):
    browser_error = None
    try:
        if webbrowser.open(url, new=2):
            return True, None
    except Exception as exc:
        browser_error = exc

    if sys.platform == "win32":
        try:
            os.startfile(url)
            return True, None
        except Exception as exc:
            if browser_error is None:
                browser_error = exc

    if browser_error is None:
        browser_error = RuntimeError("No registered browser handler reported success.")
    return False, str(browser_error)


def handle_fatal_error(exc):
    traceback_text = traceback.format_exc()
    if traceback_text.strip() == "NoneType: None":
        traceback_text = f"{type(exc).__name__}: {exc}"

    error_text = "\n".join(
        [
            "程式執行失敗。",
            "建議用 PowerShell 或 CMD 執行 `python llm_expert_bench.py`，比較容易看到完整訊息。",
            "",
            f"{type(exc).__name__}: {exc}",
            "",
            traceback_text,
        ]
    )
    crash_log_path = persist_crash_log(error_text)

    print(error_text, file=sys.stderr)
    print(f"\n錯誤記錄已寫入: {crash_log_path.resolve()}", file=sys.stderr)

    if not sys.stdin.isatty():
        dialog_message = "\n".join(
            [
                "程式執行失敗，錯誤記錄已保存。",
                str(crash_log_path.resolve()),
                "",
                f"{type(exc).__name__}: {exc}",
                "",
                "請改用 PowerShell / CMD 執行，或把 crash log 傳回來。",
            ]
        )
        show_windows_error_dialog("llm_expert_bench 啟動失敗", dialog_message)


def select_models_and_url(backend, previous_url=None, previous_models=None):
    previous_models = previous_models or []

    if backend == "ollama":
        detected_models = get_ollama_models()
        if detected_models:
            choice_list = [
                Choice(model_name, value=model_name, checked=model_name in previous_models)
                for model_name in detected_models
            ]
            selected_models = ask_checkbox_with_back(
                "Select benchmark models / 選擇測試模型:",
                choices=choice_list,
            )
            if selected_models is None:
                return None
            if selected_models == BACK_ACTION:
                return BACK_ACTION
            if selected_models:
                return "http://localhost:11434/v1", selected_models

        while True:
            manual_input = ask_text_with_back(
                "Enter Ollama model names (comma separated) / 請輸入 Ollama 模型名稱（逗號分隔）:",
                default=",".join(previous_models) if previous_models else "qwen3.5:latest",
            )
            if manual_input is None:
                return None
            if manual_input == BACK_ACTION:
                return BACK_ACTION
            models = [name.strip() for name in (manual_input or "").split(",") if name.strip()]
            if models:
                return "http://localhost:11434/v1", models
            print("At least one model is required. / 至少需要一個模型名稱。")

    port_default = "8080"
    if previous_url and previous_url.startswith("http://localhost:") and previous_url.endswith("/v1"):
        port_default = previous_url.removeprefix("http://localhost:").removesuffix("/v1") or "8080"

    while True:
        port = ask_text_with_back(
            "Enter llama-server port / 請輸入 llama-server 端口:",
            default=port_default,
        )
        if port is None:
            return None
        if port == BACK_ACTION:
            return BACK_ACTION

        model_names = ask_text_with_back(
            "Enter loaded model names (comma separated, for labeling only) / "
            "請輸入載入中的模型名稱（逗號分隔，僅供辨識）:",
            default=",".join(previous_models) if previous_models else "llama.cpp-model",
        )
        if model_names is None:
            return None
        if model_names == BACK_ACTION:
            continue

        models = [name.strip() for name in (model_names or "").split(",") if name.strip()]
        if models:
            return f"http://localhost:{port or '8080'}/v1", models
        print("At least one model is required. / 至少需要一個模型名稱。")


def interactive_config():
    print("\n" + "═" * 62)
    print("🏆 LLM 專家參數 Benchmark 工具 V3")
    print("═" * 62)

    backend = questionary.select(
        "請選擇測試後端:",
        choices=[
            Choice("🦙 Ollama", value="ollama"),
            Choice("🏗️ llama.cpp (需先啟動 llama-server)", value="llama.cpp"),
        ],
    ).ask()
    if not backend:
        return None

    url, models = select_models_and_url(backend)
    if not models:
        print("⚠️ 沒有可用模型，已取消。")
        return None

    available_groups = {
        group_name: [key for key in param_keys if backend in PARAM_INFO[key]["backends"]]
        for group_name, param_keys in PARAM_GROUPS.items()
    }
    available_groups = {name: keys for name, keys in available_groups.items() if keys}

    selected_groups = questionary.checkbox(
        "選擇想測試的參數類別 (可不選，代表只比模型預設值):",
        choices=[
            Choice(f"{group_name} ({len(param_keys)} 項)", value=group_name)
            for group_name, param_keys in available_groups.items()
        ],
    ).ask() or []

    final_params = {}
    for group_name in selected_groups:
        param_keys = available_groups[group_name]
        selected_params = questionary.checkbox(
            f"勾選 {group_name} 要測試的參數:",
            choices=[
                Choice(
                    title=(
                        f"{PARAM_INFO[key]['label']} | 範圍: {PARAM_INFO[key]['range']} | "
                        f"{PARAM_INFO[key]['desc']}"
                    ),
                    value=key,
                )
                for key in param_keys
            ],
        ).ask() or []

        for key in selected_params:
            values = ask_param_values(key)
            if values is None:
                return None
            final_params[key] = values

    prompt = questionary.text(
        "測試 Prompt:",
        default="詳細解釋 3D 列印中，PETG 材質發生蠕變 (Creep) 的溫度臨界點。",
    ).ask()
    if prompt is None:
        return None

    return {
        "backend": backend,
        "url": url,
        "models": models,
        "params": final_params,
        "prompt": prompt,
    }


def run_bench(config):
    client = OpenAI(base_url=config["url"], api_key="sk-no-key-needed")
    param_keys = list(config["params"].keys())
    param_values = [config["params"][key] for key in param_keys]
    combos = [dict(zip(param_keys, combo)) for combo in product(*param_values)] if param_keys else [{}]

    results = []
    total_runs = len(config["models"]) * len(combos)
    vram_monitoring_enabled = query_nvidia_vram_snapshot() is not None
    config["vram_monitoring"] = "nvidia-smi" if vram_monitoring_enabled else "unavailable"
    print(f"\n⚡ 啟動測試，共 {total_runs} 組配置。")
    print("📌 TPS 以串流回傳片段估算，適合做相對比較。")
    if vram_monitoring_enabled:
        print("🧠 顯存監控: 已啟用 nvidia-smi 取樣。")
    else:
        print("🧠 顯存監控: 未偵測到 nvidia-smi，報告將顯示 N/A。")

    run_index = 0
    for model in config["models"]:
        for param_set in combos:
            run_index += 1
            applied_params = build_backend_options(config["backend"], param_set)
            display_params = format_param_dict(param_set)
            print(f"[{run_index}/{total_runs}] {model} | {display_params}")

            request_kwargs = {}
            if applied_params:
                request_kwargs["extra_body"] = (
                    {"options": applied_params}
                    if config["backend"] == "ollama"
                    else applied_params
                )

            start_time = time.time()
            first_event_time = None
            first_content_time = None
            output_parts = []
            chunk_records = []
            vram_monitor = NvidiaVRAMMonitor() if vram_monitoring_enabled else None
            vram_metrics = empty_vram_metrics()
            if vram_monitor is not None and not vram_monitor.start():
                vram_monitor = None
            try:
                stream = client.chat.completions.create(
                    model=model,
                    messages=[{"role": "user", "content": config["prompt"]}],
                    stream=True,
                    **request_kwargs,
                )

                for chunk in stream:
                    event_time = time.time()
                    if first_event_time is None:
                        first_event_time = event_time

                    chunk_info = inspect_stream_chunk(chunk)
                    chunk_records.append(chunk_info)

                    if chunk_info["content"] and first_content_time is None:
                        first_content_time = event_time
                    if chunk_info["content"]:
                        output_parts.append(chunk_info["content"])

                end_time = time.time()
                if vram_monitor is not None:
                    vram_metrics = vram_monitor.stop()
                output_text = "".join(output_parts)
                classification = classify_stream_result(
                    chunk_records=chunk_records,
                    start_time=start_time,
                    end_time=end_time,
                    first_event_time=first_event_time,
                    first_content_time=first_content_time,
                    error_message=None,
                )
                results.append(
                    build_result_row(
                        run_id=run_index,
                        config=config,
                        model=model,
                        param_set=param_set,
                        applied_params=applied_params,
                        display_params=display_params,
                        classification=classification,
                        vram_metrics=vram_metrics,
                        output_text=output_text,
                        error_message=None,
                    )
                )
            except Exception as exc:
                end_time = time.time()
                if vram_monitor is not None:
                    vram_metrics = vram_monitor.stop()
                output_text = "".join(output_parts)
                classification = classify_stream_result(
                    chunk_records=chunk_records,
                    start_time=start_time,
                    end_time=end_time,
                    first_event_time=first_event_time,
                    first_content_time=first_content_time,
                    error_message=str(exc),
                )
                print(f"❌ 失敗: {exc}")
                results.append(
                    build_result_row(
                        run_id=run_index,
                        config=config,
                        model=model,
                        param_set=param_set,
                        applied_params=applied_params,
                        display_params=display_params,
                        classification=classification,
                        vram_metrics=vram_metrics,
                        output_text=output_text,
                        error_message=str(exc),
                    )
                )

    return pd.DataFrame(results)


def plot_results(df, output_path, capability="chat"):
    if df.empty:
        return None

    plot_df, used_fallback = select_plot_dataframe(df, capability=capability)
    plot_df = plot_df.copy()
    multi_prompt_mode = (
        "System_Prompt_Label" in plot_df.columns
        and plot_df["System_Prompt_Label"].fillna("N/A").nunique(dropna=False) > 1
    )

    def build_label(row):
        system_prompt_label = row.get("System_Prompt_Label", "N/A")
        if multi_prompt_mode or (system_prompt_label not in ("", "N/A", None)):
            label = f"{row['Model']} | {system_prompt_label} | {row['Config_Str']}"
        else:
            label = f"{row['Model']} | {row['Config_Str']}"
        return label if len(label) <= 42 else label[:39] + "..."

    def build_run_color(row):
        if row["Output_Category"] == "tool_call":
            return "seagreen"
        if row["Status"] == "ok":
            return "skyblue"
        if row["Status"] == "warning":
            return "darkorange"
        return "indianred"

    plot_df["Plot_Label"] = plot_df.apply(build_label, axis=1)
    plot_df["Plot_Color"] = plot_df.apply(build_run_color, axis=1)

    if capability == "tools":
        first_event_values = plot_df["First_Event_s"].fillna(0)
        duration_values = plot_df["Stream_Duration_s"].fillna(0)
        outcome_df = build_outcome_summary_dataframe(plot_df)

        fig, axes = plt.subplots(3, 1, figsize=(14, 15))

        axes[0].bar(plot_df["Plot_Label"], first_event_values, color=plot_df["Plot_Color"])
        if plot_df["First_Event_s"].notna().any():
            axes[0].axhline(
                y=plot_df["First_Event_s"].min(),
                color="green",
                linestyle="--",
                alpha=0.3,
            )
        axes[0].set_title("Tools Benchmark First Event Comparison")
        axes[0].set_ylabel("First Event (s)")

        axes[1].bar(plot_df["Plot_Label"], duration_values, color=plot_df["Plot_Color"])
        if plot_df["Stream_Duration_s"].notna().any():
            axes[1].axhline(
                y=plot_df["Stream_Duration_s"].min(),
                color="green",
                linestyle="--",
                alpha=0.3,
            )
        axes[1].set_title("Tools Benchmark Stream Duration Comparison")
        axes[1].set_ylabel("Stream Duration (s)")

        axes[2].bar(outcome_df["Output Category"], outcome_df["Count"], color="steelblue")
        axes[2].set_title("Tools Outcome Category Counts")
        axes[2].set_ylabel("Runs")
        axes[2].set_xlabel("Output Category")

        axes[0].tick_params(axis="x", rotation=35)
        axes[1].tick_params(axis="x", rotation=35)
        axes[2].tick_params(axis="x", rotation=20)
    else:
        eligible_df = filter_eligible_results(plot_df, capability=capability)
        if not eligible_df.empty:
            max_tps = eligible_df["TPS"].max()
            min_ttft = eligible_df["TTFT"].min()
            colors = ["gold" if tps == max_tps else "skyblue" for tps in eligible_df["TPS"]]
            ttft_colors = [
                "lightgreen" if ttft == min_ttft else "salmon" for ttft in eligible_df["TTFT"]
            ]
            has_vram_data = eligible_df["VRAM_Peak_MiB"].notna().any()
            has_efficiency_data = eligible_df["Efficiency_Score"].notna().any()

            subplot_count = 4 if has_efficiency_data else 3 if has_vram_data else 2
            fig, axes = plt.subplots(subplot_count, 1, figsize=(13, 5 * subplot_count), sharex=True)
            if subplot_count == 1:
                axes = [axes]

            axes[0].bar(eligible_df["Plot_Label"], eligible_df["TPS"], color=colors)
            axes[0].axhline(y=max_tps, color="red", linestyle="--", alpha=0.3)
            axes[0].set_title("Throughput Comparison")
            axes[0].set_ylabel("TPS (chunk/s)")

            axes[1].bar(eligible_df["Plot_Label"], eligible_df["TTFT"], color=ttft_colors)
            axes[1].axhline(y=min_ttft, color="green", linestyle="--", alpha=0.3)
            axes[1].set_title("First Token Latency Comparison")
            axes[1].set_ylabel("TTFT (s)")

            if has_vram_data:
                min_vram_peak = eligible_df["VRAM_Peak_MiB"].min()
                vram_colors = [
                    "lightgreen" if peak == min_vram_peak else "mediumpurple"
                    for peak in eligible_df["VRAM_Peak_MiB"]
                ]
                axes[2].bar(eligible_df["Plot_Label"], eligible_df["VRAM_Peak_MiB"], color=vram_colors)
                axes[2].axhline(y=min_vram_peak, color="green", linestyle="--", alpha=0.3)
                axes[2].set_title("VRAM Peak Comparison")
                axes[2].set_ylabel("VRAM Peak (MiB)")
                if has_efficiency_data:
                    max_efficiency = eligible_df["Efficiency_Score"].max()
                    efficiency_colors = [
                        "gold" if score == max_efficiency else "steelblue"
                        for score in eligible_df["Efficiency_Score"]
                    ]
                    axes[3].bar(
                        eligible_df["Plot_Label"],
                        eligible_df["Efficiency_Score"],
                        color=efficiency_colors,
                    )
                    axes[3].axhline(y=max_efficiency, color="orange", linestyle="--", alpha=0.3)
                    axes[3].set_title("Efficiency Score Comparison")
                    axes[3].set_ylabel("TPS/GiB Peak")
                    axes[3].set_xlabel("Model | Config")
                else:
                    axes[2].set_xlabel("Model | Config")
            else:
                axes[1].set_xlabel("Model | Config")
        else:
            status_df = (
                plot_df["Status"]
                .fillna("unknown")
                .value_counts(dropna=False)
                .rename_axis("Status")
                .reset_index(name="Count")
            )
            outcome_df = build_outcome_summary_dataframe(plot_df)
            fig, axes = plt.subplots(2, 1, figsize=(13, 10))

            axes[0].bar(status_df["Status"], status_df["Count"], color="steelblue")
            axes[0].set_title("Run Status Counts")
            axes[0].set_ylabel("Runs")

            axes[1].bar(outcome_df["Output Category"], outcome_df["Count"], color="slategray")
            title = "Outcome Category Counts"
            if used_fallback:
                title += " (Fallback)"
            axes[1].set_title(title)
            axes[1].set_ylabel("Runs")
            axes[1].set_xlabel("Output Category")
            axes[0].tick_params(axis="x", rotation=20)
            axes[1].tick_params(axis="x", rotation=20)

    plt.xticks(rotation=35, ha="right")
    plt.tight_layout()
    fig.savefig(output_path)
    plt.close(fig)

    return Path(output_path)


def save_markdown_report(df, config, report_stem):
    report_path = Path(f"{report_stem}.md")
    summary_df = build_summary_dataframe(df)
    outcome_summary_df = build_outcome_summary_dataframe(df)

    with report_path.open("w", encoding="utf-8") as file:
        file.write("# Benchmark Report\n\n")
        file.write(f"- Backend: {config['backend']}\n")
        file.write(f"- Base URL: {config['url']}\n")
        file.write(f"- Models: {', '.join(config['models'])}\n")
        file.write(f"- Prompt: {config['prompt']}\n")
        file.write("- Note: TPS is estimated from streaming content chunks for relative comparison.\n\n")
        file.write("## Environment Notes\n\n")
        file.write(f"- VRAM monitoring: {config.get('vram_monitoring', 'unavailable')}\n\n")
        file.write("## Metric Notes\n\n")
        file.write("- `TPS (chunk/s)`: 每秒收到的文字內容 chunk 數，代表輸出吞吐速度；數值越高通常越快。\n")
        file.write("- `TTFT (s)`: Time To First Token，從送出請求到收到第一段文字內容的秒數；數值越低通常越快。\n")
        file.write("- `First Event (s)`: 從送出請求到收到第一個串流事件的秒數，包含 role 或非文字事件。\n")
        file.write("- `VRAM Peak (MiB)`: 測試期間輪詢到的 NVIDIA GPU 總顯存最高占用，用來比較實際壓力。\n")
        file.write("- `Efficiency Score (TPS/GiB Peak)`: `TPS / (VRAM Peak in GiB)`，代表每 1 GiB 峰值顯存換到多少輸出速度；數值越高越划算。\n")
        file.write("- `TPS (chunk/s)` 與 `TTFT (s)` 在沒有任何文字內容輸出時會顯示 `N/A`。\n\n")
        file.write("## Output Diagnosis Notes\n\n")
        file.write("- `empty_reply`: 串流正常結束，但沒有任何文字內容。\n")
        file.write("- `non_content_stream`: 串流有事件，但只有非文字 payload，例如 `tool_calls` 或 `reasoning`。\n")
        file.write("- `early_stop`: 串流在產生任何文字前提早結束，或在前期就被異常中斷。\n\n")
        file.write("## Summary\n\n")
        file.write(summary_df.to_markdown(index=False))
        file.write("\n\n## Outcome Summary\n\n")
        file.write(outcome_summary_df.to_markdown(index=False))
        file.write("\n\n## Generated Outputs\n")

        for _, row in df.iterrows():
            params_json = json.dumps(row["Params"], ensure_ascii=False)
            applied_params_json = json.dumps(row["Applied_Params"], ensure_ascii=False)
            file.write(f"\n### Run {row['Run_ID']}\n\n")
            file.write(f"- Status: {row['Status']}\n")
            file.write(f"- Output Category: {row['Output_Category']}\n")
            file.write(f"- Diagnosis: {row['Diagnosis']}\n")
            file.write(f"- Finish Reason: {format_text_value(row['Finish_Reason'])}\n")
            file.write(f"- Backend: {row['Backend']}\n")
            file.write(f"- Model: {row['Model']}\n")
            file.write(f"- Params: `{params_json}`\n")
            file.write(f"- Applied Params: `{applied_params_json}`\n")
            file.write(f"- TPS: {format_numeric_value(row['TPS'], 2)} chunk/s\n")
            file.write(f"- TTFT: {format_numeric_value(row['TTFT'], 3)} s\n")
            file.write(f"- First Event: {format_numeric_value(row['First_Event_s'], 3)} s\n")
            file.write(f"- Stream Duration: {format_numeric_value(row['Stream_Duration_s'], 3)} s\n")
            file.write(f"- Total Chunks: {int(row['Total_Chunks'])}\n")
            file.write(f"- Content Chunks: {int(row['Content_Chunks'])}\n")
            file.write(f"- Non-Content Chunks: {int(row['Non_Content_Chunks'])}\n")
            file.write(f"- Non-Content Types: {row['Non_Content_Types']}\n")
            file.write(f"- VRAM Base: {format_mib_value(row['VRAM_Base_MiB'])}\n")
            file.write(f"- VRAM Peak: {format_mib_value(row['VRAM_Peak_MiB'])}\n")
            file.write(f"- VRAM Delta: {format_mib_value(row['VRAM_Delta_MiB'])}\n")
            file.write(f"- VRAM Detail: {row['VRAM_Detail']}\n")
            file.write(
                f"- Efficiency Score: "
                f"{format_numeric_value(row['Efficiency_Score'], 3)} TPS/GiB Peak\n"
            )
            file.write(f"- Output Chars: {row['Output_Chars']}\n")
            if row["Error"]:
                file.write(f"- Error: {row['Error']}\n")
            file.write("\n```text\n")
            file.write(normalize_output_text(row["Output_Text"]))
            file.write("\n```\n")

    return report_path


def export_best_config(df, config):
    eligible_df = filter_eligible_results(df)
    if eligible_df.empty:
        print("⚠️ 沒有正常文字輸出的結果可匯出最佳配置。")
        return None

    best_row = eligible_df.loc[eligible_df["TPS"].idxmax()]
    payload = {
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "backend": best_row["Backend"],
        "model": best_row["Model"],
        "params": best_row["Params"],
        "applied_params": best_row["Applied_Params"],
        "status": best_row["Status"],
        "output_category": best_row["Output_Category"],
        "finish_reason": best_row["Finish_Reason"],
        "diagnosis": best_row["Diagnosis"],
        "tps": best_row["TPS"],
        "ttft": best_row["TTFT"],
        "first_event_s": best_row["First_Event_s"],
        "stream_duration_s": best_row["Stream_Duration_s"],
        "vram_base_mib": best_row["VRAM_Base_MiB"],
        "vram_peak_mib": best_row["VRAM_Peak_MiB"],
        "vram_delta_mib": best_row["VRAM_Delta_MiB"],
        "vram_detail": best_row["VRAM_Detail"],
        "efficiency_score_tps_per_gib_peak": best_row["Efficiency_Score"],
        "prompt": config["prompt"],
    }

    with open("best_config.json", "w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=4)

    print("\n" + "⭐" * 18)
    print(f"🏆 性能冠軍: {best_row['Model']}")
    print(f"🚀 最高速度: {format_numeric_value(best_row['TPS'], 2)} TPS")
    print(f"⚙️ 最佳配置: {best_row['Config_Str']}")
    print(f"🧠 顯存峰值: {format_mib_value(best_row['VRAM_Peak_MiB'])}")
    if pd.notna(best_row["Efficiency_Score"]):
        print(f"📊 效率分數: {format_numeric_value(best_row['Efficiency_Score'], 3)} TPS/GiB Peak")
    print("⭐" * 18)
    print("✅ 已保存 best_config.json")

    if best_row["Backend"] == "ollama":
        with open("Ollama_Modelfile_Suggest", "w", encoding="utf-8") as file:
            file.write(f"FROM {best_row['Model']}\n")
            for key, value in best_row["Applied_Params"].items():
                file.write(f"PARAMETER {key} {value}\n")
        print("✅ 已生成 Ollama_Modelfile_Suggest")
    else:
        print("ℹ️ 本次後端不是 Ollama，略過 Modelfile 建議。")

    return payload


def plot_results(df, output_path):
    eligible_df = filter_eligible_results(df)
    if eligible_df.empty:
        return None

    def build_label(row):
        label = f"{row['Model']} | {row['Config_Str']}"
        return label if len(label) <= 42 else label[:39] + "..."

    eligible_df["Plot_Label"] = eligible_df.apply(build_label, axis=1)
    max_tps = eligible_df["TPS"].max()
    min_ttft = eligible_df["TTFT"].min()
    colors = ["gold" if tps == max_tps else "skyblue" for tps in eligible_df["TPS"]]
    ttft_colors = ["lightgreen" if ttft == min_ttft else "salmon" for ttft in eligible_df["TTFT"]]
    has_vram_data = eligible_df["VRAM_Peak_MiB"].notna().any()
    has_efficiency_data = eligible_df["Efficiency_Score"].notna().any()

    subplot_count = 4 if has_efficiency_data else 3 if has_vram_data else 2
    fig, axes = plt.subplots(subplot_count, 1, figsize=(13, 5 * subplot_count), sharex=True)
    if subplot_count == 1:
        axes = [axes]

    axes[0].bar(eligible_df["Plot_Label"], eligible_df["TPS"], color=colors)
    axes[0].axhline(y=max_tps, color="red", linestyle="--", alpha=0.3)
    axes[0].set_title("Throughput Comparison")
    axes[0].set_ylabel("TPS (chunk/s)")

    axes[1].bar(eligible_df["Plot_Label"], eligible_df["TTFT"], color=ttft_colors)
    axes[1].axhline(y=min_ttft, color="green", linestyle="--", alpha=0.3)
    axes[1].set_title("First Token Latency Comparison")
    axes[1].set_ylabel("TTFT (s)")

    if has_vram_data:
        min_vram_peak = eligible_df["VRAM_Peak_MiB"].min()
        vram_colors = [
            "lightgreen" if peak == min_vram_peak else "mediumpurple"
            for peak in eligible_df["VRAM_Peak_MiB"]
        ]
        axes[2].bar(eligible_df["Plot_Label"], eligible_df["VRAM_Peak_MiB"], color=vram_colors)
        axes[2].axhline(y=min_vram_peak, color="green", linestyle="--", alpha=0.3)
        axes[2].set_title("VRAM Peak Comparison")
        axes[2].set_ylabel("VRAM Peak (MiB)")
        if has_efficiency_data:
            max_efficiency = eligible_df["Efficiency_Score"].max()
            efficiency_colors = [
                "gold" if score == max_efficiency else "steelblue"
                for score in eligible_df["Efficiency_Score"]
            ]
            axes[3].bar(
                eligible_df["Plot_Label"],
                eligible_df["Efficiency_Score"],
                color=efficiency_colors,
            )
            axes[3].axhline(y=max_efficiency, color="orange", linestyle="--", alpha=0.3)
            axes[3].set_title("Efficiency Score Comparison")
            axes[3].set_ylabel("TPS/GiB Peak")
            axes[3].set_xlabel("Model | Config")
        else:
            axes[2].set_xlabel("Model | Config")
    else:
        axes[1].set_xlabel("Model | Config")

    plt.xticks(rotation=35, ha="right")
    plt.tight_layout()
    fig.savefig(output_path)
    plt.close(fig)

    return Path(output_path)


def main():
    config = interactive_config()
    if not config:
        return

    results_df = run_bench(config)
    if results_df.empty:
        print("⚠️ 沒有任何測試結果。")
        return

    ok_count = int((results_df["Status"] == "ok").sum())
    warning_count = int((results_df["Status"] == "warning").sum())
    error_count = int((results_df["Status"] == "error").sum())

    print("\n" + "═" * 62)
    console_df = build_summary_dataframe(results_df)[
        [
            "Run",
            "Status",
            "Output Category",
            "Finish Reason",
            "Model",
            "TPS (chunk/s)",
            "TTFT (s)",
            "First Event (s)",
            "Chunks (content/total)",
            "Config",
        ]
    ]
    print(console_df.to_markdown(index=False))
    print(f"\n📌 結果統計: ok={ok_count}, warning={warning_count}, error={error_count}")

    report_stem = f"bench_{config['backend']}_{time.strftime('%Y%m%d_%H%M%S')}"
    report_path = save_markdown_report(results_df, config, report_stem)
    chart_path = plot_results(results_df, f"{report_stem}.png")
    export_best_config(results_df, config)

    print(f"\n✅ 報告: {report_path}")
    if chart_path:
        print(f"📈 圖表: {chart_path}")
    else:
        print("ℹ️ 沒有可供繪圖的正常文字輸出結果。")


def interactive_config():
    print("\n" + "=" * 62)
    print("LLM Benchmark V3")
    print("=" * 62)

    backend = questionary.select(
        "請選擇測試後端:",
        choices=[
            Choice("Ollama", value="ollama"),
            Choice("llama.cpp（llama-server）", value="llama.cpp"),
        ],
    ).ask()
    if not backend:
        return None

    capability = questionary.select(
        "請選擇要測試的能力:",
        choices=[
            Choice(
                f"{info['label']} | {info['description']}",
                value=capability_key,
            )
            for capability_key, info in CAPABILITY_OPTIONS.items()
        ],
    ).ask()
    if not capability:
        return None

    url, models = select_models_and_url(backend)
    if not models:
        print("⚠️ 沒有可用模型，已取消。")
        return None

    available_groups = {
        group_name: [key for key in param_keys if backend in PARAM_INFO[key]["backends"]]
        for group_name, param_keys in PARAM_GROUPS.items()
    }
    available_groups = {name: keys for name, keys in available_groups.items() if keys}

    selected_groups = questionary.checkbox(
        "選擇想測試的參數類別（可不選，代表只比模型預設值）:",
        choices=[
            Choice(f"{group_name} ({len(param_keys)} 個)", value=group_name)
            for group_name, param_keys in available_groups.items()
        ],
    ).ask() or []

    final_params = {}
    for group_name in selected_groups:
        param_keys = available_groups[group_name]
        selected_params = questionary.checkbox(
            f"選擇 {group_name} 內要測試的參數:",
            choices=[
                Choice(
                    title=(
                        f"{PARAM_INFO[key]['label']} | 範圍: {PARAM_INFO[key]['range']} | "
                        f"{PARAM_INFO[key]['desc']}"
                    ),
                    value=key,
                )
                for key in param_keys
            ],
        ).ask() or []

        for key in selected_params:
            values = ask_param_values(key)
            if values is None:
                return None
            final_params[key] = values

    prompt = questionary.text(
        "測試 Prompt:",
        default=CAPABILITY_OPTIONS[capability]["default_prompt"],
    ).ask()
    if prompt is None:
        return None

    return {
        "backend": backend,
        "capability": capability,
        "url": url,
        "models": models,
        "params": final_params,
        "prompt": prompt,
    }


def run_bench(config):
    client = OpenAI(base_url=config["url"], api_key="sk-no-key-needed")
    capability = config.get("capability", "chat")
    param_keys = list(config["params"].keys())
    param_values = [config["params"][key] for key in param_keys]
    combos = [dict(zip(param_keys, combo)) for combo in product(*param_values)] if param_keys else [{}]

    results = []
    total_runs = len(config["models"]) * len(combos)
    vram_monitoring_enabled = query_nvidia_vram_snapshot() is not None
    config["vram_monitoring"] = "nvidia-smi" if vram_monitoring_enabled else "unavailable"

    capability_label = CAPABILITY_OPTIONS.get(capability, {}).get("label", capability)
    print(f"\n開始測試，共 {total_runs} 組，模式: {capability_label}")
    print("TPS 與 TTFT 適用於文字輸出；tools 模式主要看 First Event 與 tool_call 成功率。")
    if vram_monitoring_enabled:
        print("VRAM 監控: 已啟用 nvidia-smi")
    else:
        print("VRAM 監控: 未偵測到 nvidia-smi，相關欄位將顯示 N/A")

    run_index = 0
    for model in config["models"]:
        for param_set in combos:
            run_index += 1
            applied_params = build_backend_options(config["backend"], param_set)
            display_params = format_param_dict(param_set)
            print(f"[{run_index}/{total_runs}] {model} | {display_params}")

            request_kwargs = {}
            if applied_params:
                request_kwargs["extra_body"] = (
                    {"options": applied_params}
                    if config["backend"] == "ollama"
                    else applied_params
                )

            start_time = time.time()
            first_event_time = None
            first_content_time = None
            output_parts = []
            chunk_records = []
            vram_monitor = NvidiaVRAMMonitor() if vram_monitoring_enabled else None
            vram_metrics = empty_vram_metrics()
            if vram_monitor is not None and not vram_monitor.start():
                vram_monitor = None

            try:
                request_payload = build_chat_request_payload(config, model, request_kwargs)
                stream = client.chat.completions.create(**request_payload)

                for chunk in stream:
                    event_time = time.time()
                    if first_event_time is None:
                        first_event_time = event_time

                    chunk_info = inspect_stream_chunk(chunk)
                    chunk_records.append(chunk_info)

                    if chunk_info["content"] and first_content_time is None:
                        first_content_time = event_time
                    if chunk_info["content"]:
                        output_parts.append(chunk_info["content"])

                end_time = time.time()
                if vram_monitor is not None:
                    vram_metrics = vram_monitor.stop()
                output_text = "".join(output_parts)
                classification = classify_stream_result(
                    chunk_records=chunk_records,
                    start_time=start_time,
                    end_time=end_time,
                    first_event_time=first_event_time,
                    first_content_time=first_content_time,
                    error_message=None,
                )
                classification = adjust_classification_for_capability(classification, capability)
                results.append(
                    build_result_row(
                        run_id=run_index,
                        config=config,
                        model=model,
                        param_set=param_set,
                        applied_params=applied_params,
                        display_params=display_params,
                        classification=classification,
                        vram_metrics=vram_metrics,
                        output_text=output_text,
                        error_message=None,
                    )
                )
            except Exception as exc:
                end_time = time.time()
                if vram_monitor is not None:
                    vram_metrics = vram_monitor.stop()
                output_text = "".join(output_parts)
                classification = classify_stream_result(
                    chunk_records=chunk_records,
                    start_time=start_time,
                    end_time=end_time,
                    first_event_time=first_event_time,
                    first_content_time=first_content_time,
                    error_message=str(exc),
                )
                classification = adjust_classification_for_capability(classification, capability)
                print(f"錯誤: {exc}")
                results.append(
                    build_result_row(
                        run_id=run_index,
                        config=config,
                        model=model,
                        param_set=param_set,
                        applied_params=applied_params,
                        display_params=display_params,
                        classification=classification,
                        vram_metrics=vram_metrics,
                        output_text=output_text,
                        error_message=str(exc),
                    )
                )

    return pd.DataFrame(results)


def save_markdown_report(df, config, report_stem):
    report_path = Path(f"{report_stem}.html")
    summary_df = build_summary_dataframe(df)
    outcome_summary_df = build_outcome_summary_dataframe(df)
    wrapped_summary_df = wrap_markdown_table_headers(summary_df)
    wrapped_outcome_summary_df = wrap_markdown_table_headers(outcome_summary_df)
    capability = config.get("capability", "chat")
    tool_call_success_summary_df = (
        build_tool_call_success_summary_dataframe(df) if capability == "tools" else pd.DataFrame()
    )
    wrapped_tool_call_success_summary_df = wrap_markdown_table_headers(tool_call_success_summary_df)

    with report_path.open("w", encoding="utf-8") as file:
        file.write("# Benchmark Report\n\n")
        file.write(f"- Backend: {config['backend']}\n")
        file.write(f"- Capability: {capability}\n")
        file.write(f"- Base URL: {config['url']}\n")
        file.write(f"- Models: {', '.join(config['models'])}\n")
        file.write(f"- Prompt: {config['prompt']}\n")
        file.write("- Note: TPS is estimated from streaming content chunks for relative comparison.\n")
        if capability == "tools":
            file.write(
                "- Tool mode note: successful tool-calling runs may not emit text tokens, "
                "so `TPS` and `TTFT` can be `N/A`; focus on `Output Category=tool_call` "
                "and `First Event (s)`.\n"
            )
        file.write("\n## Environment Notes\n\n")
        file.write(f"- VRAM monitoring: {config.get('vram_monitoring', 'unavailable')}\n\n")
        file.write("## Metric Notes\n\n")
        file.write("- `TPS (chunk/s)`: Estimated throughput from text-bearing streaming chunks.\n")
        file.write("- `TTFT (s)`: Time to first text chunk.\n")
        file.write("- `First Event (s)`: Time to the first streamed event of any kind.\n")
        file.write("- `VRAM Peak (MiB)`: Highest observed total NVIDIA GPU memory usage during a run.\n")
        file.write(
            "- `Efficiency Score (TPS/GiB Peak)`: `TPS / (VRAM Peak in GiB)` when both values are available.\n\n"
        )
        file.write("## Output Diagnosis Notes\n\n")
        file.write("- `normal_content`: Received text output as expected for chat benchmarking.\n")
        file.write("- `tool_call`: Received `tool_calls` payload as expected for tool benchmarking.\n")
        file.write("- `text_reply_without_tool`: Returned text, but did not emit any tool call in tool mode.\n")
        file.write("- `empty_reply`: Stream completed without text output.\n")
        file.write("- `non_content_stream`: Stream only carried non-text payloads.\n")
        file.write("- `early_stop`: Stream ended before a complete reply or tool call was received.\n\n")
        file.write("## Summary\n\n")
        file.write(summary_df.to_markdown(index=False))
        file.write("\n\n## Outcome Summary\n\n")
        file.write(outcome_summary_df.to_markdown(index=False))
        file.write("\n\n## Generated Outputs\n")

        for _, row in df.iterrows():
            params_json = json.dumps(row["Params"], ensure_ascii=False)
            applied_params_json = json.dumps(row["Applied_Params"], ensure_ascii=False)
            file.write(f"\n### Run {row['Run_ID']}\n\n")
            file.write(f"- Status: {row['Status']}\n")
            file.write(f"- Capability: {row.get('Capability', capability)}\n")
            file.write(f"- Output Category: {row['Output_Category']}\n")
            file.write(f"- Diagnosis: {row['Diagnosis']}\n")
            file.write(f"- Finish Reason: {format_text_value(row['Finish_Reason'])}\n")
            file.write(f"- Backend: {row['Backend']}\n")
            file.write(f"- Model: {row['Model']}\n")
            file.write(f"- Params: `{params_json}`\n")
            file.write(f"- Applied Params: `{applied_params_json}`\n")
            file.write(f"- TPS: {format_numeric_value(row['TPS'], 2)} chunk/s\n")
            file.write(f"- TTFT: {format_numeric_value(row['TTFT'], 3)} s\n")
            file.write(f"- First Event: {format_numeric_value(row['First_Event_s'], 3)} s\n")
            file.write(f"- Stream Duration: {format_numeric_value(row['Stream_Duration_s'], 3)} s\n")
            file.write(f"- Total Chunks: {int(row['Total_Chunks'])}\n")
            file.write(f"- Content Chunks: {int(row['Content_Chunks'])}\n")
            file.write(f"- Non-Content Chunks: {int(row['Non_Content_Chunks'])}\n")
            file.write(f"- Non-Content Types: {row['Non_Content_Types']}\n")
            file.write(f"- VRAM Base: {format_mib_value(row['VRAM_Base_MiB'])}\n")
            file.write(f"- VRAM Peak: {format_mib_value(row['VRAM_Peak_MiB'])}\n")
            file.write(f"- VRAM Delta: {format_mib_value(row['VRAM_Delta_MiB'])}\n")
            file.write(f"- VRAM Detail: {row['VRAM_Detail']}\n")
            file.write(
                f"- Efficiency Score: "
                f"{format_numeric_value(row['Efficiency_Score'], 3)} TPS/GiB Peak\n"
            )
            file.write(f"- Output Chars: {row['Output_Chars']}\n")
            if row["Error"]:
                file.write(f"- Error: {row['Error']}\n")
            file.write("\n```text\n")
            file.write(normalize_output_text(row["Output_Text"]))
            file.write("\n```\n")

    return report_path


def select_best_result(eligible_df, capability):
    if capability == "tools":
        if eligible_df["First_Event_s"].notna().any():
            best_row = eligible_df.loc[eligible_df["First_Event_s"].idxmin()]
            return best_row, "first_event_s"
        best_row = eligible_df.loc[eligible_df["Stream_Duration_s"].idxmin()]
        return best_row, "stream_duration_s"

    if eligible_df["TPS"].notna().any():
        best_row = eligible_df.loc[eligible_df["TPS"].idxmax()]
        return best_row, "tps"

    best_row = eligible_df.loc[eligible_df["TTFT"].idxmin()]
    return best_row, "ttft"


def export_best_config(df, config):
    capability = config.get("capability", "chat")
    eligible_df = filter_eligible_results(df, capability=capability)
    if eligible_df.empty:
        print("⚠️ 沒有符合本次測試模式的成功結果，因此不輸出 best_config.json。")
        return None

    best_row, selection_metric = select_best_result(eligible_df, capability)
    payload = {
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "backend": best_row["Backend"],
        "capability": capability,
        "selection_metric": selection_metric,
        "model": best_row["Model"],
        "system_prompt_label": best_row.get("System_Prompt_Label", "N/A"),
        "system_prompt_text": best_row.get("System_Prompt_Text", ""),
        "params": best_row["Params"],
        "applied_params": best_row["Applied_Params"],
        "status": best_row["Status"],
        "output_category": best_row["Output_Category"],
        "finish_reason": best_row["Finish_Reason"],
        "diagnosis": best_row["Diagnosis"],
        "tps": best_row["TPS"],
        "ttft": best_row["TTFT"],
        "first_event_s": best_row["First_Event_s"],
        "stream_duration_s": best_row["Stream_Duration_s"],
        "vram_base_mib": best_row["VRAM_Base_MiB"],
        "vram_peak_mib": best_row["VRAM_Peak_MiB"],
        "vram_delta_mib": best_row["VRAM_Delta_MiB"],
        "vram_detail": best_row["VRAM_Detail"],
        "efficiency_score_tps_per_gib_peak": best_row["Efficiency_Score"],
        "prompt": config["prompt"],
    }

    with open("best_config.json", "w", encoding="utf-8") as file:
        json.dump(payload, file, ensure_ascii=False, indent=4)

    print("\n" + "=" * 18)
    print(f"最佳模型: {best_row['Model']}")
    if capability == "tools":
        print(f"最佳首事件延遲: {format_numeric_value(best_row['First_Event_s'], 3)} s")
        print(f"輸出類型: {best_row['Output_Category']}")
    else:
        print(f"最高 TPS: {format_numeric_value(best_row['TPS'], 2)} TPS")
    print(f"最佳設定: {best_row['Config_Str']}")
    print(f"VRAM Peak: {format_mib_value(best_row['VRAM_Peak_MiB'])}")
    if pd.notna(best_row["Efficiency_Score"]):
        print(f"效率分數: {format_numeric_value(best_row['Efficiency_Score'], 3)} TPS/GiB Peak")
    print("=" * 18)
    print("已保存 best_config.json")

    if best_row["Backend"] == "ollama":
        with open("Ollama_Modelfile_Suggest", "w", encoding="utf-8") as file:
            file.write(f"FROM {best_row['Model']}\n")
            for key, value in best_row["Applied_Params"].items():
                file.write(f"PARAMETER {key} {value}\n")
        print("已輸出 Ollama_Modelfile_Suggest")
    else:
        print("本次後端不是 Ollama，略過 Modelfile 建議。")

    return payload


def plot_results(df, output_path, capability="chat"):
    eligible_df = filter_eligible_results(df, capability=capability)
    if eligible_df.empty:
        return None

    eligible_df = eligible_df.copy()

    def build_label(row):
        label = f"{row['Model']} | {row['Config_Str']}"
        return label if len(label) <= 42 else label[:39] + "..."

    eligible_df["Plot_Label"] = eligible_df.apply(build_label, axis=1)

    if capability == "tools":
        min_first_event = eligible_df["First_Event_s"].min()
        min_stream_duration = eligible_df["Stream_Duration_s"].min()
        fig, axes = plt.subplots(2, 1, figsize=(13, 10), sharex=True)
        first_event_colors = [
            "lightgreen" if value == min_first_event else "steelblue"
            for value in eligible_df["First_Event_s"]
        ]
        duration_colors = [
            "lightgreen" if value == min_stream_duration else "slategray"
            for value in eligible_df["Stream_Duration_s"]
        ]

        axes[0].bar(eligible_df["Plot_Label"], eligible_df["First_Event_s"], color=first_event_colors)
        axes[0].axhline(y=min_first_event, color="green", linestyle="--", alpha=0.3)
        axes[0].set_title("Tool Call First Event Comparison")
        axes[0].set_ylabel("First Event (s)")

        axes[1].bar(
            eligible_df["Plot_Label"],
            eligible_df["Stream_Duration_s"],
            color=duration_colors,
        )
        axes[1].axhline(y=min_stream_duration, color="green", linestyle="--", alpha=0.3)
        axes[1].set_title("Tool Call Stream Duration Comparison")
        axes[1].set_ylabel("Stream Duration (s)")
        axes[1].set_xlabel("Model | Config")
    else:
        max_tps = eligible_df["TPS"].max()
        min_ttft = eligible_df["TTFT"].min()
        colors = ["gold" if tps == max_tps else "skyblue" for tps in eligible_df["TPS"]]
        ttft_colors = ["lightgreen" if ttft == min_ttft else "salmon" for ttft in eligible_df["TTFT"]]
        has_vram_data = eligible_df["VRAM_Peak_MiB"].notna().any()
        has_efficiency_data = eligible_df["Efficiency_Score"].notna().any()

        subplot_count = 4 if has_efficiency_data else 3 if has_vram_data else 2
        fig, axes = plt.subplots(subplot_count, 1, figsize=(13, 5 * subplot_count), sharex=True)
        if subplot_count == 1:
            axes = [axes]

        axes[0].bar(eligible_df["Plot_Label"], eligible_df["TPS"], color=colors)
        axes[0].axhline(y=max_tps, color="red", linestyle="--", alpha=0.3)
        axes[0].set_title("Throughput Comparison")
        axes[0].set_ylabel("TPS (chunk/s)")

        axes[1].bar(eligible_df["Plot_Label"], eligible_df["TTFT"], color=ttft_colors)
        axes[1].axhline(y=min_ttft, color="green", linestyle="--", alpha=0.3)
        axes[1].set_title("First Token Latency Comparison")
        axes[1].set_ylabel("TTFT (s)")

        if has_vram_data:
            min_vram_peak = eligible_df["VRAM_Peak_MiB"].min()
            vram_colors = [
                "lightgreen" if peak == min_vram_peak else "mediumpurple"
                for peak in eligible_df["VRAM_Peak_MiB"]
            ]
            axes[2].bar(eligible_df["Plot_Label"], eligible_df["VRAM_Peak_MiB"], color=vram_colors)
            axes[2].axhline(y=min_vram_peak, color="green", linestyle="--", alpha=0.3)
            axes[2].set_title("VRAM Peak Comparison")
            axes[2].set_ylabel("VRAM Peak (MiB)")
            if has_efficiency_data:
                max_efficiency = eligible_df["Efficiency_Score"].max()
                efficiency_colors = [
                    "gold" if score == max_efficiency else "steelblue"
                    for score in eligible_df["Efficiency_Score"]
                ]
                axes[3].bar(
                    eligible_df["Plot_Label"],
                    eligible_df["Efficiency_Score"],
                    color=efficiency_colors,
                )
                axes[3].axhline(y=max_efficiency, color="orange", linestyle="--", alpha=0.3)
                axes[3].set_title("Efficiency Score Comparison")
                axes[3].set_ylabel("TPS/GiB Peak")
                axes[3].set_xlabel("Model | Config")
            else:
                axes[2].set_xlabel("Model | Config")
        else:
            axes[1].set_xlabel("Model | Config")

    plt.xticks(rotation=35, ha="right")
    plt.tight_layout()
    fig.savefig(output_path)
    plt.close(fig)

    return Path(output_path)


def main():
    config = interactive_config()
    if not config:
        return

    results_df = run_bench(config)
    if results_df.empty:
        print("⚠️ 沒有產生任何 benchmark 結果。")
        return

    ok_count = int((results_df["Status"] == "ok").sum())
    warning_count = int((results_df["Status"] == "warning").sum())
    error_count = int((results_df["Status"] == "error").sum())

    print("\n" + "=" * 62)
    summary_df = build_summary_dataframe(results_df)
    console_columns = ["Run", "Status"]
    if "Capability" in summary_df.columns:
        console_columns.append("Capability")
    console_columns.extend(
        [
            "Output Category",
            "Finish Reason",
            "Model",
            "TPS (chunk/s)",
            "TTFT (s)",
            "First Event (s)",
            "Chunks (content/total)",
            "Config",
        ]
    )
    console_df = summary_df[console_columns]
    print(console_df.to_markdown(index=False))
    print(f"\n結果統計: ok={ok_count}, warning={warning_count}, error={error_count}")

    capability = config.get("capability", "chat")
    report_stem = f"bench_{config['backend']}_{capability}_{time.strftime('%Y%m%d_%H%M%S')}"
    report_path = save_markdown_report(results_df, config, report_stem)
    chart_path = plot_results(results_df, f"{report_stem}.png", capability=capability)
    export_best_config(results_df, config)

    print(f"\n報告已保存: {report_path}")
    if chart_path:
        print(f"圖表已保存: {chart_path}")
    else:
        print("沒有可繪圖的成功結果，略過圖表輸出。")


def interactive_config():
    capability_defaults = {
        "chat": "Explain the long-term creep risk of PETG in 3D printing and how to reduce it.",
        "tools": (
            "Check today's weather in Taipei. If you support tools or function calling, "
            "call the `lookup_weather` tool first instead of answering directly."
        ),
    }

    print("\n" + "=" * 62)
    print("LLM Benchmark V3")
    print("=" * 62)

    backend = questionary.select(
        "Select backend:",
        choices=[
            Choice("Ollama", value="ollama"),
            Choice("llama.cpp (llama-server)", value="llama.cpp"),
        ],
    ).ask()
    if not backend:
        return None

    capability = questionary.select(
        "Select benchmark mode:",
        choices=[
            Choice("Chat | Standard chat response benchmark", value="chat"),
            Choice("Tools | Check whether the model emits tool_calls", value="tools"),
        ],
    ).ask()
    if not capability:
        return None

    url, models = select_models_and_url(backend)
    if not models:
        print("No models available. Cancelled.")
        return None

    available_groups = {
        group_name: [key for key in param_keys if backend in PARAM_INFO[key]["backends"]]
        for group_name, param_keys in PARAM_GROUPS.items()
    }
    available_groups = {name: keys for name, keys in available_groups.items() if keys}

    selected_groups = questionary.checkbox(
        "Select parameter groups to benchmark (optional):",
        choices=[
            Choice(f"{group_name} ({len(param_keys)} params)", value=group_name)
            for group_name, param_keys in available_groups.items()
        ],
    ).ask() or []

    final_params = {}
    for group_name in selected_groups:
        param_keys = available_groups[group_name]
        selected_params = questionary.checkbox(
            f"Select params from {group_name}:",
            choices=[
                Choice(
                    title=(
                        f"{PARAM_INFO[key]['label']} | Range: {PARAM_INFO[key]['range']} | "
                        f"{PARAM_INFO[key]['desc']}"
                    ),
                    value=key,
                )
                for key in param_keys
            ],
        ).ask() or []

        for key in selected_params:
            values = ask_param_values(key)
            if values is None:
                return None
            final_params[key] = values

    prompt = questionary.text(
        "Benchmark prompt:",
        default=capability_defaults[capability],
    ).ask()
    if prompt is None:
        return None

    return {
        "backend": backend,
        "capability": capability,
        "url": url,
        "models": models,
        "params": final_params,
        "prompt": prompt,
    }


def run_bench(config):
    client = OpenAI(base_url=config["url"], api_key="sk-no-key-needed")
    capability = config.get("capability", "chat")
    capability_label = {"chat": "chat", "tools": "tools"}.get(capability, capability)
    param_keys = list(config["params"].keys())
    param_values = [config["params"][key] for key in param_keys]
    combos = [dict(zip(param_keys, combo)) for combo in product(*param_values)] if param_keys else [{}]

    results = []
    total_runs = len(config["models"]) * len(combos)
    vram_monitoring_enabled = query_nvidia_vram_snapshot() is not None
    config["vram_monitoring"] = "nvidia-smi" if vram_monitoring_enabled else "unavailable"

    print(f"\nStarting benchmark with {total_runs} runs. Mode: {capability_label}")
    print("Chat mode focuses on TPS/TTFT. Tools mode focuses on tool_call success and first event latency.")
    if vram_monitoring_enabled:
        print("VRAM monitoring: enabled via nvidia-smi")
    else:
        print("VRAM monitoring: nvidia-smi not available, VRAM fields will be N/A")

    run_index = 0
    for model in config["models"]:
        for param_set in combos:
            run_index += 1
            applied_params = build_backend_options(config["backend"], param_set)
            display_params = format_param_dict(param_set)
            print(f"[{run_index}/{total_runs}] {model} | {display_params}")

            request_kwargs = {}
            if applied_params:
                request_kwargs["extra_body"] = (
                    {"options": applied_params}
                    if config["backend"] == "ollama"
                    else applied_params
                )

            start_time = time.time()
            first_event_time = None
            first_content_time = None
            output_parts = []
            chunk_records = []
            vram_monitor = NvidiaVRAMMonitor() if vram_monitoring_enabled else None
            vram_metrics = empty_vram_metrics()
            if vram_monitor is not None and not vram_monitor.start():
                vram_monitor = None

            try:
                request_payload = build_chat_request_payload(config, model, request_kwargs)
                stream = client.chat.completions.create(**request_payload)

                for chunk in stream:
                    event_time = time.time()
                    if first_event_time is None:
                        first_event_time = event_time

                    chunk_info = inspect_stream_chunk(chunk)
                    chunk_records.append(chunk_info)

                    if chunk_info["content"] and first_content_time is None:
                        first_content_time = event_time
                    if chunk_info["content"]:
                        output_parts.append(chunk_info["content"])

                end_time = time.time()
                if vram_monitor is not None:
                    vram_metrics = vram_monitor.stop()
                output_text = "".join(output_parts)
                classification = classify_stream_result(
                    chunk_records=chunk_records,
                    start_time=start_time,
                    end_time=end_time,
                    first_event_time=first_event_time,
                    first_content_time=first_content_time,
                    error_message=None,
                )
                classification = adjust_classification_for_capability(classification, capability)
                results.append(
                    build_result_row(
                        run_id=run_index,
                        config=config,
                        model=model,
                        param_set=param_set,
                        applied_params=applied_params,
                        display_params=display_params,
                        classification=classification,
                        vram_metrics=vram_metrics,
                        output_text=output_text,
                        error_message=None,
                    )
                )
            except Exception as exc:
                end_time = time.time()
                if vram_monitor is not None:
                    vram_metrics = vram_monitor.stop()
                output_text = "".join(output_parts)
                classification = classify_stream_result(
                    chunk_records=chunk_records,
                    start_time=start_time,
                    end_time=end_time,
                    first_event_time=first_event_time,
                    first_content_time=first_content_time,
                    error_message=str(exc),
                )
                classification = adjust_classification_for_capability(classification, capability)
                print(f"Error: {exc}")
                results.append(
                    build_result_row(
                        run_id=run_index,
                        config=config,
                        model=model,
                        param_set=param_set,
                        applied_params=applied_params,
                        display_params=display_params,
                        classification=classification,
                        vram_metrics=vram_metrics,
                        output_text=output_text,
                        error_message=str(exc),
                    )
                )

    return pd.DataFrame(results)


def save_markdown_report(df, config, report_stem):
    report_path = Path(f"{report_stem}.html")
    summary_df = build_summary_dataframe(df)
    outcome_summary_df = build_outcome_summary_dataframe(df)
    capability = config.get("capability", "chat")
    wrapped_summary_df = wrap_markdown_table_headers(summary_df)
    wrapped_outcome_summary_df = wrap_markdown_table_headers(outcome_summary_df)
    tool_call_success_summary_df = (
        build_tool_call_success_summary_dataframe(df) if capability == "tools" else pd.DataFrame()
    )
    wrapped_tool_call_success_summary_df = wrap_markdown_table_headers(
        tool_call_success_summary_df
    )

    with report_path.open("w", encoding="utf-8") as file:
        file.write("# Benchmark Report\n\n")
        file.write(f"- Backend: {config['backend']}\n")
        file.write(f"- Capability: {capability}\n")
        file.write(f"- Base URL: {config['url']}\n")
        file.write(f"- Models: {', '.join(config['models'])}\n")
        file.write(f"- Prompt: {config['prompt']}\n")
        file.write("- Note: TPS is estimated from streaming content chunks for relative comparison.\n")
        if capability == "tools":
            file.write(
                "- Tool mode note: successful tool-calling runs may not emit text tokens, "
                "so `TPS` and `TTFT` can be `N/A`; focus on `Output Category=tool_call` "
                "and `First Event (s)`.\n"
            )
        file.write("\n## Environment Notes\n\n")
        file.write(f"- VRAM monitoring: {config.get('vram_monitoring', 'unavailable')}\n\n")
        file.write("## Metric Notes\n\n")
        file.write("- `TPS (chunk/s)`: Estimated throughput from text-bearing streaming chunks.\n")
        file.write("- `TTFT (s)`: Time to first text chunk.\n")
        file.write("- `First Event (s)`: Time to the first streamed event of any kind.\n")
        file.write("- `VRAM Peak (MiB)`: Highest observed total NVIDIA GPU memory usage during a run.\n")
        file.write(
            "- `Efficiency Score (TPS/GiB Peak)`: `TPS / (VRAM Peak in GiB)` when both values are available.\n\n"
        )
        file.write("## Output Diagnosis Notes\n\n")
        file.write("- `normal_content`: Received text output as expected for chat benchmarking.\n")
        file.write("- `tool_call`: Received `tool_calls` payload as expected for tool benchmarking.\n")
        file.write("- `text_reply_without_tool`: Returned text, but did not emit any tool call in tool mode.\n")
        file.write("- `empty_reply`: Stream completed without text output.\n")
        file.write("- `non_content_stream`: Stream only carried non-text payloads.\n")
        file.write("- `early_stop`: Stream ended before a complete reply or tool call was received.\n\n")
        file.write("## Summary\n\n")
        file.write(summary_df.to_markdown(index=False))
        file.write("\n\n## Outcome Summary\n\n")
        file.write(outcome_summary_df.to_markdown(index=False))
        file.write("\n\n## Generated Outputs\n")

        for _, row in df.iterrows():
            params_json = json.dumps(row["Params"], ensure_ascii=False)
            applied_params_json = json.dumps(row["Applied_Params"], ensure_ascii=False)
            file.write(f"\n### Run {row['Run_ID']}\n\n")
            file.write(f"- Status: {row['Status']}\n")
            file.write(f"- Capability: {row.get('Capability', capability)}\n")
            file.write(f"- Output Category: {row['Output_Category']}\n")
            file.write(f"- Diagnosis: {row['Diagnosis']}\n")
            file.write(f"- Finish Reason: {format_text_value(row['Finish_Reason'])}\n")
            file.write(f"- Backend: {row['Backend']}\n")
            file.write(f"- Model: {row['Model']}\n")
            file.write(f"- Params: `{params_json}`\n")
            file.write(f"- Applied Params: `{applied_params_json}`\n")
            file.write(f"- TPS: {format_numeric_value(row['TPS'], 2)} chunk/s\n")
            file.write(f"- TTFT: {format_numeric_value(row['TTFT'], 3)} s\n")
            file.write(f"- First Event: {format_numeric_value(row['First_Event_s'], 3)} s\n")
            file.write(f"- Stream Duration: {format_numeric_value(row['Stream_Duration_s'], 3)} s\n")
            file.write(f"- Total Chunks: {int(row['Total_Chunks'])}\n")
            file.write(f"- Content Chunks: {int(row['Content_Chunks'])}\n")
            file.write(f"- Non-Content Chunks: {int(row['Non_Content_Chunks'])}\n")
            file.write(f"- Non-Content Types: {row['Non_Content_Types']}\n")
            file.write(f"- VRAM Base: {format_mib_value(row['VRAM_Base_MiB'])}\n")
            file.write(f"- VRAM Peak: {format_mib_value(row['VRAM_Peak_MiB'])}\n")
            file.write(f"- VRAM Delta: {format_mib_value(row['VRAM_Delta_MiB'])}\n")
            file.write(f"- VRAM Detail: {row['VRAM_Detail']}\n")
            file.write(
                f"- Efficiency Score: "
                f"{format_numeric_value(row['Efficiency_Score'], 3)} TPS/GiB Peak\n"
            )
            file.write(f"- Output Chars: {row['Output_Chars']}\n")
            if row["Error"]:
                file.write(f"- Error: {row['Error']}\n")
            file.write("\n```text\n")
            file.write(normalize_output_text(row["Output_Text"]))
            file.write("\n```\n")

    return report_path


def select_best_result(eligible_df, capability):
    if capability == "tools":
        if eligible_df["First_Event_s"].notna().any():
            best_row = eligible_df.loc[eligible_df["First_Event_s"].idxmin()]
            return best_row, "first_event_s"
        best_row = eligible_df.loc[eligible_df["Stream_Duration_s"].idxmin()]
        return best_row, "stream_duration_s"

    if eligible_df["TPS"].notna().any():
        best_row = eligible_df.loc[eligible_df["TPS"].idxmax()]
        return best_row, "tps"

    best_row = eligible_df.loc[eligible_df["TTFT"].idxmin()]
    return best_row, "ttft"


def export_best_config(df, config, output_dir="."):
    capability = config.get("capability", "chat")
    eligible_df = filter_eligible_results(df, capability=capability)
    if eligible_df.empty:
        print("No successful result for this benchmark mode, so best_config.json was not written.")
        return None

    best_row, selection_metric = select_best_result(eligible_df, capability)
    payload = {
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "backend": best_row["Backend"],
        "capability": capability,
        "selection_metric": selection_metric,
        "model": best_row["Model"],
        "params": best_row["Params"],
        "applied_params": best_row["Applied_Params"],
        "status": best_row["Status"],
        "output_category": best_row["Output_Category"],
        "finish_reason": best_row["Finish_Reason"],
        "diagnosis": best_row["Diagnosis"],
        "tps": best_row["TPS"],
        "ttft": best_row["TTFT"],
        "first_event_s": best_row["First_Event_s"],
        "stream_duration_s": best_row["Stream_Duration_s"],
        "vram_base_mib": best_row["VRAM_Base_MiB"],
        "vram_peak_mib": best_row["VRAM_Peak_MiB"],
        "vram_delta_mib": best_row["VRAM_Delta_MiB"],
        "vram_detail": best_row["VRAM_Detail"],
        "efficiency_score_tps_per_gib_peak": best_row["Efficiency_Score"],
        "prompt": config["prompt"],
        "system_prompts": config.get("system_prompts", []),
    }
    serializable_payload = {
        key: serialize_result_value(value) for key, value in payload.items()
    }

    output_dir_path = Path(output_dir)
    output_dir_path.mkdir(parents=True, exist_ok=True)
    best_config_path = output_dir_path / "best_config.json"
    with best_config_path.open("w", encoding="utf-8") as file:
        json.dump(serializable_payload, file, ensure_ascii=False, indent=4)

    print("\n" + "=" * 18)
    print(f"Best model: {best_row['Model']}")
    if capability == "tools":
        print(f"Best first-event latency: {format_numeric_value(best_row['First_Event_s'], 3)} s")
        print(f"Output category: {best_row['Output_Category']}")
    else:
        print(f"Highest TPS: {format_numeric_value(best_row['TPS'], 2)} TPS")
    print(f"Best config: {best_row['Config_Str']}")
    print(f"VRAM Peak: {format_mib_value(best_row['VRAM_Peak_MiB'])}")
    if pd.notna(best_row["Efficiency_Score"]):
        print(f"Efficiency score: {format_numeric_value(best_row['Efficiency_Score'], 3)} TPS/GiB Peak")
    print("=" * 18)
    print(f"Saved best_config.json: {best_config_path}")

    modelfile_path = None
    if best_row["Backend"] == "ollama":
        modelfile_path = output_dir_path / "Ollama_Modelfile_Suggest"
        modelfile_params = build_ollama_modelfile_params(best_row["Params"])
        with modelfile_path.open("w", encoding="utf-8") as file:
            file.write(f"FROM {best_row['Model']}\n")
            for key, value in modelfile_params.items():
                file.write(f"PARAMETER {key} {value}\n")
        print(f"Saved Ollama_Modelfile_Suggest: {modelfile_path}")
    else:
        print("Backend is not Ollama, so Modelfile output was skipped.")

    return {
        "payload": serializable_payload,
        "best_config_path": best_config_path,
        "modelfile_path": modelfile_path,
    }


def plot_results(df, output_path, capability="chat"):
    eligible_df = filter_eligible_results(df, capability=capability)
    if eligible_df.empty:
        return None

    eligible_df = eligible_df.copy()

    def build_label(row):
        label = f"{row['Model']} | {row['Config_Str']}"
        return label if len(label) <= 42 else label[:39] + "..."

    eligible_df["Plot_Label"] = eligible_df.apply(build_label, axis=1)

    if capability == "tools":
        min_first_event = eligible_df["First_Event_s"].min()
        min_stream_duration = eligible_df["Stream_Duration_s"].min()
        fig, axes = plt.subplots(2, 1, figsize=(13, 10), sharex=True)
        first_event_colors = [
            "lightgreen" if value == min_first_event else "steelblue"
            for value in eligible_df["First_Event_s"]
        ]
        duration_colors = [
            "lightgreen" if value == min_stream_duration else "slategray"
            for value in eligible_df["Stream_Duration_s"]
        ]

        axes[0].bar(eligible_df["Plot_Label"], eligible_df["First_Event_s"], color=first_event_colors)
        axes[0].axhline(y=min_first_event, color="green", linestyle="--", alpha=0.3)
        axes[0].set_title("Tool Call First Event Comparison")
        axes[0].set_ylabel("First Event (s)")

        axes[1].bar(
            eligible_df["Plot_Label"],
            eligible_df["Stream_Duration_s"],
            color=duration_colors,
        )
        axes[1].axhline(y=min_stream_duration, color="green", linestyle="--", alpha=0.3)
        axes[1].set_title("Tool Call Stream Duration Comparison")
        axes[1].set_ylabel("Stream Duration (s)")
        axes[1].set_xlabel("Model | Config")
    else:
        max_tps = eligible_df["TPS"].max()
        min_ttft = eligible_df["TTFT"].min()
        colors = ["gold" if tps == max_tps else "skyblue" for tps in eligible_df["TPS"]]
        ttft_colors = ["lightgreen" if ttft == min_ttft else "salmon" for ttft in eligible_df["TTFT"]]
        has_vram_data = eligible_df["VRAM_Peak_MiB"].notna().any()
        has_efficiency_data = eligible_df["Efficiency_Score"].notna().any()

        subplot_count = 4 if has_efficiency_data else 3 if has_vram_data else 2
        fig, axes = plt.subplots(subplot_count, 1, figsize=(13, 5 * subplot_count), sharex=True)
        if subplot_count == 1:
            axes = [axes]

        axes[0].bar(eligible_df["Plot_Label"], eligible_df["TPS"], color=colors)
        axes[0].axhline(y=max_tps, color="red", linestyle="--", alpha=0.3)
        axes[0].set_title("Throughput Comparison")
        axes[0].set_ylabel("TPS (chunk/s)")

        axes[1].bar(eligible_df["Plot_Label"], eligible_df["TTFT"], color=ttft_colors)
        axes[1].axhline(y=min_ttft, color="green", linestyle="--", alpha=0.3)
        axes[1].set_title("First Token Latency Comparison")
        axes[1].set_ylabel("TTFT (s)")

        if has_vram_data:
            min_vram_peak = eligible_df["VRAM_Peak_MiB"].min()
            vram_colors = [
                "lightgreen" if peak == min_vram_peak else "mediumpurple"
                for peak in eligible_df["VRAM_Peak_MiB"]
            ]
            axes[2].bar(eligible_df["Plot_Label"], eligible_df["VRAM_Peak_MiB"], color=vram_colors)
            axes[2].axhline(y=min_vram_peak, color="green", linestyle="--", alpha=0.3)
            axes[2].set_title("VRAM Peak Comparison")
            axes[2].set_ylabel("VRAM Peak (MiB)")
            if has_efficiency_data:
                max_efficiency = eligible_df["Efficiency_Score"].max()
                efficiency_colors = [
                    "gold" if score == max_efficiency else "steelblue"
                    for score in eligible_df["Efficiency_Score"]
                ]
                axes[3].bar(
                    eligible_df["Plot_Label"],
                    eligible_df["Efficiency_Score"],
                    color=efficiency_colors,
                )
                axes[3].axhline(y=max_efficiency, color="orange", linestyle="--", alpha=0.3)
                axes[3].set_title("Efficiency Score Comparison")
                axes[3].set_ylabel("TPS/GiB Peak")
                axes[3].set_xlabel("Model | Config")
            else:
                axes[2].set_xlabel("Model | Config")
        else:
            axes[1].set_xlabel("Model | Config")

    plt.xticks(rotation=35, ha="right")
    plt.tight_layout()
    fig.savefig(output_path)
    plt.close(fig)

    return Path(output_path)


def main():
    config = interactive_config()
    if not config:
        return

    results_df = run_bench(config)
    if results_df.empty:
        print("No benchmark rows were produced.")
        return

    ok_count = int((results_df["Status"] == "ok").sum())
    warning_count = int((results_df["Status"] == "warning").sum())
    error_count = int((results_df["Status"] == "error").sum())

    print("\n" + "=" * 62)
    summary_df = build_summary_dataframe(results_df)
    console_columns = ["Run", "Status"]
    if "Capability" in summary_df.columns:
        console_columns.append("Capability")
    console_columns.extend(
        [
            "Output Category",
            "Finish Reason",
            "Model",
            "TPS (chunk/s)",
            "TTFT (s)",
            "First Event (s)",
            "Chunks (content/total)",
            "Config",
        ]
    )
    console_df = summary_df[console_columns]
    print(console_df.to_markdown(index=False))
    print(f"\nResult counts: ok={ok_count}, warning={warning_count}, error={error_count}")

    capability = config.get("capability", "chat")
    report_stem = f"bench_{config['backend']}_{capability}_{time.strftime('%Y%m%d_%H%M%S')}"
    report_path = save_markdown_report(results_df, config, report_stem)
    chart_path = plot_results(results_df, f"{report_stem}.png", capability=capability)
    export_best_config(results_df, config)

    print(f"\nSaved report: {report_path}")
    if chart_path:
        print(f"Saved chart: {chart_path}")
    else:
        print("Skipped chart output because there were no eligible successful results.")


def save_markdown_report(df, config, report_stem):
    report_path = Path(f"{report_stem}.html")
    summary_df = build_summary_dataframe(df)
    outcome_summary_df = build_outcome_summary_dataframe(df)
    capability = config.get("capability", "chat")
    wrapped_summary_df = wrap_markdown_table_headers(summary_df)
    wrapped_outcome_summary_df = wrap_markdown_table_headers(outcome_summary_df)
    tool_call_success_summary_df = (
        build_tool_call_success_summary_dataframe(df) if capability == "tools" else pd.DataFrame()
    )
    wrapped_tool_call_success_summary_df = wrap_markdown_table_headers(
        tool_call_success_summary_df
    )

    with report_path.open("w", encoding="utf-8") as file:
        file.write("# Benchmark Report\n\n")
        file.write(f"- Backend: {config['backend']}\n")
        file.write(f"- Capability: {capability}\n")
        file.write(f"- Base URL: {config['url']}\n")
        file.write(f"- Models: {', '.join(config['models'])}\n")
        file.write(f"- Prompt: {config['prompt']}\n")
        file.write("- Note: TPS is estimated from streaming content chunks for relative comparison.\n")
        if capability == "tools":
            file.write(
                "- Tool mode note: successful tool-calling runs may not emit text tokens, "
                "so `TPS` and `TTFT` can be `N/A`; focus on `Output Category=tool_call` "
                "and `First Event (s)`.\n"
            )
        file.write("\n## Environment Notes\n\n")
        file.write(f"- VRAM monitoring: {config.get('vram_monitoring', 'unavailable')}\n\n")
        file.write("## Metric Notes\n\n")
        file.write("- `TPS (chunk/s)`: Estimated throughput from text-bearing streaming chunks.\n")
        file.write("- `TTFT (s)`: Time to first text chunk.\n")
        file.write("- `First Event (s)`: Time to the first streamed event of any kind.\n")
        file.write("- `VRAM Peak (MiB)`: Highest observed total NVIDIA GPU memory usage during a run.\n")
        file.write(
            "- `Efficiency Score (TPS/GiB Peak)`: `TPS / (VRAM Peak in GiB)` when both values are available.\n\n"
        )
        file.write("## Output Diagnosis Notes\n\n")
        file.write("- `normal_content`: Received text output as expected for chat benchmarking.\n")
        file.write("- `tool_call`: Received `tool_calls` payload as expected for tool benchmarking.\n")
        file.write("- `text_reply_without_tool`: Returned text, but did not emit any tool call in tool mode.\n")
        file.write("- `empty_reply`: Stream completed without text output.\n")
        file.write("- `non_content_stream`: Stream only carried non-text payloads.\n")
        file.write("- `early_stop`: Stream ended before a complete reply or tool call was received.\n\n")
        file.write("## Summary\n\n")
        file.write(dataframe_to_report_table(wrapped_summary_df))
        file.write("\n\n## Outcome Summary\n\n")
        file.write(dataframe_to_report_table(wrapped_outcome_summary_df))
        if capability == "tools" and not tool_call_success_summary_df.empty:
            file.write("\n\n## Tool Call Success by Model\n\n")
            file.write(dataframe_to_report_table(wrapped_tool_call_success_summary_df))
        file.write("\n\n## Generated Outputs\n")

        for _, row in df.iterrows():
            params_json = json.dumps(row["Params"], ensure_ascii=False)
            applied_params_json = json.dumps(row["Applied_Params"], ensure_ascii=False)
            file.write(f"\n### Run {row['Run_ID']}\n\n")
            file.write(f"- Status: {row['Status']}\n")
            file.write(f"- Capability: {row.get('Capability', capability)}\n")
            file.write(f"- Output Category: {row['Output_Category']}\n")
            file.write(f"- Diagnosis: {row['Diagnosis']}\n")
            file.write(f"- Finish Reason: {format_text_value(row['Finish_Reason'])}\n")
            file.write(f"- Backend: {row['Backend']}\n")
            file.write(f"- Model: {row['Model']}\n")
            file.write(f"- Params: `{params_json}`\n")
            file.write(f"- Applied Params: `{applied_params_json}`\n")
            file.write(f"- TPS: {format_numeric_value(row['TPS'], 2)} chunk/s\n")
            file.write(f"- TTFT: {format_numeric_value(row['TTFT'], 3)} s\n")
            file.write(f"- First Event: {format_numeric_value(row['First_Event_s'], 3)} s\n")
            file.write(f"- Stream Duration: {format_numeric_value(row['Stream_Duration_s'], 3)} s\n")
            file.write(f"- Total Chunks: {int(row['Total_Chunks'])}\n")
            file.write(f"- Content Chunks: {int(row['Content_Chunks'])}\n")
            file.write(f"- Non-Content Chunks: {int(row['Non_Content_Chunks'])}\n")
            file.write(f"- Non-Content Types: {row['Non_Content_Types']}\n")
            file.write(f"- VRAM Base: {format_mib_value(row['VRAM_Base_MiB'])}\n")
            file.write(f"- VRAM Peak: {format_mib_value(row['VRAM_Peak_MiB'])}\n")
            file.write(f"- VRAM Delta: {format_mib_value(row['VRAM_Delta_MiB'])}\n")
            file.write(f"- VRAM Detail: {row['VRAM_Detail']}\n")
            file.write(
                f"- Efficiency Score: "
                f"{format_numeric_value(row['Efficiency_Score'], 3)} TPS/GiB Peak\n"
            )
            file.write(f"- Output Chars: {row['Output_Chars']}\n")
            if row["Error"]:
                file.write(f"- Error: {row['Error']}\n")
            file.write("\n```text\n")
            file.write(normalize_output_text(row["Output_Text"]))
            file.write("\n```\n")

    return report_path


def main():
    config = interactive_config()
    if not config:
        return

    results_df = run_bench(config)
    if results_df.empty:
        print("No benchmark rows were produced.")
        return

    ok_count = int((results_df["Status"] == "ok").sum())
    warning_count = int((results_df["Status"] == "warning").sum())
    error_count = int((results_df["Status"] == "error").sum())
    capability = config.get("capability", "chat")
    report_dir = ensure_report_output_dir()
    report_stem = report_dir / f"bench_{config['backend']}_{capability}_{time.strftime('%Y%m%d_%H%M%S')}"

    raw_outputs_path = save_raw_outputs(results_df, report_stem)

    print("\n" + "=" * 62)
    summary_df = build_summary_dataframe(results_df)
    console_columns = ["Run", "Status"]
    if "Capability" in summary_df.columns:
        console_columns.append("Capability")
    console_columns.extend(
        [
            "Output Category",
            "Finish Reason",
            "Model",
            "TPS (chunk/s)",
            "TTFT (s)",
            "First Event (s)",
            "Chunks (content/total)",
            "Config",
        ]
    )
    console_df = summary_df[console_columns]
    print(dataframe_to_text_table(console_df))
    print(f"\nResult counts: ok={ok_count}, warning={warning_count}, error={error_count}")

    report_path = None
    chart_path = None
    summary_excel_path = None

    try:
        summary_excel_path = save_summary_excel_workbook(results_df, config, report_stem)
    except Exception as exc:
        print(f"Summary Excel export failed: {exc}")

    try:
        report_path = save_markdown_report(
            results_df,
            config,
            report_stem,
            summary_excel_path=summary_excel_path,
        )
    except Exception as exc:
        print(f"Report generation failed: {exc}")

    try:
        chart_path = plot_results(results_df, f"{report_stem}.png", capability=capability)
    except Exception as exc:
        print(f"Chart generation failed: {exc}")

    try:
        best_config_artifacts = export_best_config(results_df, config, output_dir=report_dir)
    except Exception as exc:
        best_config_artifacts = None
        print(f"best_config export failed: {exc}")

    print(f"\nSaved artifacts directory: {report_dir}")
    if report_path:
        print(f"\nSaved report: {report_path}")
    else:
        print("\nMarkdown report was not saved.")
    if summary_excel_path:
        print(f"Saved summary Excel: {summary_excel_path}")

    print(f"Saved raw outputs: {raw_outputs_path}")
    if chart_path:
        print(f"Saved chart: {chart_path}")
    else:
        print("Skipped chart output because there were no eligible successful results.")
    if best_config_artifacts:
        print(f"Saved best config: {best_config_artifacts['best_config_path']}")
        if best_config_artifacts["modelfile_path"]:
            print(f"Saved Modelfile suggestion: {best_config_artifacts['modelfile_path']}")


def plot_results(df, output_path, capability="chat"):
    if df.empty:
        return None

    plot_df, used_fallback = select_plot_dataframe(df, capability=capability)
    plot_df = plot_df.copy()

    def build_label(row):
        label = f"{row['Model']} | {row['Config_Str']}"
        return label if len(label) <= 42 else label[:39] + "..."

    def build_run_color(row):
        if row["Output_Category"] == "tool_call":
            return "seagreen"
        if row["Status"] == "ok":
            return "skyblue"
        if row["Status"] == "warning":
            return "darkorange"
        return "indianred"

    plot_df["Plot_Label"] = plot_df.apply(build_label, axis=1)
    plot_df["Plot_Color"] = plot_df.apply(build_run_color, axis=1)

    if capability == "tools":
        first_event_values = plot_df["First_Event_s"].fillna(0)
        duration_values = plot_df["Stream_Duration_s"].fillna(0)
        outcome_df = build_outcome_summary_dataframe(plot_df)

        fig, axes = plt.subplots(3, 1, figsize=(14, 15))

        axes[0].bar(plot_df["Plot_Label"], first_event_values, color=plot_df["Plot_Color"])
        if plot_df["First_Event_s"].notna().any():
            axes[0].axhline(
                y=plot_df["First_Event_s"].min(),
                color="green",
                linestyle="--",
                alpha=0.3,
            )
        axes[0].set_title("Tools Benchmark First Event Comparison")
        axes[0].set_ylabel("First Event (s)")

        axes[1].bar(plot_df["Plot_Label"], duration_values, color=plot_df["Plot_Color"])
        if plot_df["Stream_Duration_s"].notna().any():
            axes[1].axhline(
                y=plot_df["Stream_Duration_s"].min(),
                color="green",
                linestyle="--",
                alpha=0.3,
            )
        axes[1].set_title("Tools Benchmark Stream Duration Comparison")
        axes[1].set_ylabel("Stream Duration (s)")

        axes[2].bar(outcome_df["Output Category"], outcome_df["Count"], color="steelblue")
        axes[2].set_title("Tools Outcome Category Counts")
        axes[2].set_ylabel("Runs")
        axes[2].set_xlabel("Output Category")

        axes[0].tick_params(axis="x", rotation=35)
        axes[1].tick_params(axis="x", rotation=35)
        axes[2].tick_params(axis="x", rotation=20)
    else:
        eligible_df = filter_eligible_results(plot_df, capability=capability)
        if not eligible_df.empty:
            max_tps = eligible_df["TPS"].max()
            min_ttft = eligible_df["TTFT"].min()
            colors = ["gold" if tps == max_tps else "skyblue" for tps in eligible_df["TPS"]]
            ttft_colors = [
                "lightgreen" if ttft == min_ttft else "salmon" for ttft in eligible_df["TTFT"]
            ]
            has_vram_data = eligible_df["VRAM_Peak_MiB"].notna().any()
            has_efficiency_data = eligible_df["Efficiency_Score"].notna().any()

            subplot_count = 4 if has_efficiency_data else 3 if has_vram_data else 2
            fig, axes = plt.subplots(subplot_count, 1, figsize=(13, 5 * subplot_count), sharex=True)
            if subplot_count == 1:
                axes = [axes]

            axes[0].bar(eligible_df["Plot_Label"], eligible_df["TPS"], color=colors)
            axes[0].axhline(y=max_tps, color="red", linestyle="--", alpha=0.3)
            axes[0].set_title("Throughput Comparison")
            axes[0].set_ylabel("TPS (chunk/s)")

            axes[1].bar(eligible_df["Plot_Label"], eligible_df["TTFT"], color=ttft_colors)
            axes[1].axhline(y=min_ttft, color="green", linestyle="--", alpha=0.3)
            axes[1].set_title("First Token Latency Comparison")
            axes[1].set_ylabel("TTFT (s)")

            if has_vram_data:
                min_vram_peak = eligible_df["VRAM_Peak_MiB"].min()
                vram_colors = [
                    "lightgreen" if peak == min_vram_peak else "mediumpurple"
                    for peak in eligible_df["VRAM_Peak_MiB"]
                ]
                axes[2].bar(eligible_df["Plot_Label"], eligible_df["VRAM_Peak_MiB"], color=vram_colors)
                axes[2].axhline(y=min_vram_peak, color="green", linestyle="--", alpha=0.3)
                axes[2].set_title("VRAM Peak Comparison")
                axes[2].set_ylabel("VRAM Peak (MiB)")
                if has_efficiency_data:
                    max_efficiency = eligible_df["Efficiency_Score"].max()
                    efficiency_colors = [
                        "gold" if score == max_efficiency else "steelblue"
                        for score in eligible_df["Efficiency_Score"]
                    ]
                    axes[3].bar(
                        eligible_df["Plot_Label"],
                        eligible_df["Efficiency_Score"],
                        color=efficiency_colors,
                    )
                    axes[3].axhline(y=max_efficiency, color="orange", linestyle="--", alpha=0.3)
                    axes[3].set_title("Efficiency Score Comparison")
                    axes[3].set_ylabel("TPS/GiB Peak")
                    axes[3].set_xlabel("Model | Config")
                else:
                    axes[2].set_xlabel("Model | Config")
            else:
                axes[1].set_xlabel("Model | Config")
        else:
            status_df = (
                plot_df["Status"]
                .fillna("unknown")
                .value_counts(dropna=False)
                .rename_axis("Status")
                .reset_index(name="Count")
            )
            outcome_df = build_outcome_summary_dataframe(plot_df)
            fig, axes = plt.subplots(2, 1, figsize=(13, 10))

            axes[0].bar(status_df["Status"], status_df["Count"], color="steelblue")
            axes[0].set_title("Run Status Counts")
            axes[0].set_ylabel("Runs")

            axes[1].bar(outcome_df["Output Category"], outcome_df["Count"], color="slategray")
            title = "Outcome Category Counts"
            if used_fallback:
                title += " (Fallback)"
            axes[1].set_title(title)
            axes[1].set_ylabel("Runs")
            axes[1].set_xlabel("Output Category")
            axes[0].tick_params(axis="x", rotation=20)
            axes[1].tick_params(axis="x", rotation=20)

    plt.xticks(rotation=35, ha="right")
    plt.tight_layout()
    fig.savefig(output_path)
    plt.close(fig)

    return Path(output_path)


def normalize_reasoning_content(value):
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, list):
        return "".join(normalize_reasoning_content(item) for item in value)
    if isinstance(value, dict):
        preferred_keys = (
            "text",
            "content",
            "reasoning",
            "reasoning_content",
            "thinking",
            "summary",
            "output_text",
        )
        parts = []
        for key in preferred_keys:
            if key in value:
                text = normalize_reasoning_content(value[key])
                if text:
                    parts.append(text)
        if parts:
            return "".join(parts)
        return "".join(normalize_reasoning_content(item) for item in value.values())
    return str(value)


def build_retained_sections(thinking_text, dialogue_output_text):
    retained_sections = []
    if (thinking_text or "").strip():
        retained_sections.append("thinking")
    if (dialogue_output_text or "").strip():
        retained_sections.append("dialogue_output")
    return retained_sections


def normalize_non_content_type(field_name):
    field_map = {
        "role": "role",
        "tool_calls": "tool_calls",
        "reasoning": "reasoning",
        "reasoning_content": "reasoning",
        "thinking": "reasoning",
        "refusal": "refusal",
        "audio": "audio",
        "function_call": "tool_calls",
    }
    return field_map.get(field_name, "other")


def inspect_stream_chunk(chunk):
    chunk_info = {
        "content": "",
        "thinking": "",
        "non_content_types": [],
        "finish_reason": None,
        "content_tokens": 0,
        "thinking_tokens": 0,
        "usage_metrics": extract_stream_usage_metrics(chunk),
    }

    choices = getattr(chunk, "choices", None) or []
    if not choices:
        return chunk_info

    choice = choices[0]
    delta_payload = extract_delta_payload(getattr(choice, "delta", None))
    chunk_info["content"] = normalize_text_content(delta_payload.pop("content", None))

    reasoning_parts = []
    for reasoning_key in ("reasoning", "reasoning_content", "thinking"):
        reasoning_value = delta_payload.pop(reasoning_key, None)
        if reasoning_value is not None:
            reasoning_parts.append(normalize_reasoning_content(reasoning_value))
    chunk_info["thinking"] = "".join(reasoning_parts)
    chunk_info["content_tokens"] = estimate_token_count(chunk_info["content"])
    chunk_info["thinking_tokens"] = estimate_token_count(chunk_info["thinking"])

    chunk_info["non_content_types"] = sorted(
        {normalize_non_content_type(field_name) for field_name in delta_payload}
        | ({"reasoning"} if chunk_info["thinking"] else set())
    )
    chunk_info["finish_reason"] = getattr(choice, "finish_reason", None) or None
    return chunk_info


def build_result_row(
    run_id,
    config,
    model,
    param_set,
    applied_params,
    display_params,
    classification,
    vram_metrics,
    dialogue_output_text,
    thinking_text,
    error_message,
    system_prompt_label="N/A",
    system_prompt_text="",
    question=None,
):
    efficiency_score = calculate_efficiency_score(classification["TPS"], vram_metrics["VRAM_Peak_MiB"])
    retained_sections = build_retained_sections(thinking_text, dialogue_output_text)
    thinking_chars = classification.get("Thinking_Chars", len(thinking_text))
    output_chars = classification.get("Output_Chars", len(dialogue_output_text))
    question = question or resolve_benchmark_questions(config)[0]
    reported_prompt_tokens = classification.get("Prompt_Tokens")
    reported_completion_tokens = classification.get("Completion_Tokens")
    reported_total_tokens = classification.get("Total_Tokens")
    estimated_prompt_tokens = estimate_token_count(
        "\n".join(
            text
            for text in (system_prompt_text, question.get("prompt", config.get("prompt", "")))
            if text
        )
    )
    thinking_tokens = classification.get("Thinking_Tokens", estimate_token_count(thinking_text))
    answer_tokens = classification.get("Output_Tokens", estimate_token_count(dialogue_output_text))
    prompt_tokens = reported_prompt_tokens if reported_prompt_tokens is not None else estimated_prompt_tokens
    completion_tokens = (
        reported_completion_tokens
        if reported_completion_tokens is not None
        else thinking_tokens + answer_tokens
    )
    total_tokens = (
        reported_total_tokens
        if reported_total_tokens is not None
        else prompt_tokens + completion_tokens
    )
    if reported_total_tokens is not None or (
        reported_prompt_tokens is not None and reported_completion_tokens is not None
    ):
        token_count_source = "backend_usage"
    elif reported_prompt_tokens is not None or reported_completion_tokens is not None:
        token_count_source = "mixed"
    else:
        token_count_source = "estimated"
    return {
        "Run_ID": run_id,
        "Status": classification["Status"],
        "Capability": config.get("capability", "chat"),
        "Suite_ID": question.get("suite_id", ""),
        "Suite_Version": question.get("suite_version", ""),
        "Question_ID": question.get("id", ""),
        "Question_Category": question.get("category", ""),
        "Question_Title": question.get("title", ""),
        "Question_Prompt": question.get("prompt", config.get("prompt", "")),
        "Expected_Output": question.get("expected_output", ""),
        "Evaluation_Guide": question.get("evaluation_guide", ""),
        "Question_Auto_Checks": question.get("auto_checks", {}),
        "Question_Wiki_Source": question.get("wiki_source_path", question.get("wiki_source", "")),
        "Question_Wiki_Excerpt_Chars": question.get("wiki_excerpt_chars", ""),
        "Output_Category": classification["Output_Category"],
        "Diagnosis": classification["Diagnosis"],
        "Finish_Reason": classification["Finish_Reason"],
        "Backend": config["backend"],
        "Model": model,
        "System_Prompt_Label": system_prompt_label,
        "System_Prompt_Text": system_prompt_text,
        "Thinking_Mode": get_thinking_mode_for_run(param_set),
        "Params": param_set.copy(),
        "Applied_Params": applied_params.copy(),
        "Config_Str": display_params,
        "TPS": classification["TPS"],
        "TTFT": classification["TTFT"],
        "First_Event_s": classification["First_Event_s"],
        "Stream_Duration_s": classification["Stream_Duration_s"],
        "Thinking_Time_s": classification.get("Thinking_Time_s"),
        "Answer_Time_s": classification.get("Answer_Time_s"),
        "Prompt_Tokens": prompt_tokens,
        "Completion_Tokens": completion_tokens,
        "Total_Tokens": total_tokens,
        "Token_Count_Source": token_count_source,
        "Total_Chunks": classification["Total_Chunks"],
        "Content_Chunks": classification["Content_Chunks"],
        "Non_Content_Chunks": classification["Non_Content_Chunks"],
        "Non_Content_Types": classification["Non_Content_Types"],
        "VRAM_Base_MiB": vram_metrics["VRAM_Base_MiB"],
        "VRAM_Peak_MiB": vram_metrics["VRAM_Peak_MiB"],
        "VRAM_Delta_MiB": vram_metrics["VRAM_Delta_MiB"],
        "VRAM_Detail": vram_metrics["VRAM_Detail"],
        "Efficiency_Score": efficiency_score,
        "Thinking_TPS": classification.get("Thinking_TPS"),
        "Output_TPS": classification.get("Output_TPS"),
        "Output_Thinking_Ratio": classification.get("Output_Thinking_Ratio"),
        "Output_Time_s": classification.get("Output_Time_s"),
        "Retained_Sections": retained_sections,
        "Thinking_Chars": thinking_chars,
        "Thinking_Tokens": thinking_tokens,
        "Thinking_Text": thinking_text,
        "Dialogue_Output_Chars": output_chars,
        "Dialogue_Output_Tokens": answer_tokens,
        "Dialogue_Output_Text": dialogue_output_text,
        "Output_Chars": output_chars,
        "Output_Tokens": answer_tokens,
        "Answer_Tokens": answer_tokens,
        "Output_Text": dialogue_output_text,
        "Error": error_message or "",
    }


def run_bench(config):
    client = OpenAI(base_url=config["url"], api_key="sk-no-key-needed")
    capability = config.get("capability", "chat")
    capability_label = {
        "chat": "chat",
        "tools": "tools",
        "suite-smoke-7": "suite-smoke-7",
        LOCAL_EXPERT_BATTLE_SUITE_ID: "local-expert-battle-48",
    }.get(capability, capability)
    param_keys = list(config["params"].keys())
    param_values = [config["params"][key] for key in param_keys]
    combos = [dict(zip(param_keys, combo)) for combo in product(*param_values)] if param_keys else [{}]
    system_prompt_variants = build_system_prompt_variants(config.get("system_prompts", []))
    benchmark_questions = resolve_benchmark_questions(config)
    include_system_prompt_in_label = len(system_prompt_variants) > 1 or system_prompt_variants[0]["label"] != "N/A"

    results = []
    total_runs = (
        len(config["models"])
        * len(combos)
        * len(system_prompt_variants)
        * len(benchmark_questions)
    )
    vram_monitoring_enabled = query_nvidia_vram_snapshot() is not None
    config["vram_monitoring"] = "nvidia-smi" if vram_monitoring_enabled else "unavailable"

    print(f"\nStarting benchmark with {total_runs} runs. Mode: {capability_label}")
    print("Chat mode focuses on TPS/TTFT. Tools mode focuses on tool_call success and first event latency.")
    print(f"System prompt variants: {len(system_prompt_variants)}")
    print(f"Questions per configuration: {len(benchmark_questions)}")
    if config.get("use_current_llama_cpp_model"):
        print(f"llama.cpp direct mode: using the model currently served by {config['url']}")
    if vram_monitoring_enabled:
        print("VRAM monitoring: enabled via nvidia-smi")
    else:
        print("VRAM monitoring: nvidia-smi not available, VRAM fields will be N/A")

    run_index = 0
    for model in config["models"]:
        if config.get("backend") == "llama.cpp" and config.get("llama_cpp_auto_switch"):
            model_entry = next(
                (
                    item
                    for item in config.get("llama_cpp_model_entries", [])
                    if item.get("name") == model
                ),
                None,
            )
            if model_entry is None:
                raise RuntimeError(f"No easy_llamacpp catalog entry resolved for {model}.")
            print(f"Switching llama.cpp model: {model} ({model_entry['path']})")
            switch_started_at = time.monotonic()
            start_llama_cpp_model(model_entry, config["url"])
            print(f"llama.cpp ready: {model} ({time.monotonic() - switch_started_at:.1f}s)")
            # Rebuild the client after the launcher has replaced the HTTP server process.
            client = OpenAI(base_url=config["url"], api_key="sk-no-key-needed")
        for system_prompt_variant, question, param_set in product(
            system_prompt_variants,
            benchmark_questions,
            combos,
        ):
                run_index += 1
                applied_params = build_backend_options(config["backend"], param_set)
                display_params = format_param_dict(param_set)
                if include_system_prompt_in_label:
                    display_params = f"{display_params} | system_prompt={system_prompt_variant['label']}"
                if question.get("id"):
                    display_params = f"{display_params} | question={question['id']}"
                print(
                    f"[{run_index}/{total_runs}] {model} | {display_params}"
                )

                request_kwargs = {}
                extra_body = build_backend_extra_body(config["backend"], param_set)
                if extra_body:
                    request_kwargs["extra_body"] = extra_body

                start_time = time.time()
                first_event_time = None
                first_content_time = None
                first_thinking_time = None
                dialogue_output_parts = []
                thinking_parts = []
                chunk_records = []
                vram_monitor = NvidiaVRAMMonitor() if vram_monitoring_enabled else None
                vram_metrics = empty_vram_metrics()
                if vram_monitor is not None and not vram_monitor.start():
                    vram_monitor = None

                try:
                    request_payload = build_chat_request_payload(
                        config,
                        model,
                        request_kwargs,
                        system_prompt_text=system_prompt_variant["text"],
                        prompt=question["prompt"],
                    )
                    stream = client.chat.completions.create(**request_payload)

                    for chunk in stream:
                        event_time = time.time()
                        if first_event_time is None:
                            first_event_time = event_time

                        chunk_info = inspect_stream_chunk(chunk)
                        chunk_records.append(chunk_info)

                        if chunk_info["content"] and first_content_time is None:
                            first_content_time = event_time
                        if chunk_info["thinking"] and first_thinking_time is None:
                            first_thinking_time = event_time
                        if chunk_info["content"]:
                            dialogue_output_parts.append(chunk_info["content"])
                        if chunk_info["thinking"]:
                            thinking_parts.append(chunk_info["thinking"])

                    end_time = time.time()
                    if vram_monitor is not None:
                        vram_metrics = vram_monitor.stop()
                    dialogue_output_text = "".join(dialogue_output_parts)
                    thinking_text = "".join(thinking_parts)
                    classification = classify_stream_result(
                        chunk_records=chunk_records,
                        start_time=start_time,
                        end_time=end_time,
                        first_event_time=first_event_time,
                        first_content_time=first_content_time,
                        first_thinking_time=first_thinking_time,
                        error_message=None,
                    )
                    classification = adjust_classification_for_capability(classification, capability)
                    results.append(
                        build_result_row(
                            run_id=run_index,
                            config=config,
                            model=model,
                            param_set=param_set,
                            applied_params=applied_params,
                            display_params=display_params,
                            system_prompt_label=system_prompt_variant["label"],
                            system_prompt_text=system_prompt_variant["text"],
                            classification=classification,
                            vram_metrics=vram_metrics,
                            dialogue_output_text=dialogue_output_text,
                            thinking_text=thinking_text,
                            error_message=None,
                            question=question,
                        )
                    )
                except Exception as exc:
                    end_time = time.time()
                    if vram_monitor is not None:
                        vram_metrics = vram_monitor.stop()
                    dialogue_output_text = "".join(dialogue_output_parts)
                    thinking_text = "".join(thinking_parts)
                    classification = classify_stream_result(
                        chunk_records=chunk_records,
                        start_time=start_time,
                        end_time=end_time,
                        first_event_time=first_event_time,
                        first_content_time=first_content_time,
                        first_thinking_time=first_thinking_time,
                        error_message=str(exc),
                    )
                    classification = adjust_classification_for_capability(classification, capability)
                    print(f"Error: {exc}")
                    results.append(
                        build_result_row(
                            run_id=run_index,
                            config=config,
                            model=model,
                            param_set=param_set,
                            applied_params=applied_params,
                            display_params=display_params,
                            system_prompt_label=system_prompt_variant["label"],
                            system_prompt_text=system_prompt_variant["text"],
                            classification=classification,
                            vram_metrics=vram_metrics,
                            dialogue_output_text=dialogue_output_text,
                            thinking_text=thinking_text,
                            error_message=str(exc),
                            question=question,
                        )
                    )

    return pd.DataFrame(results)


def save_markdown_report(df, config, report_stem, summary_excel_path=None):
    report_path = Path(f"{report_stem}.html")
    summary_df = build_summary_dataframe(df)
    outcome_summary_df = build_outcome_summary_dataframe(df)
    capability = config.get("capability", "chat")
    system_prompt_variants = build_system_prompt_variants(config.get("system_prompts", []))
    localized_summary_df = localize_report_dataframe(summary_df)
    localized_outcome_summary_df = localize_report_dataframe(outcome_summary_df)
    localized_suite_questions_df = localize_report_dataframe(
        build_suite_questions_dataframe(config)
    )
    localized_question_statistics_df = localize_report_dataframe(
        build_question_statistics_dataframe(df)
    )
    tool_call_success_summary_df = (
        build_tool_call_success_summary_dataframe(df) if capability == "tools" else pd.DataFrame()
    )
    localized_tool_call_success_summary_df = localize_report_dataframe(tool_call_success_summary_df)
    summary_download_href = Path(summary_excel_path).name if summary_excel_path else None

    matrix_rows = [
        (bilingual_text("Backend", "後端"), config["backend"]),
        (bilingual_text("Capability", "能力模式"), localize_capability_value(capability)),
        (bilingual_text("Base URL", "基礎網址"), config.get("url", "N/A")),
        (bilingual_text("Models", "模型"), ", ".join(config.get("models", []))),
        (bilingual_text("Prompt", "使用者提示"), config.get("prompt", "")),
        (
            bilingual_text("System Prompt Count", "系統提示數量"),
            len(system_prompt_variants),
        ),
        (
            bilingual_text("System Prompt Variants", "系統提示變體"),
            ", ".join(variant["label"] for variant in system_prompt_variants),
        ),
        (
            bilingual_text("Note", "備註"),
            "TPS is estimated from streaming content chunks for relative comparison. / "
            "TPS 以帶文字內容的串流片段估算，適合做相對比較。",
        ),
    ]
    if capability == "tools":
        matrix_rows.append(
            (
                bilingual_text("Tool Mode Note", "工具模式備註"),
                "Successful tool-calling runs may not emit text tokens, so `TPS` and `TTFT` can be `N/A`; "
                "focus on `Output Category=tool_call` and `First Event (s)`. / "
                "工具呼叫成功時可能不會輸出文字，因此 `TPS` 和 `TTFT` 可能是 `N/A`；"
                "請優先看 `Output Category=tool_call` 與 `First Event (s)`。",
            )
        )
    if capability in BUILTIN_SUITES or capability == LOCAL_EXPERT_BATTLE_SUITE_ID:
        suite = get_suite_definition(capability)
        matrix_rows.extend(
            [
                (bilingual_text("Suite ID", "套裝識別碼"), suite["id"]),
                (bilingual_text("Suite Version", "套裝版本"), suite["version"]),
                (bilingual_text("Question Count", "題目數量"), len(suite["questions"])),
                (bilingual_text("Suite Description", "套裝說明"), suite["description"]),
            ]
        )

    environment_notes = [
        f"VRAM monitoring: {config.get('vram_monitoring', 'unavailable')} / "
        f"顯存監控: {config.get('vram_monitoring', 'unavailable')}"
    ]
    metric_notes = [
        "`TPS (chunk/s)`: Estimated throughput from text-bearing streaming chunks. / 以含文字的串流片段估算輸出速度。",
        "`Prompt Tokens`: Prompt-side token count reported by the backend when available. / Prompt 端 token 數，優先使用後端回傳值。",
        "`Thinking Time (s)`: Time from request dispatch until the first visible answer token; includes prefill and hidden or exposed reasoning. / 從送出請求到第一個可見回答 token，包含預填充及隱藏或顯式推理。",
        "`Answer Time (s)`: Time from the first visible answer token to stream end. / 從第一個可見回答 token 到串流結束。",
        "`Total Tokens`: Prompt plus completion tokens for the question. Backend usage is preferred; otherwise the value is estimated. / 每題提示與完成 token 總和，優先採後端 usage，否則使用估算值。",
        "`Token Count Source`: `backend_usage`, `mixed`, or `estimated`. / Token 數來源分為後端 usage、混合或估算。",
        "`Prefill TPS (tok/s)`: Prompt token throughput. Uses Ollama prompt-eval timing when available; otherwise falls back to `prompt_tokens / TTFT`. / 預填充速度，優先使用 Ollama 的 prompt_eval_duration，否則回退為 `prompt_tokens / TTFT`。",
        "`Total Output (chars)`: Total visible dialogue output character count retained for the run. / 本次保留的可見回覆總字數。",
        "`Total Output Time (s)`: Time from the first output text chunk to stream end. / 從第一段輸出文字到串流結束的時間。",
        "`Thinking TPS (tok/s)`: Estimated retained-thinking token throughput from the first thinking payload to stream end. / 從第一段 thinking 到結束的思考 token 速率。",
        "`Output TPS (tok/s)`: Estimated visible dialogue-output token throughput from the first output text chunk to stream end. / 從第一段輸出文字到結束的回覆 token 速率。",
        "`Output/Thinking Ratio`: `Dialogue Output Chars / Thinking Chars`; higher means more visible answer text per retained thinking text. / 回覆字數除以 thinking 字數，越高代表可見答案佔比越高。",
        "`TTFT (s)`: Time to first text chunk. / 首段文字輸出的延遲。",
        "`First Event (s)`: Time to the first streamed event of any kind. / 第一個串流事件出現的延遲。",
        "`VRAM Peak (MiB)`: Highest observed total NVIDIA GPU memory usage during a run. / 單次測試觀測到的最高總顯存使用量。",
        "`Efficiency Score (TPS/GiB Peak)`: `TPS / (VRAM Peak in GiB)` when both values are available. / 兩者都有值時，以 `TPS / 顯存峰值 GiB` 計算效率。",
    ]
    retained_text_notes = [
        "`thinking`: Reasoning/thinking text captured from non-content reasoning payloads when the backend exposed them. / 後端有提供時，從 reasoning 類 payload 保留的思考文字。",
        "`dialogue_output`: Final conversational text emitted in normal content chunks. / 一般 content chunk 中輸出的最終對話文字。",
        "If neither exists for a run, the retained sections list will be `none`. / 若兩者都沒有，保留欄位會顯示 `none`。",
    ]
    output_diagnosis_notes = [
        "`normal_content`: Received text output as expected for chat benchmarking. / 收到正常文字輸出。",
        "`tool_call`: Received `tool_calls` payload as expected for tool benchmarking. / 收到工具呼叫 payload。",
        "`text_reply_without_tool`: Returned text, but did not emit any tool call in tool mode. / 工具模式下只回文字、沒有呼叫工具。",
        "`empty_reply`: Stream completed without text output. / 串流結束但沒有文字輸出。",
        "`non_content_stream`: Stream only carried non-text payloads. / 串流只有非文字 payload。",
        "`early_stop`: Stream ended before a complete reply or tool call was received. / 在完整回覆或工具呼叫前就中斷。",
    ]

    page_parts = [
        "<!DOCTYPE html>",
        '<html lang="zh-Hant">',
        "<head>",
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        "<title>Benchmark Report / 基準測試報告</title>",
        """<style>
body {
    margin: 0;
    background: #f4efe6;
    color: #1f2933;
    font-family: "Segoe UI", "Noto Sans TC", sans-serif;
    line-height: 1.6;
}
.page {
    max-width: 1500px;
    margin: 0 auto;
    padding: 32px 24px 64px;
}
h1, h2, h3, h4 {
    margin: 0;
    color: #13212b;
}
h1 {
    font-size: 2.1rem;
    margin-bottom: 8px;
}
h2 {
    font-size: 1.3rem;
    margin-bottom: 14px;
}
h3 {
    font-size: 1.05rem;
    margin-bottom: 12px;
}
h4 {
    font-size: 0.98rem;
    margin: 14px 0 8px;
}
p {
    margin: 0;
}
.lead {
    color: #566370;
    margin-bottom: 24px;
}
.section {
    background: #fffdfa;
    border: 1px solid #dccfbb;
    border-radius: 18px;
    padding: 22px 24px;
    margin-top: 20px;
    box-shadow: 0 10px 28px rgba(76, 61, 36, 0.06);
}
.section-header {
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 12px;
    flex-wrap: wrap;
    margin-bottom: 14px;
}
.table-wrap {
    overflow-x: auto;
}
table {
    width: 100%;
    border-collapse: collapse;
    table-layout: auto;
}
th,
td {
    padding: 10px 12px;
    border-bottom: 1px solid #e8dece;
    text-align: left;
    vertical-align: top;
}
th {
    background: #f6efe2;
    color: #263746;
    font-weight: 700;
    white-space: nowrap;
}
.matrix-table th:first-child,
.kv-table th:first-child,
.kv-table td:first-child {
    width: 220px;
}
.note-list {
    margin: 0;
    padding-left: 20px;
}
.note-list li + li {
    margin-top: 8px;
}
.run-grid {
    display: grid;
    gap: 16px;
}
.run-card {
    border: 1px solid #e5dac8;
    border-radius: 16px;
    padding: 18px;
    background: #fff;
}
.text-block {
    background: #17212b;
    color: #f8f4ed;
    padding: 16px;
    border-radius: 14px;
    overflow-x: auto;
    white-space: pre-wrap;
    word-break: break-word;
    font-family: "Cascadia Code", "Consolas", monospace;
    font-size: 0.92rem;
}
.empty-note {
    color: #6f7c88;
    font-style: italic;
}
.empty-cell {
    color: #6f7c88;
    text-align: center;
}
.metric-chart-grid {
    display: grid;
    grid-template-columns: repeat(2, minmax(0, 1fr));
    gap: 18px;
}
.metric-chart {
    border-top: 3px solid #b88a44;
    padding-top: 14px;
}
.metric-chart > p {
    color: #6f7c88;
    font-size: 0.84rem;
    margin: -6px 0 14px;
}
.metric-bar-row {
    display: grid;
    grid-template-columns: minmax(110px, 0.8fr) minmax(140px, 2fr) 72px;
    align-items: center;
    gap: 10px;
    min-height: 34px;
}
.metric-bar-label {
    overflow: hidden;
    text-overflow: ellipsis;
    white-space: nowrap;
    font-size: 0.86rem;
    font-weight: 650;
}
.metric-bar-track {
    height: 13px;
    background: #eee5d7;
    overflow: hidden;
}
.metric-bar-fill {
    height: 100%;
    min-width: 0;
}
.metric-bar-value {
    font-family: "Cascadia Code", "Consolas", monospace;
    text-align: right;
    font-variant-numeric: tabular-nums;
    font-size: 0.84rem;
}
@media (max-width: 760px) {
    .metric-chart-grid { grid-template-columns: 1fr; }
    .metric-bar-row { grid-template-columns: minmax(90px, 0.8fr) minmax(100px, 2fr) 64px; }
}
.download-button {
    display: inline-flex;
    align-items: center;
    justify-content: center;
    min-height: 40px;
    padding: 0 14px;
    border-radius: 999px;
    border: 1px solid #b88a44;
    background: #d9a84c;
    color: #1f2933;
    text-decoration: none;
    font-weight: 700;
    white-space: nowrap;
}
.download-button:hover {
    background: #e4b45e;
}
        </style>""",
        "</head>",
        "<body>",
        '<main class="page">',
        f"<h1>{bilingual_text('Benchmark Report', '基準測試報告')}</h1>",
        '<p class="lead">Same benchmark content, rendered as HTML for clearer sections and more stable table layout. / 保留原本 benchmark 內容，只將排版改成較穩定且較容易閱讀的 HTML 報告。</p>',
        '<section class="section">',
        f"<h2>{bilingual_text('Test Matrix', '測試矩陣')}</h2>",
        '<div class="table-wrap">',
        key_value_rows_to_html_table(matrix_rows, table_class="matrix-table"),
        "</div>",
        "</section>",
    ]

    if not localized_suite_questions_df.empty:
        page_parts.extend(
            [
                '<section class="section">',
                f"<h2>{bilingual_text('Suite Questions', '套裝題庫')}</h2>",
                '<div class="table-wrap">',
                dataframe_to_html_table(localized_suite_questions_df),
                "</div>",
                "</section>",
            ]
        )

    page_parts.extend(
        [
            '<section class="section">',
            f"<h2>{bilingual_text('System Prompt Variants', '系統提示變體')}</h2>",
        ]
    )

    if system_prompt_variants[0]["label"] == "N/A" and len(system_prompt_variants) == 1:
        page_parts.append(
            '<p class="empty-note">No extra system prompt was used. / 本次未額外加入 system prompt。</p>'
        )
    else:
        page_parts.append('<div class="run-grid">')
        for variant in system_prompt_variants:
            page_parts.extend(
                [
                    '<article class="run-card">',
                    f"<h3>{bilingual_text(variant['label'], '系統提示變體')}</h3>",
                    f'<pre class="text-block">{html_escape_text(normalize_output_text(variant["text"]))}</pre>',
                    "</article>",
                ]
            )
        page_parts.append("</div>")

    if not localized_question_statistics_df.empty:
        page_parts.extend(
            [
                "</section>",
                '<section class="section">',
                f"<h2>{bilingual_text('Question Statistics', '題目統計')}</h2>",
                '<div class="table-wrap">',
                dataframe_to_html_table(localized_question_statistics_df),
                "</div>",
            ]
        )

    page_parts.extend(
        [
        "</section>",
        '<section class="section">',
        f"<h2>{bilingual_text('Environment Notes', '環境說明')}</h2>",
        bullet_list_to_html(environment_notes),
        "</section>",
        '<section class="section">',
        f"<h2>{bilingual_text('Metric Notes', '指標說明')}</h2>",
        bullet_list_to_html(metric_notes),
        "</section>",
        '<section class="section">',
        f"<h2>{bilingual_text('Model Comparison', '模型長條圖對比')}</h2>",
        build_model_metric_charts_html(df),
        "</section>",
        '<section class="section">',
        f"<h2>{bilingual_text('Retained Text Notes', '保留文字說明')}</h2>",
        bullet_list_to_html(retained_text_notes),
        "</section>",
        '<section class="section">',
        f"<h2>{bilingual_text('Output Diagnosis Notes', '輸出診斷說明')}</h2>",
        bullet_list_to_html(output_diagnosis_notes),
        "</section>",
        '<section class="section">',
        '<div class="section-header">',
        f"<h2>{bilingual_text('Summary', '摘要')}</h2>",
        build_download_button(
            summary_download_href,
            "Download Excel / 下載 Excel 摘要",
        ),
        "</div>",
        '<div class="table-wrap">',
        dataframe_to_html_table(localized_summary_df),
        "</div>",
        "</section>",
        '<section class="section">',
        f"<h2>{bilingual_text('Outcome Summary', '結果摘要')}</h2>",
        '<div class="table-wrap">',
        dataframe_to_html_table(localized_outcome_summary_df),
        "</div>",
        "</section>",
    ]
    )

    if capability == "tools" and not tool_call_success_summary_df.empty:
        page_parts.extend(
            [
                '<section class="section">',
                f"<h2>{bilingual_text('Tool Call Success by Model', '各模型工具呼叫成功統計')}</h2>",
                '<div class="table-wrap">',
                dataframe_to_html_table(localized_tool_call_success_summary_df),
                "</div>",
                "</section>",
            ]
        )

    page_parts.extend(
        [
            '<section class="section">',
            f"<h2>{bilingual_text('Generated Outputs', '各次執行輸出')}</h2>",
            '<div class="run-grid">',
        ]
    )

    for _, row in df.iterrows():
        params_json = json.dumps(row["Params"], ensure_ascii=False)
        applied_params_json = json.dumps(row["Applied_Params"], ensure_ascii=False)
        retained_sections = row.get("Retained_Sections", []) or []
        retained_sections_label = (
            ", ".join(
                {
                    "thinking": bilingual_text("thinking", "思考內容"),
                    "dialogue_output": bilingual_text("dialogue_output", "對話輸出"),
                }.get(section, section)
                for section in retained_sections
            )
            if retained_sections
            else bilingual_text("none", "無")
        )
        thinking_text = row.get("Thinking_Text", "")
        dialogue_output_text = row.get("Dialogue_Output_Text", row.get("Output_Text", ""))
        system_prompt_text = row.get("System_Prompt_Text", "")

        detail_rows = [
            (bilingual_text("Status", "狀態"), localize_status_value(row["Status"])),
            (bilingual_text("Capability", "能力模式"), localize_capability_value(row.get("Capability", capability))),
            (bilingual_text("Output Category", "輸出分類"), localize_output_category_value(row["Output_Category"])),
            (bilingual_text("Diagnosis", "診斷"), row["Diagnosis"]),
            (bilingual_text("Finish Reason", "結束原因"), localize_finish_reason_value(format_text_value(row["Finish_Reason"]))),
            (bilingual_text("Backend", "後端"), row["Backend"]),
            (bilingual_text("Model", "模型"), row["Model"]),
            (bilingual_text("System Prompt", "系統提示"), localize_system_prompt_label(row.get("System_Prompt_Label", "N/A"))),
            (
                bilingual_text("Thinking Mode", "思考模式"),
                localize_thinking_mode_value(row.get("Thinking_Mode", "default")),
            ),
            (bilingual_text("System Prompt Chars", "系統提示字數"), len(system_prompt_text)),
            (bilingual_text("Params", "參數"), params_json),
            (bilingual_text("Applied Params", "實際套用參數"), applied_params_json),
            (bilingual_text("Retained Sections", "保留區塊"), retained_sections_label),
            (
                bilingual_text("Prompt Tokens", "提示詞 Token 數"),
                format_token_count(row.get("Prompt_Tokens")),
            ),
            (bilingual_text("Thinking Chars", "思考字數"), int(row.get("Thinking_Chars", len(thinking_text)))),
            (
                bilingual_text("Thinking Tokens", "思考 Token 數"),
                format_token_count(
                    row.get("Thinking_Tokens"),
                    estimate_token_count(thinking_text),
                ),
            ),
            (
                bilingual_text("Answer Tokens", "回答 Token 數"),
                format_token_count(
                    row.get("Answer_Tokens"),
                    estimate_token_count(dialogue_output_text),
                ),
            ),
            (
                bilingual_text("Completion Tokens", "完成 Token 數"),
                format_token_count(row.get("Completion_Tokens")),
            ),
            (
                bilingual_text("Total Tokens", "總 Token 數"),
                format_token_count(row.get("Total_Tokens")),
            ),
            (
                bilingual_text("Token Count Source", "Token 計數來源"),
                format_text_value(row.get("Token_Count_Source")),
            ),
            (
                bilingual_text("Dialogue Output Chars", "對話輸出字數"),
                int(row.get("Dialogue_Output_Chars", len(dialogue_output_text))),
            ),
            (
                bilingual_text("Dialogue Output Tokens", "對話輸出 Token 數"),
                format_token_count(
                    row.get("Dialogue_Output_Tokens"),
                    estimate_token_count(dialogue_output_text),
                ),
            ),
            (bilingual_text("TPS", "輸出速率"), f"{format_numeric_value(row['TPS'], 2)} chunk/s"),
            (bilingual_text("Prefill TPS", "預填充速率"), f"{format_numeric_value(row.get('Prefill_TPS'), 2)} tok/s"),
            (bilingual_text("Thinking TPS", "思考速率"), f"{format_numeric_value(row.get('Thinking_TPS'), 2)} tok/s"),
            (bilingual_text("Output TPS", "回覆速率"), f"{format_numeric_value(row.get('Output_TPS'), 2)} tok/s"),
            (bilingual_text("Output/Thinking Ratio", "輸出思考比"), format_numeric_value(row.get("Output_Thinking_Ratio"), 3)),
            (bilingual_text("Thinking Time", "思考時間"), f"{format_numeric_value(row.get('Thinking_Time_s'), 3)} s"),
            (bilingual_text("Answer Time", "回答時間"), f"{format_numeric_value(row.get('Answer_Time_s'), 3)} s"),
            (bilingual_text("TTFT", "首字延遲"), f"{format_numeric_value(row['TTFT'], 3)} s"),
            (bilingual_text("First Event", "首事件時間"), f"{format_numeric_value(row['First_Event_s'], 3)} s"),
            (bilingual_text("Stream Duration", "串流總時長"), f"{format_numeric_value(row['Stream_Duration_s'], 3)} s"),
            (bilingual_text("Total Chunks", "總片段數"), int(row["Total_Chunks"])),
            (bilingual_text("Content Chunks", "文字片段數"), int(row["Content_Chunks"])),
            (bilingual_text("Non-Content Chunks", "非文字片段數"), int(row["Non_Content_Chunks"])),
            (bilingual_text("Non-Content Types", "非文字類型"), row["Non_Content_Types"]),
            (bilingual_text("VRAM Base", "起始顯存"), format_mib_value(row["VRAM_Base_MiB"])),
            (bilingual_text("VRAM Peak", "顯存峰值"), format_mib_value(row["VRAM_Peak_MiB"])),
            (bilingual_text("VRAM Delta", "顯存增量"), format_mib_value(row["VRAM_Delta_MiB"])),
            (bilingual_text("VRAM Detail", "顯存細節"), row["VRAM_Detail"]),
            (
                bilingual_text("Efficiency Score", "效率分數"),
                f"{format_numeric_value(row['Efficiency_Score'], 3)} TPS/GiB Peak",
            ),
        ]
        if row.get("Question_ID"):
            detail_rows[2:2] = [
                (bilingual_text("Suite ID", "套裝識別碼"), row.get("Suite_ID", "")),
                (bilingual_text("Suite Version", "套裝版本"), row.get("Suite_Version", "")),
                (bilingual_text("Question ID", "題目識別碼"), row.get("Question_ID", "")),
                (
                    bilingual_text("Question Category", "題目分類"),
                    QUESTION_CATEGORY_BILINGUAL_MAP.get(
                        row.get("Question_Category"),
                        row.get("Question_Category", ""),
                    ),
                ),
                (bilingual_text("Question Title", "題目名稱"), row.get("Question_Title", "")),
                (bilingual_text("Question Prompt", "題目內容"), row.get("Question_Prompt", "")),
                (bilingual_text("Expected Output", "預期輸出"), row.get("Expected_Output", "")),
                (bilingual_text("Evaluation Guide", "評估提示"), row.get("Evaluation_Guide", "")),
            ]
        if row["Error"]:
            detail_rows.append((bilingual_text("Error", "錯誤"), row["Error"]))

        page_parts.extend(
            [
                '<article class="run-card">',
                f"<h3>{bilingual_text('Run ' + str(row['Run_ID']), '第 ' + str(row['Run_ID']) + ' 次執行')}</h3>",
                key_value_rows_to_html_table(detail_rows),
            ]
        )

        if thinking_text:
            page_parts.extend(
                [
                    f"<h4>{bilingual_text('thinking', '思考內容')}</h4>",
                    f'<pre class="text-block">{html_escape_text(normalize_output_text(thinking_text))}</pre>',
                ]
            )

        if dialogue_output_text:
            page_parts.extend(
                [
                    f"<h4>{bilingual_text('dialogue_output', '對話輸出')}</h4>",
                    f'<pre class="text-block">{html_escape_text(normalize_output_text(dialogue_output_text))}</pre>',
                ]
            )

        if not thinking_text and not dialogue_output_text:
            page_parts.append(
                '<p class="empty-note">No retained thinking or dialogue output text for this run. / 這次執行沒有保留 thinking 或 dialogue output 文字。</p>'
            )

        page_parts.append("</article>")

    page_parts.extend(["</div>", "</section>", "</main>", "</body>", "</html>"])

    with report_path.open("w", encoding="utf-8") as file:
        file.write("\n".join(page_parts))

    return report_path

import sys
from dataclasses import dataclass





CAPABILITY_DEFAULTS = {
    "chat": "Explain the long-term creep risk of PETG in 3D printing and how to reduce it.",
    "tools": (
        "Check today's weather in Taipei. If you support tools or function calling, "
        "call the `lookup_weather` tool first instead of answering directly."
    ),
    "suite-smoke-7": (
        "Built-in suite-smoke-7 uses seven fixed questions. / "
        "內建 suite-smoke-7 會依序執行七道固定題目。"
    ),
    LOCAL_EXPERT_BATTLE_SUITE_ID: (
        "Built-in Local Expert Battle runs 48 fixed questions: PLC, engineering calculations, "
        "Traditional Chinese context, and long summaries grounded in ~/wiki. / "
        "內建 Local Expert Battle 會執行 48 題固定題目，長文摘要直接擷取 ~/wiki。"
    ),
}


TABLE_COLUMNS = (
    ("idx", 4),
    ("param_key", 18),
    ("state", 8),
    ("count", 7),
    ("values", 24),
    ("range_text", 16),
)


ALLOWED_VALUE_CHARS = set("0123456789,.-+eEabcdefghijklmnopqrstuvwxyzABCDEFGHIJKLMNOPQRSTUVWXYZ_")


@dataclass
class ParamGridRow:
    key: str
    group: str
    label: str
    range_text: str
    desc: str
    supported: bool
    default_value: str
    enabled: bool = False
    raw_value: str = ""


def ordered_param_keys():
    ordered_keys = []
    for param_keys in PARAM_GROUPS.values():
        for key in param_keys:
            if key not in ordered_keys:
                ordered_keys.append(key)

    for key in PARAM_INFO:
        if key not in ordered_keys:
            ordered_keys.append(key)
    return ordered_keys


def build_param_rows(backend):
    group_by_key = {}
    for group_name, param_keys in PARAM_GROUPS.items():
        for key in param_keys:
            group_by_key[key] = group_name

    rows = []
    for key in ordered_param_keys():
        info = PARAM_INFO[key]
        rows.append(
            ParamGridRow(
                key=key,
                group=group_by_key.get(key, "Other"),
                label=info["label"],
                range_text=info["range"],
                desc=info["desc"],
                supported=backend in info["backends"],
                default_value=str(info["default"]),
                raw_value=str(info["default"]),
            )
        )
    return rows


def truncate_text(text, width):
    text = str(text)
    if width <= 0:
        return ""
    if len(text) <= width:
        return text.ljust(width)
    if width == 1:
        return text[:1]
    return text[: width - 1] + "…"


def estimate_combo_count(params):
    combo_count = 1
    for values in params.values():
        combo_count *= len(values)
    return combo_count


def validate_param_rows(rows):
    final_params = {}
    for index, row in enumerate(rows):
        if not row.supported or not row.enabled:
            continue
        try:
            final_params[row.key] = parse_csv_values(row.raw_value, param_key=row.key)
        except ValueError as exc:
            return None, index, f"{row.label}: {exc}"
    return final_params, None, None


def row_value_count(row):
    if not row.supported:
        return "LOCK"
    if not row.enabled:
        return "-"

    try:
        return str(len(parse_csv_values(row.raw_value, param_key=row.key)))
    except ValueError:
        return "ERR"


def build_grid_fragments(rows, selected_row_index, selected_column_index, message, message_style):
    from prompt_toolkit.formatted_text import to_formatted_text

    selected_row = rows[selected_row_index]
    preview_params, _, preview_error = validate_param_rows(rows)
    combo_count = "ERR" if preview_error else str(estimate_combo_count(preview_params or {}))
    selected_count = sum(1 for row in rows if row.supported and row.enabled)
    active_column_name = "state" if selected_column_index == 0 else "values"

    fragments = []
    fragments.extend(
        to_formatted_text(
            [
                ("class:title", "LLM Benchmark | Full-Page Parameter Grid\n"),
                (
                    "class:subtitle",
                    "Arrow keys move | Left/Right switch cell | Space toggles N/A/TEST | "
                    "Type values in Values (numbers or enable/disable) | Backspace deletes or goes back from State | "
                    "d restores default | Enter/Ctrl-S saves | Esc cancels\n\n",
                ),
            ]
        )
    )

    header_cells = {
        "idx": "#",
        "param_key": "Param Key",
        "state": "State",
        "count": "Count",
        "values": "Values",
        "range_text": "Range",
    }
    for column_name, width in TABLE_COLUMNS:
        fragments.append(("class:table.header", truncate_text(header_cells[column_name], width)))
        fragments.append(("class:table.header", " "))
    fragments.append(("", "\n"))
    fragments.append(("class:table.rule", "-" * (sum(width for _, width in TABLE_COLUMNS) + len(TABLE_COLUMNS))))
    fragments.append(("", "\n"))

    for row_index, row in enumerate(rows):
        row_style = "class:table.row"
        if row_index == selected_row_index:
            row_style = "class:table.row.selected"

        state_text = "LOCK" if not row.supported else "TEST" if row.enabled else "N/A"
        value_text = "backend n/a" if not row.supported else row.raw_value if row.enabled else "N/A"
        cell_values = {
            "idx": f"{row_index + 1:02d}",
            "param_key": row.key,
            "state": state_text,
            "count": row_value_count(row),
            "values": value_text,
            "range_text": row.range_text,
        }

        for column_name, width in TABLE_COLUMNS:
            cell_style = row_style
            if row_index == selected_row_index and column_name == active_column_name:
                cell_style = "class:table.cell.current"
            fragments.append((cell_style, truncate_text(cell_values[column_name], width)))
            fragments.append((row_style, " "))
        fragments.append(("", "\n"))

    support_text = ", ".join(PARAM_INFO[selected_row.key]["backends"])
    status_style = "class:status.ok" if not preview_error else "class:status.error"
    fragments.extend(
        to_formatted_text(
            [
                ("", "\n"),
                ("class:panel.title", "Selected Parameter\n"),
                ("class:panel.label", f"Key: {selected_row.key}\n"),
                ("class:panel.label", f"Label: {selected_row.label}\n"),
                ("class:panel.label", f"Groups: {selected_row.group}\n"),
                ("class:panel.label", f"Supports: {support_text}\n"),
                ("class:panel.label", f"Default values: {selected_row.default_value}\n"),
                ("class:panel.label", f"Description: {selected_row.desc}\n\n"),
                ("class:panel.title", "Config Preview\n"),
                ("class:panel.label", f"Selected params: {selected_count}\n"),
                ("class:panel.label", f"Combination count: {combo_count}\n"),
                (status_style, f"Validation: {'OK' if not preview_error else preview_error}\n"),
            ]
        )
    )

    if message:
        fragments.append((message_style, f"\n{message}\n"))

    return fragments


def edit_param_grid(backend, initial_params=None):
    from prompt_toolkit.application import Application
    from prompt_toolkit.key_binding import KeyBindings
    from prompt_toolkit.keys import Keys
    from prompt_toolkit.layout import HSplit, Layout, Window
    from prompt_toolkit.layout.controls import FormattedTextControl
    from prompt_toolkit.styles import Style

    rows = build_param_rows(backend)
    for row in rows:
        if initial_params and row.key in initial_params:
            row.enabled = True
            row.raw_value = format_param_values_for_display(row.key, initial_params[row.key])
    state = {
        "row_index": 0,
        "column_index": 0,
        "message": "",
        "message_style": "class:hint",
    }

    def set_message(text, style="class:hint"):
        state["message"] = text
        state["message_style"] = style

    def refresh(event):
        if event is not None and getattr(event, "app", None):
            event.app.invalidate()

    def current_row():
        return rows[state["row_index"]]

    def move_row(delta):
        state["row_index"] = max(0, min(len(rows) - 1, state["row_index"] + delta))

    def move_column(delta):
        state["column_index"] = (state["column_index"] + delta) % 2

    def toggle_current_row():
        row = current_row()
        if not row.supported:
            set_message(f"{row.key} is not supported on {backend}.", "class:status.warning")
            return

        row.enabled = not row.enabled
        if row.enabled and not row.raw_value.strip():
            row.raw_value = row.default_value
        set_message(f"{row.key} -> {'TEST' if row.enabled else 'N/A'}", "class:hint")

    def restore_default():
        row = current_row()
        if not row.supported:
            set_message(f"{row.key} is locked for {backend}.", "class:status.warning")
            return

        row.enabled = True
        row.raw_value = row.default_value
        set_message(f"{row.key} restored to default values.", "class:hint")

    def append_value(char):
        row = current_row()
        if not row.supported:
            set_message(f"{row.key} is locked for {backend}.", "class:status.warning")
            return

        if char not in ALLOWED_VALUE_CHARS:
            set_message(f"Unsupported character: {char!r}", "class:status.warning")
            return

        if not row.enabled:
            row.enabled = True
            if row.raw_value == "N/A":
                row.raw_value = ""
        row.raw_value += char
        set_message(f"Editing {row.key}", "class:hint")

    def backspace_value():
        row = current_row()
        if not row.supported:
            set_message(f"{row.key} is locked for {backend}.", "class:status.warning")
            return

        if not row.enabled:
            row.enabled = True
            row.raw_value = row.default_value
            set_message(f"{row.key} enabled with default values.", "class:hint")
            return

        row.raw_value = row.raw_value[:-1]
        set_message(f"Editing {row.key}", "class:hint")

    def accept(event):
        params, error_row_index, error_message = validate_param_rows(rows)
        if error_message:
            state["row_index"] = error_row_index
            state["column_index"] = 1
            set_message(error_message, "class:status.error")
            refresh(event)
            return
        event.app.exit(result=params or {})

    table_control = FormattedTextControl(
        lambda: build_grid_fragments(
            rows=rows,
            selected_row_index=state["row_index"],
            selected_column_index=state["column_index"],
            message=state["message"],
            message_style=state["message_style"],
        ),
        focusable=True,
        show_cursor=False,
    )

    root_container = HSplit([Window(content=table_control, always_hide_cursor=True)])
    style = Style.from_dict(
        {
            "title": "bold ansicyan",
            "subtitle": "ansibrightblack",
            "table.header": "bold ansiyellow",
            "table.rule": "ansibrightblack",
            "table.row": "",
            "table.row.selected": "bg:ansiblue ansiwhite",
            "table.cell.current": "reverse",
            "panel.title": "bold ansigreen",
            "panel.label": "",
            "status.ok": "ansigreen",
            "status.warning": "ansiyellow",
            "status.error": "ansired",
            "hint": "ansicyan",
        }
    )

    kb = KeyBindings()

    @kb.add("up")
    def _(event):
        move_row(-1)
        refresh(event)

    @kb.add("down")
    def _(event):
        move_row(1)
        refresh(event)

    @kb.add("left")
    def _(event):
        move_column(-1)
        refresh(event)

    @kb.add("right")
    def _(event):
        move_column(1)
        refresh(event)

    @kb.add("tab")
    def _(event):
        move_column(1)
        refresh(event)

    @kb.add("s-tab")
    def _(event):
        move_column(-1)
        refresh(event)

    @kb.add("space")
    def _(event):
        if state["column_index"] == 0:
            toggle_current_row()
            refresh(event)

    @kb.add("backspace")
    @kb.add("c-h")
    def _(event):
        if state["column_index"] == 1:
            backspace_value()
            refresh(event)
        else:
            event.app.exit(result=BACK_ACTION)

    @kb.add("delete")
    def _(event):
        if state["column_index"] == 1:
            current_row().raw_value = ""
            current_row().enabled = True
            set_message(f"Cleared {current_row().key}.", "class:hint")
            refresh(event)

    @kb.add("d")
    def _(event):
        restore_default()
        refresh(event)

    @kb.add("enter")
    @kb.add("c-s")
    def _(event):
        accept(event)

    @kb.add("escape")
    @kb.add("c-c")
    def _(event):
        event.app.exit(result=None)

    @kb.add(Keys.Any)
    def _(event):
        if state["column_index"] != 1:
            return

        char = event.data
        if not char or char not in ALLOWED_VALUE_CHARS:
            return
        append_value(char)
        refresh(event)

    application = Application(
        layout=Layout(root_container),
        key_bindings=kb,
        full_screen=True,
        mouse_support=False,
        style=style,
    )
    return application.run()


def parse_system_prompt_blocks(raw_text, expected_count):
    normalized = (raw_text or "").replace("\r\n", "\n").strip()
    if expected_count <= 0:
        return []

    blocks = [
        block.strip()
        for block in re.split(r"(?m)^\s*---\s*$", normalized)
        if block.strip()
    ]
    if len(blocks) != expected_count:
        raise ValueError(
            f"Expected {expected_count} system prompt blocks, but found {len(blocks)}. "
            "Use a line containing only --- between prompts."
        )
    return blocks


def edit_system_prompt_blocks(expected_count):
    from prompt_toolkit.application import Application
    from prompt_toolkit.key_binding import KeyBindings
    from prompt_toolkit.layout import HSplit, Layout, Window
    from prompt_toolkit.layout.controls import FormattedTextControl
    from prompt_toolkit.styles import Style
    from prompt_toolkit.widgets import Frame, TextArea

    text_area = TextArea(
        text="",
        multiline=True,
        scrollbar=True,
        line_numbers=True,
        wrap_lines=False,
        focus_on_click=True,
    )
    state = {
        "message": (
            f"Paste {expected_count} system prompt block(s). Use --- on its own line as a separator. "
            "Ctrl-S saves, Esc cancels."
        ),
        "style": "class:hint",
    }

    def set_message(text, style):
        state["message"] = text
        state["style"] = style

    kb = KeyBindings()

    @kb.add("c-s")
    def save_editor(event):
        try:
            prompts = parse_system_prompt_blocks(text_area.text, expected_count)
        except ValueError as exc:
            set_message(str(exc), "class:status.error")
            return
        event.app.exit(result=prompts)

    @kb.add("escape")
    @kb.add("c-c")
    def cancel_editor(event):
        event.app.exit(result=None)

    @kb.add("backspace")
    @kb.add("c-h")
    def go_back(event):
        if text_area.text:
            return
        event.app.exit(result=BACK_ACTION)

    root_container = HSplit(
        [
            Window(
                height=4,
                content=FormattedTextControl(
                    lambda: [
                        ("class:title", "System Prompt Editor / 系統提示編輯器\n"),
                        (
                            "class:subtitle",
                            "Paste multiple system prompts here. Use --- on its own line as a separator.\n"
                            "可在這裡直接貼上多段 system prompt，段落之間用單獨一行的 --- 分隔。\n",
                        ),
                    ]
                ),
            ),
            Frame(text_area, title="System Prompt Blocks / 系統提示區塊"),
            Window(
                height=2,
                content=FormattedTextControl(lambda: [(state["style"], state["message"])]),
            ),
        ]
    )

    app = Application(
        layout=Layout(root_container, focused_element=text_area),
        key_bindings=kb,
        full_screen=True,
        mouse_support=True,
        style=Style.from_dict(
            {
                "title": "bold ansicyan",
                "subtitle": "ansibrightblack",
                "frame.label": "bold ansiyellow",
                "status.error": "bold ansired",
                "hint": "ansicyan",
            }
        ),
    )
    return app.run()


def merge_instruction_text(instruction, back_hint):
    if instruction:
        return f"{instruction} {back_hint}"
    return back_hint


def attach_backspace_binding(question, empty_text_only=False):
    from prompt_toolkit.filters import Condition
    from prompt_toolkit.key_binding import KeyBindings, merge_key_bindings
    from prompt_toolkit.keys import Keys

    bindings = KeyBindings()
    app = question.application

    if empty_text_only:
        active_filter = Condition(
            lambda: not getattr(getattr(app, "current_buffer", None), "text", "")
        )
    else:
        active_filter = True

    @bindings.add(Keys.Backspace, eager=True, filter=active_filter)
    @bindings.add(Keys.ControlH, eager=True, filter=active_filter)
    def go_back(event):
        event.app.exit(result=BACK_ACTION)

    app.key_bindings = merge_key_bindings([app.key_bindings, bindings])
    return question


def ask_select_with_back(message, choices, default=None, instruction=None):
    question = questionary.select(
        message,
        choices=choices,
        default=default,
        instruction=merge_instruction_text(
            instruction,
            "(Backspace: previous step / Backspace 返回上一階段)",
        ),
    )
    return attach_backspace_binding(question).ask()


def ask_checkbox_with_back(message, choices, instruction=None):
    question = questionary.checkbox(
        message,
        choices=choices,
        instruction=merge_instruction_text(
            instruction,
            "(Space: select | Backspace: previous step / Space 勾選 | Backspace 返回上一階段)",
        ),
    )
    return attach_backspace_binding(question).ask()


def ask_text_with_back(message, default="", instruction=None):
    question = questionary.text(
        message,
        default=default,
        instruction=merge_instruction_text(
            instruction,
            "(Backspace on empty input: previous step / 空白時按 Backspace 返回上一階段)",
        ),
    )
    return attach_backspace_binding(question, empty_text_only=True).ask()


def ask_confirm_with_back(message, default=True, instruction=None):
    question = questionary.confirm(
        message,
        default=default,
        instruction=merge_instruction_text(
            instruction,
            "(Enter: confirm | Backspace: previous step / Enter 確認 | Backspace 返回上一階段)",
        ),
    )
    return attach_backspace_binding(question).ask()


def select_system_prompt_variants(existing_prompts=None):
    existing_prompts = existing_prompts or []
    if not existing_prompts:
        default_selection = 0
    elif len(existing_prompts) in (1, 2, 3):
        default_selection = len(existing_prompts)
    else:
        default_selection = "custom"

    while True:
        selection = ask_select_with_back(
            "System prompt variants / 系統提示變體:",
            choices=[
                Choice("N/A | no extra system prompt / 不額外加入 system prompt", value=0),
                Choice("1 variant / 1 種", value=1),
                Choice("2 variants / 2 種", value=2),
                Choice("3 variants / 3 種", value=3),
                Choice("Custom count / 自訂數量", value="custom"),
            ],
            default=default_selection,
        )
        if selection is None:
            return None
        if selection == BACK_ACTION:
            return BACK_ACTION
        if selection == 0:
            return []

        expected_count = selection
        if selection == "custom":
            while True:
                raw_count = ask_text_with_back(
                    "How many system prompt variants? / 要測幾種 system prompt？",
                    default=str(len(existing_prompts) or 1),
                )
                if raw_count is None:
                    return None
                if raw_count == BACK_ACTION:
                    break
                try:
                    expected_count = int((raw_count or "").strip())
                except ValueError:
                    print("System prompt count must be an integer. / system prompt 數量必須是整數。")
                    continue
                if expected_count < 0:
                    print("System prompt count cannot be negative. / system prompt 數量不能小於 0。")
                    continue
                if expected_count == 0:
                    return []
                break
            if raw_count == BACK_ACTION:
                continue

        prompts = edit_system_prompt_blocks(expected_count)
        if prompts is None:
            print("System prompt editor cancelled. / system prompt 編輯已取消。")
            return None
        if prompts == BACK_ACTION:
            continue
        return prompts


def print_config_review(config):
    params = config["params"]
    combo_count = estimate_combo_count(params) if params else 1
    system_prompt_variants = build_system_prompt_variants(config.get("system_prompts", []))
    questions = resolve_benchmark_questions(config)
    total_run_count = (
        len(config["models"]) * combo_count * len(system_prompt_variants) * len(questions)
    )

    print("\n" + "=" * 62)
    print("Config Review")
    print("=" * 62)
    print(f"- Backend: {config['backend']}")
    print(f"- Capability: {config['capability']}")
    print(f"- Base URL: {config['url']}")
    print(f"- Models: {', '.join(config['models'])}")
    print(f"- Param count: {len(params)}")
    print(f"- Combination count: {combo_count}")
    print(f"- Questions per combination: {len(questions)}")
    print(f"- Total runs: {total_run_count}")
    print(f"- System prompt variants: {len(system_prompt_variants)}")
    if params:
        print("- Parameter values:")
        for key, values in params.items():
            print(f"  - {key}: {values}")
    else:
        print("- Parameter values: use backend defaults only")
    if system_prompt_variants[0]["label"] == "N/A" and len(system_prompt_variants) == 1:
        print("- System prompts: N/A")
    else:
        print("- System prompt previews:")
        for variant in system_prompt_variants:
            preview = variant["text"].splitlines()[0] if variant["text"] else ""
            preview = preview[:90] + ("..." if len(preview) > 90 else "")
            print(f"  - {variant['label']}: {preview} ({len(variant['text'])} chars)")


def build_console_summary_dataframe(results_df):
    summary_df = build_summary_dataframe(results_df)
    console_columns = ["Run", "Status"]
    if "Capability" in summary_df.columns:
        console_columns.append("Capability")
    for question_column in ("Question ID", "Question Category", "Question Title"):
        if question_column in summary_df.columns:
            console_columns.append(question_column)
    if "System Prompt" in summary_df.columns:
        console_columns.append("System Prompt")
    console_columns.extend(
        [
            "Output Category",
            "Finish Reason",
            "Model",
            "Prompt Tokens",
            "Thinking Tokens",
            "Answer Tokens",
            "Completion Tokens",
            "Total Tokens",
            "Token Count Source",
            "Thinking Time (s)",
            "Answer Time (s)",
            "Prefill TPS (tok/s)",
            "TPS (chunk/s)",
            "Thinking TPS (tok/s)",
            "Output TPS (tok/s)",
            "Output/Thinking Ratio",
            "TTFT (s)",
            "First Event (s)",
            "Chunks (content/total)",
            "Config",
        ]
    )
    return summary_df[console_columns]


UI_DEFAULTS_FILENAME = "llm_expert_bench_ui_defaults.json"
SUPPORTED_BACKENDS = ("ollama", "llama.cpp")
DEFAULT_BACKEND = "llama.cpp"


def get_ui_defaults_file_path():
    return Path(UI_DEFAULTS_FILENAME)


def get_builtin_base_url(backend):
    return "http://localhost:11434/v1" if backend == "ollama" else "http://localhost:8080/v1"


def format_system_prompts_for_textarea(system_prompts):
    prompts = [str(item).strip() for item in (system_prompts or []) if str(item).strip()]
    if not prompts:
        return ""
    return f"\n{SYSTEM_PROMPT_BLOCK_SEPARATOR}\n".join(prompts)


def build_builtin_ui_defaults():
    return {
        "default_backend": DEFAULT_BACKEND,
        "default_capability": "chat",
        "backend_defaults": {
            backend: {
                "url": get_builtin_base_url(backend),
                "models": "",
                "params": {
                    row.key: {"enabled": False, "raw_value": row.default_value}
                    for row in build_param_rows(backend)
                    if row.supported
                },
            }
            for backend in SUPPORTED_BACKENDS
        },
        "capability_defaults": {
            capability: {
                "prompt": prompt_text,
                "system_prompts": [],
            }
            for capability, prompt_text in CAPABILITY_DEFAULTS.items()
        },
    }


def normalize_ui_defaults_payload(payload):
    builtin_defaults = build_builtin_ui_defaults()
    if not isinstance(payload, dict):
        payload = {}

    default_backend = payload.get("default_backend", builtin_defaults["default_backend"])
    if default_backend not in SUPPORTED_BACKENDS:
        default_backend = builtin_defaults["default_backend"]

    default_capability = payload.get("default_capability", builtin_defaults["default_capability"])
    if default_capability not in CAPABILITY_DEFAULTS:
        default_capability = builtin_defaults["default_capability"]

    raw_backend_defaults = payload.get("backend_defaults")
    if not isinstance(raw_backend_defaults, dict):
        raw_backend_defaults = {}

    normalized_backend_defaults = {}
    for backend in SUPPORTED_BACKENDS:
        builtin_backend_defaults = builtin_defaults["backend_defaults"][backend]
        submitted_backend_defaults = raw_backend_defaults.get(backend)
        if not isinstance(submitted_backend_defaults, dict):
            submitted_backend_defaults = {}

        normalized_param_defaults = {}
        submitted_param_defaults = submitted_backend_defaults.get("params")
        if not isinstance(submitted_param_defaults, dict):
            submitted_param_defaults = {}

        rows = build_param_rows(backend)
        for row in rows:
            if not row.supported:
                continue
            submitted_row_defaults = submitted_param_defaults.get(row.key)
            if not isinstance(submitted_row_defaults, dict):
                submitted_row_defaults = {}
            raw_value = str(submitted_row_defaults.get("raw_value") or row.default_value).strip()
            normalized_param_defaults[row.key] = {
                "enabled": bool(submitted_row_defaults.get("enabled")),
                "raw_value": raw_value or row.default_value,
            }

        normalized_backend_defaults[backend] = {
            "url": str(submitted_backend_defaults.get("url") or builtin_backend_defaults["url"]).strip()
            or builtin_backend_defaults["url"],
            "models": ", ".join(split_model_names(submitted_backend_defaults.get("models"))),
            "params": normalized_param_defaults,
        }

    raw_capability_defaults = payload.get("capability_defaults")
    if not isinstance(raw_capability_defaults, dict):
        raw_capability_defaults = {}

    normalized_capability_defaults = {}
    for capability, builtin_values in builtin_defaults["capability_defaults"].items():
        submitted_capability_defaults = raw_capability_defaults.get(capability)
        if not isinstance(submitted_capability_defaults, dict):
            submitted_capability_defaults = {}
        raw_system_prompts = submitted_capability_defaults.get("system_prompts", builtin_values["system_prompts"])
        if isinstance(raw_system_prompts, str):
            normalized_system_prompts = parse_system_prompt_text(raw_system_prompts)
        elif isinstance(raw_system_prompts, list):
            normalized_system_prompts = [
                str(item).strip() for item in raw_system_prompts if str(item).strip()
            ]
        else:
            normalized_system_prompts = list(builtin_values["system_prompts"])

        normalized_capability_defaults[capability] = {
            "prompt": str(submitted_capability_defaults.get("prompt") or builtin_values["prompt"]).strip()
            or builtin_values["prompt"],
            "system_prompts": normalized_system_prompts,
        }

    return {
        "default_backend": default_backend,
        "default_capability": default_capability,
        "backend_defaults": normalized_backend_defaults,
        "capability_defaults": normalized_capability_defaults,
    }


def validate_ui_defaults_payload(payload):
    normalized = normalize_ui_defaults_payload(payload)
    for backend in SUPPORTED_BACKENDS:
        rows = build_param_rows(backend)
        backend_param_defaults = normalized["backend_defaults"][backend]["params"]
        for row in rows:
            if not row.supported:
                row.enabled = False
                row.raw_value = row.default_value
                continue
            submitted_defaults = backend_param_defaults.get(row.key, {})
            row.enabled = bool(submitted_defaults.get("enabled"))
            row.raw_value = str(submitted_defaults.get("raw_value") or row.default_value).strip()
            if row.enabled and not row.raw_value:
                row.raw_value = row.default_value
        _, _, error_message = validate_param_rows(rows)
        if error_message:
            raise ValueError(f"{backend}: {error_message}")
    return normalized


def load_ui_defaults():
    defaults_path = get_ui_defaults_file_path()
    if not defaults_path.exists():
        return build_builtin_ui_defaults()
    try:
        payload = json.loads(defaults_path.read_text(encoding="utf-8"))
    except Exception:
        return build_builtin_ui_defaults()
    try:
        return validate_ui_defaults_payload(payload)
    except Exception:
        return build_builtin_ui_defaults()


def save_ui_defaults(payload):
    normalized = validate_ui_defaults_payload(payload)
    defaults_path = get_ui_defaults_file_path()
    defaults_path.write_text(
        json.dumps(normalized, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return normalized


def reset_ui_defaults():
    defaults = build_builtin_ui_defaults()
    return save_ui_defaults(defaults)


def get_default_base_url(backend, ui_defaults=None):
    normalized_backend = backend if backend in SUPPORTED_BACKENDS else DEFAULT_BACKEND
    defaults = ui_defaults or load_ui_defaults()
    return (
        str(
            defaults.get("backend_defaults", {})
            .get(normalized_backend, {})
            .get("url", get_builtin_base_url(normalized_backend))
        ).strip()
        or get_builtin_base_url(normalized_backend)
    )


def get_default_models_text(backend, ui_defaults=None):
    normalized_backend = backend if backend in SUPPORTED_BACKENDS else DEFAULT_BACKEND
    defaults = ui_defaults or load_ui_defaults()
    return str(
        defaults.get("backend_defaults", {})
        .get(normalized_backend, {})
        .get("models", "")
    ).strip()


def get_default_prompt_for_capability(capability, ui_defaults=None):
    normalized_capability = capability if capability in CAPABILITY_DEFAULTS else "chat"
    defaults = ui_defaults or load_ui_defaults()
    capability_defaults = defaults.get("capability_defaults", {}).get(normalized_capability, {})
    return str(capability_defaults.get("prompt") or CAPABILITY_DEFAULTS[normalized_capability]).strip()


def get_default_system_prompts_for_capability(capability, ui_defaults=None):
    normalized_capability = capability if capability in CAPABILITY_DEFAULTS else "chat"
    defaults = ui_defaults or load_ui_defaults()
    system_prompts = defaults.get("capability_defaults", {}).get(normalized_capability, {}).get(
        "system_prompts", []
    )
    if not isinstance(system_prompts, list):
        return []
    return [str(item).strip() for item in system_prompts if str(item).strip()]


def split_model_names(raw_models):
    if isinstance(raw_models, list):
        return [str(item).strip() for item in raw_models if str(item).strip()]
    return [name.strip() for name in str(raw_models or "").split(",") if name.strip()]


def parse_system_prompt_text(raw_text):
    normalized = (raw_text or "").replace("\r\n", "\n").strip()
    if not normalized:
        return []
    return [
        block.strip()
        for block in re.split(r"(?m)^\s*---\s*$", normalized)
        if block.strip()
    ]


def serialize_param_rows_for_ui(backend, initial_params=None, row_overrides=None):
    rows = build_param_rows(backend)
    for row in rows:
        override_values = (row_overrides or {}).get(row.key, {})
        if row.supported and isinstance(override_values, dict):
            row.enabled = bool(override_values.get("enabled"))
            row.raw_value = str(override_values.get("raw_value") or row.default_value).strip() or row.default_value
        if initial_params and row.key in initial_params:
            row.enabled = True
            row.raw_value = format_param_values_for_display(row.key, initial_params[row.key])
    return [
        {
            "key": row.key,
            "group": row.group,
            "label": row.label,
            "range_text": row.range_text,
            "desc": row.desc,
            "supported": row.supported,
            "default_value": row.default_value,
            "enabled": row.enabled,
            "raw_value": row.raw_value,
        }
        for row in rows
    ]


def build_web_ui_backend_state(backend, initial_params=None, ui_defaults=None):
    normalized_backend = backend if backend in SUPPORTED_BACKENDS else DEFAULT_BACKEND
    defaults = ui_defaults or load_ui_defaults()
    backend_defaults = defaults.get("backend_defaults", {}).get(normalized_backend, {})
    is_llama_cpp = normalized_backend == "llama.cpp"
    detected_models = get_llama_cpp_models() if is_llama_cpp else get_ollama_models()
    return {
        "backend": normalized_backend,
        "default_url": get_default_base_url(normalized_backend, ui_defaults=defaults),
        "default_models_text": get_default_models_text(normalized_backend, ui_defaults=defaults),
        "detected_models": detected_models,
        "model_catalog_label": (
            "easy_llamacpp GGUF Catalog / easy_llamacpp GGUF 模型目錄"
            if is_llama_cpp
            else "Detected Ollama Models / 偵測到的 Ollama 模型"
        ),
        "model_catalog_note": (
            f"讀取 {get_llama_cpp_launcher_root() / 'json' / 'model-index.json'}；"
            "按 Refresh 重新讀取。若不勾選模型，會直接測試 Base URL 目前已載入的 llama.cpp 模型。"
            if is_llama_cpp
            else "Click a detected model to add it to the benchmark list. / 點選模型即可加入測試清單。"
        ),
        "param_rows": serialize_param_rows_for_ui(
            normalized_backend,
            initial_params=initial_params,
            row_overrides=backend_defaults.get("params"),
        ),
    }


def summarize_config_for_ui(config):
    params = config.get("params", {}) or {}
    question_count = len(resolve_benchmark_questions(config))
    param_combination_count = estimate_combo_count(params) if params else 1
    system_prompt_count = max(1, len(config.get("system_prompts", []) or []))
    return {
        "backend": config.get("backend", DEFAULT_BACKEND),
        "capability": config.get("capability", "chat"),
        "url": config.get("url", ""),
        "models": config.get("models", []),
        "param_count": len(params),
        "combination_count": param_combination_count,
        "question_count": question_count,
        "estimated_run_count": (
            len(config.get("models", []))
            * param_combination_count
            * system_prompt_count
            * question_count
        ),
        "prompt_length": len(config.get("prompt", "")),
        "system_prompt_count": len(config.get("system_prompts", []) or []),
        "use_current_llama_cpp_model": bool(config.get("use_current_llama_cpp_model")),
        "llama_cpp_auto_switch": bool(config.get("llama_cpp_auto_switch")),
    }


def normalize_web_ui_config(payload):
    if not isinstance(payload, dict):
        raise ValueError("Invalid request payload.")

    ui_defaults = load_ui_defaults()

    backend = payload.get("backend", DEFAULT_BACKEND)
    if backend not in ("ollama", "llama.cpp"):
        raise ValueError("Unsupported backend.")

    capability = payload.get("capability", "chat")
    if capability not in CAPABILITY_DEFAULTS:
        raise ValueError("Unsupported benchmark mode.")

    url = str(payload.get("url") or get_default_base_url(backend, ui_defaults=ui_defaults)).strip()
    if not url:
        raise ValueError("Base URL is required.")

    selected_models = split_model_names(payload.get("models"))
    if not selected_models and backend != "llama.cpp":
        raise ValueError("At least one model is required.")
    use_current_llama_cpp_model = backend == "llama.cpp" and not selected_models
    models = selected_models
    if use_current_llama_cpp_model:
        models = get_openai_compatible_models(url) or [CURRENT_LLAMA_CPP_MODEL_FALLBACK]
    llama_cpp_auto_switch = (
        backend == "llama.cpp"
        and bool(selected_models)
        and bool(payload.get("llama_cpp_auto_switch"))
    )
    llama_cpp_model_entries = (
        resolve_llama_cpp_auto_switch_models(models) if llama_cpp_auto_switch else []
    )
    if llama_cpp_auto_switch:
        get_llama_cpp_switch_port(url)

    prompt = str(payload.get("prompt") or "").strip() or get_default_prompt_for_capability(
        capability, ui_defaults=ui_defaults
    )
    if capability in BUILTIN_SUITES or capability == LOCAL_EXPERT_BATTLE_SUITE_ID:
        prompt = CAPABILITY_DEFAULTS[capability]
    system_prompts = parse_system_prompt_text(payload.get("system_prompts"))
    if not system_prompts:
        system_prompts = get_default_system_prompts_for_capability(capability, ui_defaults=ui_defaults)

    submitted_params = payload.get("params") or {}
    if not isinstance(submitted_params, dict):
        raise ValueError("Invalid parameter payload.")

    rows = build_param_rows(backend)
    for row in rows:
        if not row.supported:
            row.enabled = False
            row.raw_value = row.default_value
            continue

        submitted_row = submitted_params.get(row.key) or {}
        row.enabled = bool(submitted_row.get("enabled"))
        row.raw_value = str(submitted_row.get("raw_value") or row.default_value).strip()
        if row.enabled and not row.raw_value:
            row.raw_value = row.default_value

    final_params, _, error_message = validate_param_rows(rows)
    if error_message:
        raise ValueError(error_message)

    return {
        "backend": backend,
        "capability": capability,
        "url": url,
        "models": models,
        "use_current_llama_cpp_model": use_current_llama_cpp_model,
        "llama_cpp_auto_switch": llama_cpp_auto_switch,
        "llama_cpp_model_entries": llama_cpp_model_entries,
        "params": final_params or {},
        "prompt": prompt,
        "system_prompts": system_prompts,
    }


def build_report_file_url(path_or_name):
    file_name = path_or_name.name if isinstance(path_or_name, Path) else str(path_or_name)
    return f"/report-files/{quote(file_name)}"


def build_artifact_link(label, path):
    if not path:
        return None
    artifact_path = Path(path)
    if not artifact_path.exists():
        return None
    return {
        "label": label,
        "name": artifact_path.name,
        "url": build_report_file_url(artifact_path),
    }


def build_report_entry(report_path):
    html_path = Path(report_path)
    if not html_path.exists():
        return None

    artifact_links = []
    for label, candidate_path in (
        ("HTML Report", html_path),
        ("Chart", html_path.with_suffix(".png")),
        ("Summary Excel", html_path.with_name(f"{html_path.stem}_summary.xlsx")),
        ("Raw Outputs", html_path.with_name(f"{html_path.stem}_outputs.jsonl")),
        ("Best Config", html_path.with_name("best_config.json")),
        ("Ollama Modelfile", html_path.with_name("Ollama_Modelfile_Suggest")),
    ):
        link = build_artifact_link(label, candidate_path)
        if link:
            artifact_links.append(link)

    stat = html_path.stat()
    return {
        "id": html_path.stem,
        "title": html_path.stem,
        "html_name": html_path.name,
        "html_url": build_report_file_url(html_path),
        "modified_at": time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(stat.st_mtime)),
        "modified_ts": stat.st_mtime,
        "size_kib": round(stat.st_size / 1024, 1),
        "artifact_links": artifact_links,
    }


def list_report_entries(report_dir=None, limit=30):
    report_root = Path(report_dir) if report_dir else ensure_report_output_dir()
    if not report_root.exists():
        return []

    html_paths = sorted(
        report_root.glob("bench_*.html"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    entries = []
    for html_path in html_paths[:limit]:
        entry = build_report_entry(html_path)
        if entry:
            entries.append(entry)
    return entries


def render_local_expert_battle_report(raw_outputs_path, report_stem):
    script_path = Path(__file__).with_name("report.py")
    if not script_path.is_file():
        raise RuntimeError(f"Local Expert Battle reporter is missing: {script_path}")
    battle_path = Path(f"{report_stem}_battle.html")
    review_path = Path(f"{report_stem}_manual_review.json")
    completed = subprocess.run(
        [
            sys.executable,
            str(script_path),
            "--input",
            str(raw_outputs_path),
            "--output",
            str(battle_path),
            "--review",
            str(review_path),
        ],
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    if completed.returncode != 0:
        detail = (completed.stderr or completed.stdout).strip()
        raise RuntimeError(detail or "report.py exited without a diagnostic")
    return battle_path, review_path


def run_benchmark_workflow(config, progress_callback=None):
    warnings = []

    def report_progress(message):
        print(message)
        if progress_callback:
            progress_callback(message)

    report_progress("Starting benchmark run...")
    results_df = run_bench(config)
    if results_df.empty:
        raise RuntimeError("No benchmark rows were produced.")

    ok_count = int((results_df["Status"] == "ok").sum())
    warning_count = int((results_df["Status"] == "warning").sum())
    error_count = int((results_df["Status"] == "error").sum())
    capability = config.get("capability", "chat")
    report_dir = ensure_report_output_dir()
    report_stem = report_dir / f"bench_{config['backend']}_{capability}_{time.strftime('%Y%m%d_%H%M%S')}"

    report_progress("Saving raw outputs...")
    raw_outputs_path = save_raw_outputs(results_df, report_stem)
    battle_report_path = None
    battle_review_path = None
    if capability == LOCAL_EXPERT_BATTLE_SUITE_ID:
        try:
            report_progress("Scoring objective Battle questions and creating manual review template...")
            battle_report_path, battle_review_path = render_local_expert_battle_report(
                raw_outputs_path,
                report_stem,
            )
        except Exception as exc:
            warning_text = f"Battle report generation failed: {exc}"
            warnings.append(warning_text)
            print(warning_text)

    console_df = build_console_summary_dataframe(results_df)
    console_summary_text = dataframe_to_text_table(console_df)

    report_path = None
    chart_path = None
    summary_excel_path = None
    best_config_artifacts = None

    try:
        report_progress("Exporting summary workbook...")
        summary_excel_path = save_summary_excel_workbook(results_df, config, report_stem)
    except Exception as exc:
        warning_text = f"Summary Excel export failed: {exc}"
        warnings.append(warning_text)
        print(warning_text)

    try:
        report_progress("Rendering HTML report...")
        report_path = save_markdown_report(
            results_df,
            config,
            report_stem,
            summary_excel_path=summary_excel_path,
        )
    except Exception as exc:
        warning_text = f"Report generation failed: {exc}"
        warnings.append(warning_text)
        print(warning_text)

    try:
        report_progress("Rendering chart...")
        chart_path = plot_results(results_df, f"{report_stem}.png", capability=capability)
    except Exception as exc:
        warning_text = f"Chart generation failed: {exc}"
        warnings.append(warning_text)
        print(warning_text)

    try:
        report_progress("Exporting best config...")
        best_config_artifacts = export_best_config(results_df, config, output_dir=report_dir)
    except Exception as exc:
        warning_text = f"best_config export failed: {exc}"
        warnings.append(warning_text)
        print(warning_text)

    latest_report_entry = build_report_entry(report_path) if report_path else None
    artifact_links = []
    for label, path in (
        ("HTML Report", report_path),
        ("Battle Result", battle_report_path),
        ("Manual Review Template", battle_review_path),
        ("Summary Excel", summary_excel_path),
        ("Chart", chart_path),
        ("Raw Outputs", raw_outputs_path),
    ):
        link = build_artifact_link(label, path)
        if link:
            artifact_links.append(link)

    if best_config_artifacts:
        for label, path in (
            ("Best Config", best_config_artifacts.get("best_config_path")),
            ("Ollama Modelfile", best_config_artifacts.get("modelfile_path")),
        ):
            link = build_artifact_link(label, path)
            if link:
                artifact_links.append(link)

    report_progress("Benchmark run complete.")
    return {
        "config_summary": summarize_config_for_ui(config),
        "counts": {
            "ok": ok_count,
            "warning": warning_count,
            "error": error_count,
        },
        "console_summary_text": console_summary_text,
        "report_dir": str(report_dir),
        "warnings": warnings,
        "artifact_links": artifact_links,
        "latest_report": latest_report_entry,
    }


class BenchmarkWebUiState:
    def __init__(self):
        self._lock = threading.Lock()
        self._job = self._build_idle_job()

    def _build_idle_job(self):
        return {
            "id": None,
            "status": "idle",
            "message": "Ready. / 就緒。",
            "error": "",
            "logs": [],
            "started_at": None,
            "ended_at": None,
            "config_summary": None,
            "result": None,
        }

    def snapshot(self):
        with self._lock:
            return copy.deepcopy(self._job)

    def start_job(self, config):
        with self._lock:
            if self._job.get("status") == "running":
                raise RuntimeError("A benchmark is already running. / 目前已有 benchmark 正在執行。")
            job_id = f"job-{int(time.time() * 1000)}"
            self._job = {
                "id": job_id,
                "status": "running",
                "message": "Preparing benchmark run... / 正在準備 benchmark...",
                "error": "",
                "logs": ["Preparing benchmark run... / 正在準備 benchmark..."],
                "started_at": time.strftime("%Y-%m-%d %H:%M:%S"),
                "ended_at": None,
                "config_summary": summarize_config_for_ui(config),
                "result": None,
            }
            return job_id

    def append_log(self, job_id, message):
        with self._lock:
            if self._job.get("id") != job_id:
                return
            self._job["message"] = message
            self._job["logs"].append(message)
            self._job["logs"] = self._job["logs"][-120:]

    def complete_job(self, job_id, result):
        with self._lock:
            if self._job.get("id") != job_id:
                return
            self._job["status"] = "succeeded"
            self._job["message"] = "Benchmark completed. / Benchmark 已完成。"
            self._job["ended_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
            self._job["result"] = result
            self._job["logs"].append("Benchmark completed. / Benchmark 已完成。")
            self._job["logs"] = self._job["logs"][-120:]

    def fail_job(self, job_id, error_text):
        with self._lock:
            if self._job.get("id") != job_id:
                return
            self._job["status"] = "failed"
            self._job["message"] = "Benchmark failed. / Benchmark 失敗。"
            self._job["error"] = error_text
            self._job["ended_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
            self._job["logs"].append(error_text)
            self._job["logs"] = self._job["logs"][-120:]


def build_ui_defaults_api_payload(ui_defaults=None):
    defaults = ui_defaults or load_ui_defaults()
    return {
        "ui_defaults": defaults,
        "ui_defaults_path": str(get_ui_defaults_file_path().resolve()),
        "backend_catalog": {
            backend: build_web_ui_backend_state(backend, ui_defaults=defaults)
            for backend in SUPPORTED_BACKENDS
        },
        "capability_defaults": {
            capability: get_default_prompt_for_capability(capability, ui_defaults=defaults)
            for capability in CAPABILITY_DEFAULTS
        },
        "capability_system_prompt_defaults": {
            capability: format_system_prompts_for_textarea(
                get_default_system_prompts_for_capability(capability, ui_defaults=defaults)
            )
            for capability in CAPABILITY_DEFAULTS
        },
        "default_backend": defaults.get("default_backend", DEFAULT_BACKEND),
        "default_capability": defaults.get("default_capability", "chat"),
    }


def build_web_ui_bootstrap_payload(app_state):
    defaults_payload = build_ui_defaults_api_payload()
    default_backend = defaults_payload["default_backend"]
    payload = {
        "app_title": "DIY LLM Benchmark / DIY LLM Benchmark 控制台",
        "backends": [
            {"value": "ollama", "label": "Ollama / Ollama"},
            {"value": "llama.cpp", "label": "llama.cpp / llama.cpp"},
        ],
        "capabilities": [
            {
                "value": "chat",
                "label": "Chat / 對話",
                "description": "Standard response benchmark. / 一般對話輸出 benchmark。",
            },
            {
                "value": "tools",
                "label": "Tools / 工具呼叫",
                "description": "Check tool_call output. / 檢查是否輸出 tool_call。",
            },
            {
                "value": "suite-smoke-7",
                "label": "Suite Smoke 7 / 七項能力套裝",
                "description": "Seven fixed questions covering core capabilities. / 以七道固定題目涵蓋核心能力。",
            },
            {
                "value": LOCAL_EXPERT_BATTLE_SUITE_ID,
                "label": "Local Expert Battle 48 / 在地工程專家對戰",
                "description": "48 questions: PLC, engineering calculations, Traditional Chinese, and ~/wiki summaries. / PLC、工程計算、繁中語境與 ~/wiki 長文摘要各 12 題。",
            },
        ],
        "backend_state": defaults_payload["backend_catalog"][default_backend],
        "reports": list_report_entries(),
        "job": app_state.snapshot(),
    }
    payload.update(defaults_payload)
    return payload


def build_single_file_benchmark_ui_html():
    return """<!DOCTYPE html>
<html lang="zh-Hant">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>DIY LLM Benchmark / DIY LLM Benchmark 控制台</title>
  <style>
    :root {
      --paper: #f4efe6;
      --paper-strong: #efe7d8;
      --ink: #1f2430;
      --ink-soft: #586274;
      --panel: rgba(255, 251, 245, 0.92);
      --panel-strong: rgba(255, 248, 239, 0.98);
      --line: rgba(59, 45, 30, 0.14);
      --accent: #bd6741;
      --accent-strong: #944424;
      --accent-soft: rgba(189, 103, 65, 0.14);
      --ok: #2f7d4d;
      --warn: #b2751a;
      --danger: #a23b3b;
      --shadow: 0 24px 60px rgba(43, 34, 24, 0.12);
      --radius-xl: 28px;
      --radius-lg: 20px;
      --radius-md: 14px;
      --radius-sm: 10px;
      --mono: "Cascadia Code", "JetBrains Mono", "SFMono-Regular", Consolas, monospace;
      --display: "Iowan Old Style", "Palatino Linotype", "Book Antiqua", "Noto Serif TC", serif;
      --body: "Aptos", "Segoe UI", "Noto Sans TC", sans-serif;
    }

    * { box-sizing: border-box; }
    html, body { margin: 0; min-height: 100%; }
    body {
      font-family: var(--body);
      color: var(--ink);
      background:
        radial-gradient(circle at top left, rgba(189, 103, 65, 0.16), transparent 32%),
        radial-gradient(circle at 80% 20%, rgba(45, 109, 113, 0.12), transparent 28%),
        linear-gradient(180deg, #f8f4ed 0%, var(--paper) 100%);
      letter-spacing: 0.01em;
    }

    .shell {
      width: min(1680px, calc(100vw - 40px));
      margin: 24px auto 40px;
    }

    .hero {
      display: grid;
      grid-template-columns: minmax(0, 1.4fr) minmax(320px, 0.8fr);
      gap: 22px;
      padding: 24px 28px;
      border: 1px solid var(--line);
      border-radius: var(--radius-xl);
      background: linear-gradient(135deg, rgba(255,255,255,0.75), rgba(255,247,237,0.92));
      box-shadow: var(--shadow);
      backdrop-filter: blur(20px) saturate(115%);
    }

    .hero h1 {
      margin: 0 0 10px;
      font-family: var(--display);
      font-size: clamp(2.1rem, 4vw, 3.6rem);
      line-height: 0.95;
      letter-spacing: -0.03em;
      text-wrap: balance;
    }

    .hero p {
      margin: 0;
      max-width: 62ch;
      color: var(--ink-soft);
      font-size: 1rem;
      line-height: 1.7;
      text-wrap: pretty;
    }

    .hero-meta {
      display: grid;
      gap: 12px;
      align-content: end;
      justify-items: stretch;
    }

    .meta-card, .panel {
      border: 1px solid var(--line);
      border-radius: var(--radius-lg);
      background: var(--panel);
      box-shadow: 0 16px 38px rgba(43, 34, 24, 0.08);
    }

    .meta-card {
      padding: 16px 18px;
    }

    .meta-label {
      display: block;
      margin-bottom: 6px;
      color: var(--ink-soft);
      font-size: 0.78rem;
      text-transform: uppercase;
      letter-spacing: 0.14em;
    }

    .meta-value {
      font-size: 1.15rem;
      font-weight: 700;
    }

    .layout {
      display: grid;
      grid-template-columns: minmax(420px, 0.95fr) minmax(0, 1.25fr);
      gap: 22px;
      margin-top: 22px;
      align-items: start;
    }

    .panel {
      padding: 22px;
      overflow: hidden;
    }

    .panel h2,
    .panel h3 {
      margin: 0;
      font-family: var(--display);
      line-height: 1;
      letter-spacing: -0.02em;
    }

    .panel h2 { font-size: 1.65rem; margin-bottom: 8px; }
    .panel h3 { font-size: 1.15rem; }

    .section-note {
      margin: 0;
      color: var(--ink-soft);
      line-height: 1.6;
      text-wrap: pretty;
    }

    .stack {
      display: grid;
      gap: 18px;
    }

    .form-grid {
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 14px;
    }

    .field {
      display: grid;
      gap: 7px;
    }

    .field-wide {
      grid-column: 1 / -1;
    }

    label,
    .field-label {
      font-size: 0.86rem;
      font-weight: 700;
      color: var(--ink-soft);
      letter-spacing: 0.02em;
    }

    input,
    select,
    textarea,
    button {
      font: inherit;
    }

    input,
    select,
    textarea {
      width: 100%;
      padding: 12px 14px;
      border: 1px solid rgba(54, 43, 29, 0.16);
      border-radius: var(--radius-sm);
      background: rgba(255, 255, 255, 0.78);
      color: var(--ink);
      transition: border-color 160ms ease, box-shadow 160ms ease, background 160ms ease;
    }

    input:focus,
    select:focus,
    textarea:focus {
      outline: none;
      border-color: rgba(148, 68, 36, 0.45);
      box-shadow: 0 0 0 4px rgba(189, 103, 65, 0.12);
      background: rgba(255, 255, 255, 0.96);
    }

    textarea {
      min-height: 112px;
      resize: vertical;
      line-height: 1.55;
    }

    .hint {
      color: var(--ink-soft);
      font-size: 0.78rem;
      line-height: 1.55;
    }

    .button-row {
      display: flex;
      gap: 12px;
      flex-wrap: wrap;
    }

    .button {
      appearance: none;
      border: 0;
      border-radius: 999px;
      padding: 12px 18px;
      cursor: pointer;
      font-weight: 700;
      letter-spacing: 0.02em;
      transition: transform 140ms ease, background 140ms ease, opacity 140ms ease;
    }

    .button:hover { transform: translateY(-1px); }
    .button:disabled {
      cursor: wait;
      opacity: 0.66;
      transform: none;
    }

    .button-primary {
      background: linear-gradient(135deg, var(--accent), var(--accent-strong));
      color: #fff9f5;
      box-shadow: 0 16px 34px rgba(148, 68, 36, 0.22);
    }

    .button-secondary {
      background: rgba(35, 46, 61, 0.08);
      color: var(--ink);
      border: 1px solid rgba(35, 46, 61, 0.08);
    }

    .pill-row {
      display: flex;
      gap: 8px;
      flex-wrap: wrap;
    }

    .pill {
      display: inline-flex;
      align-items: center;
      gap: 6px;
      padding: 8px 12px;
      border-radius: 999px;
      border: 1px solid rgba(54, 43, 29, 0.14);
      background: rgba(255,255,255,0.72);
      color: var(--ink);
      font-size: 0.86rem;
    }

    .pill input {
      width: auto;
      margin: 0;
      padding: 0;
      box-shadow: none;
    }

    .switch-line {
      display: inline-flex;
      align-items: center;
      gap: 9px;
      width: fit-content;
      color: var(--ink);
      cursor: pointer;
    }

    .switch-line input {
      width: auto;
      margin: 0;
      padding: 0;
      box-shadow: none;
    }

    .status-band {
      display: grid;
      gap: 14px;
    }

    .status-header {
      display: flex;
      justify-content: space-between;
      align-items: flex-start;
      gap: 12px;
      flex-wrap: wrap;
    }

    .status-chip {
      display: inline-flex;
      align-items: center;
      gap: 8px;
      padding: 8px 12px;
      border-radius: 999px;
      font-weight: 700;
      letter-spacing: 0.02em;
      text-transform: capitalize;
      background: rgba(35, 46, 61, 0.08);
    }

    .status-idle { color: var(--ink-soft); }
    .status-running { color: var(--warn); background: rgba(178, 117, 26, 0.12); }
    .status-succeeded { color: var(--ok); background: rgba(47, 125, 77, 0.12); }
    .status-failed { color: var(--danger); background: rgba(162, 59, 59, 0.12); }

    .summary-strip {
      display: grid;
      grid-template-columns: repeat(4, minmax(0, 1fr));
      gap: 10px;
    }

    .summary-item {
      padding: 14px 16px;
      border: 1px solid var(--line);
      border-radius: var(--radius-md);
      background: rgba(255,255,255,0.68);
    }

    .summary-item strong {
      display: block;
      font-size: 1.3rem;
      margin-bottom: 4px;
    }

    .summary-item span {
      color: var(--ink-soft);
      font-size: 0.82rem;
      text-transform: uppercase;
      letter-spacing: 0.12em;
    }

    .param-table-wrap,
    .report-list {
      border: 1px solid var(--line);
      border-radius: var(--radius-md);
      background: rgba(255,255,255,0.66);
      overflow: hidden;
    }

    .param-table-wrap {
      max-height: 460px;
      overflow: auto;
    }

    table {
      width: 100%;
      border-collapse: collapse;
    }

    th,
    td {
      text-align: left;
      padding: 12px 12px;
      border-bottom: 1px solid rgba(54, 43, 29, 0.1);
      vertical-align: top;
    }

    th {
      position: sticky;
      top: 0;
      background: rgba(247, 241, 232, 0.96);
      z-index: 1;
      font-size: 0.78rem;
      text-transform: uppercase;
      letter-spacing: 0.08em;
      color: var(--ink-soft);
    }

    tr:last-child td { border-bottom: 0; }
    .unsupported-row { opacity: 0.58; }

    .count-strip {
      display: flex;
      gap: 10px;
      flex-wrap: wrap;
      margin-top: 12px;
    }

    .count-badge {
      padding: 8px 12px;
      border-radius: 999px;
      background: rgba(255,255,255,0.8);
      border: 1px solid var(--line);
      font-size: 0.82rem;
      color: var(--ink-soft);
    }

    .log-box,
    .console-box {
      border: 1px solid var(--line);
      border-radius: var(--radius-md);
      padding: 14px;
      background: #201a16;
      color: #f9ead7;
      font-family: var(--mono);
      font-size: 0.84rem;
      line-height: 1.6;
      max-height: 220px;
      overflow: auto;
      white-space: pre-wrap;
      word-break: break-word;
    }

    .report-grid {
      display: grid;
      gap: 12px;
      padding: 12px;
    }

    .report-card {
      display: grid;
      gap: 10px;
      padding: 14px;
      border: 1px solid rgba(54, 43, 29, 0.1);
      border-radius: var(--radius-md);
      background: linear-gradient(180deg, rgba(255,255,255,0.8), rgba(255,249,243,0.92));
    }

    .report-card.active {
      border-color: rgba(148, 68, 36, 0.28);
      box-shadow: inset 0 0 0 1px rgba(148, 68, 36, 0.18);
    }

    .report-card h4 {
      margin: 0;
      font-size: 1rem;
      line-height: 1.35;
      word-break: break-word;
    }

    .report-meta {
      color: var(--ink-soft);
      font-size: 0.82rem;
      display: flex;
      gap: 8px 14px;
      flex-wrap: wrap;
    }

    .artifact-links {
      display: flex;
      gap: 8px;
      flex-wrap: wrap;
    }

    .artifact-links a,
    .artifact-links button {
      text-decoration: none;
      background: rgba(189, 103, 65, 0.1);
      color: var(--accent-strong);
      border: 1px solid rgba(189, 103, 65, 0.14);
      padding: 8px 10px;
      border-radius: 999px;
      font-size: 0.78rem;
      font-weight: 700;
    }

    .artifact-links button {
      cursor: pointer;
      font: inherit;
    }

    .viewer-shell {
      display: grid;
      gap: 12px;
    }

    .viewer-frame {
      width: 100%;
      min-height: 720px;
      border: 1px solid var(--line);
      border-radius: var(--radius-lg);
      background: white;
    }

    .maintenance-panel {
      margin-top: 22px;
    }

    .maintenance-panel[hidden] {
      display: none;
    }

    .maintenance-toggle-row {
      display: flex;
      justify-content: flex-end;
      margin-top: 22px;
    }

    .maintenance-grid {
      display: grid;
      grid-template-columns: repeat(2, minmax(0, 1fr));
      gap: 18px;
    }

    .maintenance-path {
      display: inline-flex;
      align-items: center;
      min-height: 40px;
      padding: 8px 12px;
      border-radius: 999px;
      border: 1px solid rgba(54, 43, 29, 0.12);
      background: rgba(255,255,255,0.78);
      color: var(--ink-soft);
      font-size: 0.78rem;
      word-break: break-all;
    }

    .empty-state {
      padding: 22px;
      color: var(--ink-soft);
      text-align: center;
      line-height: 1.7;
    }

    .top-alert {
      display: none;
      margin-top: 18px;
      padding: 14px 16px;
      border-radius: var(--radius-md);
      border: 1px solid rgba(162, 59, 59, 0.2);
      background: rgba(162, 59, 59, 0.08);
      color: var(--danger);
      font-weight: 700;
    }

    .top-alert.visible { display: block; }

    @media (max-width: 1220px) {
      .hero,
      .layout {
        grid-template-columns: 1fr;
      }
      .summary-strip {
        grid-template-columns: repeat(2, minmax(0, 1fr));
      }
      .maintenance-grid {
        grid-template-columns: 1fr;
      }
    }

    @media (max-width: 760px) {
      .shell {
        width: min(100vw - 20px, 1680px);
        margin: 12px auto 22px;
      }
      .hero,
      .panel {
        padding: 18px;
      }
      .form-grid,
      .summary-strip {
        grid-template-columns: 1fr;
      }
      th:nth-child(3),
      td:nth-child(3),
      th:nth-child(4),
      td:nth-child(4) {
        display: none;
      }
      .viewer-frame {
        min-height: 520px;
      }
    }
  </style>
</head>
<body>
  <div class="shell">
    <section class="hero">
      <div>
        <h1>DIY LLM Benchmark Control Room / DIY LLM Benchmark 控制台</h1>
        <p>把原本 terminal 式設定流程收斂成一個單檔 HTML 控制台，左邊編 benchmark 配置，右邊直接看執行狀態、產物連結與歷史報告，不用來回切視窗。</p>
      </div>
      <div class="hero-meta">
        <div class="meta-card">
          <span class="meta-label">UI Mode / 介面模式</span>
          <div class="meta-value">Single-file HTML / 單檔 HTML</div>
        </div>
        <div class="meta-card">
          <span class="meta-label">Report Flow / 報告流程</span>
          <div class="meta-value">Preview + download + history / 預覽 + 下載 + 歷史紀錄</div>
        </div>
      </div>
    </section>

    <div id="top-alert" class="top-alert"></div>

    <main class="layout">
      <section class="panel stack">
        <div>
          <h2>Benchmark Setup / 測試設定</h2>
          <p class="section-note">保留原本 Python benchmark 核心，只把配置與操作面改成瀏覽器控制台。參數區會依 backend 自動顯示支援狀態。</p>
        </div>

        <div class="form-grid">
          <div class="field">
            <label for="backend">Backend / 後端</label>
            <select id="backend"></select>
          </div>
          <div class="field">
            <label for="capability">Benchmark Mode / 測試模式</label>
            <select id="capability"></select>
          </div>
          <div class="field field-wide">
            <label for="base-url">Base URL / 基礎網址</label>
            <input id="base-url" type="text" spellcheck="false">
          </div>
          <div class="field field-wide">
            <span id="model-catalog-label" class="field-label">Detected Ollama Models / 偵測到的 Ollama 模型</span>
            <div id="detected-models" class="pill-row"></div>
            <div id="model-catalog-note" class="hint"></div>
            <div class="button-row">
              <button id="refresh-models" class="button button-secondary" type="button">Refresh Model List / 重新整理模型清單</button>
            </div>
          </div>
          <div class="field field-wide">
            <label for="models-input">Models / 模型</label>
            <input id="models-input" type="text" spellcheck="false" placeholder="qwen3.5:latest, llama.cpp-model">
            <div class="hint">用逗號分隔。llama.cpp 留空時會直接連線 Base URL，測試目前 llama-server 已載入的模型。</div>
          </div>
          <div id="llama-batch-switch-field" class="field field-wide" hidden>
            <label class="switch-line" for="llama-cpp-auto-switch">
              <input id="llama-cpp-auto-switch" type="checkbox" checked>
              <span>Batch switch selected GGUF models / 自動依序切換所選 GGUF 並測試</span>
            </label>
            <div class="hint">有勾選模型時才會由 easy_llamacpp 逐一換模；未勾選時不換模，直接使用目前 Base URL 的 llama-server。</div>
          </div>
          <div class="field field-wide">
            <label for="prompt-input">Benchmark Prompt / 測試提示</label>
            <textarea id="prompt-input"></textarea>
            <div id="prompt-hint" class="hint">This prompt is sent once per configuration. / 此提示會依每組設定送出一次。</div>
          </div>
          <div class="field field-wide">
            <label for="system-prompts-input">System Prompt Variants / System Prompt 變體</label>
            <textarea id="system-prompts-input" placeholder="每段 system prompt 之間用一行 --- 分隔"></textarea>
            <div class="hint">留空表示不使用額外 system prompt。若有多段，請用單獨一行 <code>---</code> 分隔。</div>
          </div>
        </div>

        <div class="stack">
          <div>
            <h3>Parameter Matrix / 參數矩陣</h3>
            <p class="section-note">支援欄位可直接勾選啟用，值欄延續原本 CSV 寫法，例如 <code>0.1, 0.8</code> 或 <code>enable, disable</code>。</p>
          </div>
          <div class="param-table-wrap">
            <table>
              <thead>
                <tr>
                  <th>Use / 啟用</th>
                  <th>Key / 鍵</th>
                  <th>Group / 群組</th>
                  <th>Range / 範圍</th>
                  <th>Values / 值</th>
                </tr>
              </thead>
              <tbody id="param-table-body"></tbody>
            </table>
          </div>
          <div id="param-summary" class="count-strip"></div>
        </div>

        <div class="button-row">
          <button id="start-button" class="button button-primary" type="button">Start Benchmark / 開始測試</button>
          <button id="refresh-reports" class="button button-secondary" type="button">Refresh Reports / 重新整理報告</button>
          <button id="shutdown-ui" class="button button-secondary" type="button">Close Local UI / 關閉本機 UI</button>
        </div>
      </section>

      <section class="stack">
        <section class="panel status-band">
          <div class="status-header">
            <div>
              <h2>Run Status / 執行狀態</h2>
              <p id="job-message" class="section-note">Ready. / 就緒。</p>
            </div>
            <div id="job-status-chip" class="status-chip status-idle">idle / 就緒</div>
          </div>

          <div id="job-summary-strip" class="summary-strip"></div>

          <div class="stack">
            <div>
              <h3>Progress Log / 進度紀錄</h3>
            </div>
            <div id="job-logs" class="log-box">No job started yet. / 尚未開始任務。</div>
          </div>

          <div class="stack">
            <div>
              <h3>Console Summary / 終端摘要</h3>
            </div>
            <div id="console-summary" class="console-box">Benchmark table output will appear here after a run. / 執行後會在這裡顯示 benchmark 表格摘要。</div>
          </div>

          <div>
            <h3>Run Artifacts / 執行產物</h3>
            <div id="run-artifacts" class="artifact-links"></div>
          </div>
        </section>

        <section class="panel stack">
          <div class="status-header">
            <div>
              <h2>Reports / 報告列表</h2>
              <p class="section-note">這裡會列出最新產生的 HTML report，點選後可直接在下方預覽，也可以開新分頁或下載同一組產物。</p>
            </div>
          </div>
          <div class="report-list">
            <div id="report-list" class="report-grid"></div>
          </div>
        </section>

        <section class="panel viewer-shell">
          <div class="status-header">
            <div>
              <h2>Report Viewer / 報告檢視器</h2>
              <p id="viewer-caption" class="section-note">Select a report to preview it here. / 請選擇一份報告在此預覽。</p>
            </div>
            <div id="viewer-actions" class="artifact-links"></div>
          </div>
          <iframe id="report-frame" class="viewer-frame" title="Benchmark report preview / Benchmark 報告預覽"></iframe>
        </section>
      </section>
    </main>

    <div class="maintenance-toggle-row">
      <button
        id="toggle-maintenance"
        class="button button-secondary"
        type="button"
        aria-controls="maintenance-panel"
        aria-expanded="false"
      >Show Maintenance / 顯示維護設定</button>
    </div>

    <section id="maintenance-panel" class="panel stack maintenance-panel" hidden>
      <div class="status-header">
        <div>
          <h2>Maintenance Page / 維護頁面</h2>
          <p class="section-note">Edit and save the default values shown by the benchmark UI, including backend URLs, models, prompts, system prompts, and parameter presets. / 編輯並儲存 benchmark UI 顯示的預設值，包含後端網址、模型、prompt、system prompt 與參數預設。</p>
        </div>
        <div id="ui-defaults-path" class="maintenance-path"></div>
      </div>

      <div class="maintenance-grid">
        <section class="stack">
          <div class="form-grid">
            <div class="field">
              <label for="defaults-backend">Default Backend / 預設後端</label>
              <select id="defaults-backend"></select>
            </div>
            <div class="field">
              <label for="defaults-capability">Default Benchmark Mode / 預設測試模式</label>
              <select id="defaults-capability"></select>
            </div>
          </div>

          <div class="stack">
            <div class="status-header">
              <div>
                <h3>Mode Defaults / 模式預設</h3>
                <p class="section-note">Set the default prompt and default system prompt blocks for each benchmark mode. / 為每個 benchmark mode 設定預設 prompt 與預設 system prompt 區塊。</p>
              </div>
              <div class="field">
                <label for="maintenance-capability">Mode / 模式</label>
                <select id="maintenance-capability"></select>
              </div>
            </div>

            <div class="field field-wide">
              <label for="maintenance-prompt">Default Prompt / 預設 Prompt</label>
              <textarea id="maintenance-prompt"></textarea>
            </div>

            <div class="field field-wide">
              <label for="maintenance-system-prompts">Default System Prompts / 預設 System Prompts</label>
              <textarea id="maintenance-system-prompts" placeholder="Separate each system prompt with a line containing --- / 每段 system prompt 之間用一行 --- 分隔"></textarea>
              <div class="hint">Use <code>---</code> as the separator between multiple system prompt blocks. / 多段 system prompt 之間請使用 <code>---</code> 分隔。</div>
            </div>
          </div>
        </section>

        <section class="stack">
          <div class="status-header">
            <div>
              <h3>Backend Defaults / 後端預設</h3>
              <p class="section-note">Set the default base URL, model list, and parameter matrix for each backend. / 為各後端設定預設 Base URL、模型清單與參數矩陣。</p>
            </div>
            <div class="field">
              <label for="maintenance-backend">Backend / 後端</label>
              <select id="maintenance-backend"></select>
            </div>
          </div>

          <div class="form-grid">
            <div class="field field-wide">
              <label for="maintenance-base-url">Default Base URL / 預設 Base URL</label>
              <input id="maintenance-base-url" type="text" spellcheck="false">
            </div>
            <div class="field field-wide">
              <label for="maintenance-models">Default Models / 預設模型</label>
              <input id="maintenance-models" type="text" spellcheck="false" placeholder="qwen3.5:latest, llama.cpp-model">
            </div>
          </div>

          <div class="stack">
            <div>
              <h3>Default Parameter Matrix / 預設參數矩陣</h3>
              <p class="section-note">These values become the initial parameter table state when the main benchmark page loads. / 這些值會成為主 benchmark 頁面初次載入時的參數表格狀態。</p>
            </div>
            <div class="param-table-wrap">
              <table>
                <thead>
                  <tr>
                    <th>Use / 啟用</th>
                    <th>Key / 鍵</th>
                    <th>Group / 群組</th>
                    <th>Range / 範圍</th>
                    <th>Values / 值</th>
                  </tr>
                </thead>
                <tbody id="maintenance-param-table-body"></tbody>
              </table>
            </div>
            <div id="maintenance-param-summary" class="count-strip"></div>
          </div>
        </section>
      </div>

      <div class="button-row">
        <button id="save-ui-defaults" class="button button-primary" type="button">Save UI Defaults / 儲存 UI 預設</button>
        <button id="reload-ui-defaults" class="button button-secondary" type="button">Reload Saved Defaults / 重新載入已儲存預設</button>
        <button id="reset-ui-defaults" class="button button-secondary" type="button">Reset Built-in Defaults / 還原內建預設</button>
      </div>
    </section>
  </div>

  <script>
    const appState = {
      capabilityDefaults: {},
      capabilitySystemPromptDefaults: {},
      backendState: null,
      backendCatalog: {},
      uiDefaults: null,
      maintenanceDraft: null,
      job: null,
      reports: [],
      activeReportId: null,
      lastPromptDefault: "",
      lastSystemPromptsDefault: "",
      reportRefreshToken: "",
    };

    const els = {
      alert: document.getElementById("top-alert"),
      backend: document.getElementById("backend"),
      capability: document.getElementById("capability"),
      baseUrl: document.getElementById("base-url"),
      llamaBatchSwitchField: document.getElementById("llama-batch-switch-field"),
      llamaCppAutoSwitch: document.getElementById("llama-cpp-auto-switch"),
      modelCatalogLabel: document.getElementById("model-catalog-label"),
      modelCatalogNote: document.getElementById("model-catalog-note"),
      detectedModels: document.getElementById("detected-models"),
      refreshModels: document.getElementById("refresh-models"),
      modelsInput: document.getElementById("models-input"),
      promptInput: document.getElementById("prompt-input"),
      promptHint: document.getElementById("prompt-hint"),
      systemPromptsInput: document.getElementById("system-prompts-input"),
      paramBody: document.getElementById("param-table-body"),
      paramSummary: document.getElementById("param-summary"),
      startButton: document.getElementById("start-button"),
      refreshReports: document.getElementById("refresh-reports"),
      shutdownUi: document.getElementById("shutdown-ui"),
      jobStatusChip: document.getElementById("job-status-chip"),
      jobMessage: document.getElementById("job-message"),
      jobSummaryStrip: document.getElementById("job-summary-strip"),
      jobLogs: document.getElementById("job-logs"),
      consoleSummary: document.getElementById("console-summary"),
      runArtifacts: document.getElementById("run-artifacts"),
      reportList: document.getElementById("report-list"),
      reportFrame: document.getElementById("report-frame"),
      viewerCaption: document.getElementById("viewer-caption"),
      viewerActions: document.getElementById("viewer-actions"),
      toggleMaintenance: document.getElementById("toggle-maintenance"),
      maintenancePanel: document.getElementById("maintenance-panel"),
      uiDefaultsPath: document.getElementById("ui-defaults-path"),
      defaultsBackend: document.getElementById("defaults-backend"),
      defaultsCapability: document.getElementById("defaults-capability"),
      maintenanceCapability: document.getElementById("maintenance-capability"),
      maintenancePrompt: document.getElementById("maintenance-prompt"),
      maintenanceSystemPrompts: document.getElementById("maintenance-system-prompts"),
      maintenanceBackend: document.getElementById("maintenance-backend"),
      maintenanceBaseUrl: document.getElementById("maintenance-base-url"),
      maintenanceModels: document.getElementById("maintenance-models"),
      maintenanceParamBody: document.getElementById("maintenance-param-table-body"),
      maintenanceParamSummary: document.getElementById("maintenance-param-summary"),
      saveUiDefaults: document.getElementById("save-ui-defaults"),
      reloadUiDefaults: document.getElementById("reload-ui-defaults"),
      resetUiDefaults: document.getElementById("reset-ui-defaults"),
    };

    const statusLabels = {
      idle: "idle / 就緒",
      running: "running / 執行中",
      succeeded: "succeeded / 完成",
      failed: "failed / 失敗",
    };

    function escapeHtml(value) {
      return String(value ?? "")
        .replace(/&/g, "&amp;")
        .replace(/</g, "&lt;")
        .replace(/>/g, "&gt;")
        .replace(/"/g, "&quot;");
    }

    async function fetchJson(url, options = {}) {
      const response = await fetch(url, {
        ...options,
        headers: {
          "Content-Type": "application/json",
          ...(options.headers || {}),
        },
      });
      const text = await response.text();
      const data = text ? JSON.parse(text) : {};
      if (!response.ok) {
        throw new Error(data.error || "Request failed. / 請求失敗。");
      }
      return data;
    }

    function showAlert(message) {
      if (!message) {
        els.alert.classList.remove("visible");
        els.alert.textContent = "";
        return;
      }
      els.alert.textContent = message;
      els.alert.classList.add("visible");
    }

    function setMaintenancePanelVisible(isVisible) {
      els.maintenancePanel.hidden = !isVisible;
      els.toggleMaintenance.setAttribute("aria-expanded", String(isVisible));
      els.toggleMaintenance.textContent = isVisible
        ? "Hide Maintenance / 隱藏維護設定"
        : "Show Maintenance / 顯示維護設定";
    }

    function cloneData(value) {
      return JSON.parse(JSON.stringify(value ?? null));
    }

    function parseSystemPromptBlocks(text) {
      const normalized = String(text || "").replace(/\\r\\n/g, "\\n").trim();
      if (!normalized) {
        return [];
      }
      return normalized
        .split(/\\n\\s*---\\s*\\n/g)
        .map((item) => item.trim())
        .filter(Boolean);
    }

    function formatSystemPromptBlocks(prompts) {
      const items = Array.isArray(prompts)
        ? prompts.map((item) => String(item || "").trim()).filter(Boolean)
        : [];
      return items.join("\\n---\\n");
    }

    function currentParamValuesFor(bodyEl) {
      const rowMap = new Map();
      bodyEl.querySelectorAll("tr[data-param-key]").forEach((row) => {
        const key = row.dataset.paramKey;
        const enabled = row.querySelector("input[type='checkbox']")?.checked || false;
        const rawValue = row.querySelector("input[type='text']")?.value || "";
        rowMap.set(key, { enabled, rawValue });
      });
      return rowMap;
    }

    function renderParamSummaryFor(bodyEl, summaryEl) {
      const rows = [...bodyEl.querySelectorAll("tr[data-param-key]")];
      let selectedCount = 0;
      let comboCount = 1;
      let hasError = false;

      rows.forEach((row) => {
        const checkbox = row.querySelector("input[type='checkbox']");
        const valueInput = row.querySelector("input[type='text']");
        if (!checkbox || checkbox.disabled || !checkbox.checked) {
          return;
        }

        selectedCount += 1;
        const parts = String(valueInput.value || "")
          .split(",")
          .map((item) => item.trim())
          .filter(Boolean);
        if (!parts.length) {
          hasError = true;
          return;
        }
        comboCount *= parts.length;
      });

      const badges = [
        `<div class="count-badge">Selected Params / 已選參數: <strong>${selectedCount}</strong></div>`,
        `<div class="count-badge">Combination Count / 組合數量: <strong>${hasError ? "ERR" : comboCount}</strong></div>`,
      ];
      summaryEl.innerHTML = badges.join("");
    }

    function renderDetectedModels(models) {
      const backendState = appState.backendState || {};
      els.llamaBatchSwitchField.hidden = backendState.backend !== "llama.cpp";
      els.modelCatalogLabel.textContent = backendState.model_catalog_label || "Detected Models / 偵測到的模型";
      els.modelCatalogNote.textContent = backendState.model_catalog_note || "";
      if (!models || !models.length) {
        const isLlamaCpp = backendState.backend === "llama.cpp";
        els.detectedModels.innerHTML = isLlamaCpp
          ? '<div class="hint">No GGUF catalog entries found. You can still leave Models empty to test the model currently served by llama.cpp. / 找不到 GGUF 模型索引；仍可留空並直接測試目前 llama.cpp 模型。</div>'
          : '<div class="hint">No Ollama models detected right now. / 目前沒有偵測到 Ollama 模型。</div>';
        return;
      }

      const selected = new Set(
        String(els.modelsInput.value || "")
          .split(",")
          .map((item) => item.trim())
          .filter(Boolean)
      );

      els.detectedModels.innerHTML = models.map((item) => {
        const model = typeof item === "string" ? { name: item, available: true } : item;
        const name = String(model.name || "").trim();
        const path = String(model.path || "").trim();
        const suffix = model.available === false ? " (missing)" : "";
        return `
        <label class="pill">
          <input type="checkbox" value="${escapeHtml(name)}" ${selected.has(name) ? "checked" : ""}>
          <span title="${escapeHtml(path)}">${escapeHtml(name + suffix)}</span>
        </label>
      `;
      }).join("");

      els.detectedModels.querySelectorAll("input[type='checkbox']").forEach((checkbox) => {
        checkbox.addEventListener("change", syncModelsFromDetectedSelection);
      });
    }

    function syncModelsFromDetectedSelection() {
      const selected = [...els.detectedModels.querySelectorAll("input[type='checkbox']:checked")]
        .map((checkbox) => checkbox.value.trim())
        .filter(Boolean);
      els.modelsInput.value = selected.join(", ");
    }

    function buildRowsWithParamDefaults(rows, paramDefaults) {
      return (rows || []).map((row) => {
        const overrideValues = paramDefaults?.[row.key] || {};
        return {
          ...row,
          enabled: row.supported ? Boolean(overrideValues.enabled ?? row.enabled) : false,
          raw_value: row.supported
            ? String(overrideValues.raw_value ?? row.raw_value ?? row.default_value)
            : row.default_value,
        };
      });
    }

    function renderParamRowsInto(bodyEl, summaryEl, rows, { onInputChange = null, preserveCurrentValues = false } = {}) {
      const currentValues = preserveCurrentValues ? currentParamValuesFor(bodyEl) : new Map();
      bodyEl.innerHTML = (rows || []).map((row) => {
        const preserved = currentValues.get(row.key) || {};
        const enabled = row.supported
          ? (Object.prototype.hasOwnProperty.call(preserved, "enabled") ? preserved.enabled : row.enabled)
          : false;
        const rawValue = row.supported
          ? (Object.prototype.hasOwnProperty.call(preserved, "rawValue") ? preserved.rawValue : (row.raw_value ?? row.default_value))
          : row.default_value;

        return `
          <tr data-param-key="${escapeHtml(row.key)}" class="${row.supported ? "" : "unsupported-row"}">
            <td>
              <input type="checkbox" ${enabled ? "checked" : ""} ${row.supported ? "" : "disabled"}>
            </td>
            <td>
              <strong>${escapeHtml(row.key)}</strong><br>
              <span class="hint">${escapeHtml(row.label)}</span>
            </td>
            <td>${escapeHtml(row.group)}</td>
            <td>
              ${escapeHtml(row.range_text)}<br>
              <span class="hint">${row.supported ? "Supported / 支援" : "Unsupported on this backend / 此後端不支援"}</span>
            </td>
            <td>
              <input type="text" value="${escapeHtml(rawValue)}" ${row.supported ? "" : "disabled"}>
              <div class="hint">${escapeHtml(row.desc)}</div>
            </td>
          </tr>
        `;
      }).join("");

      bodyEl.querySelectorAll("input").forEach((input) => {
        input.addEventListener("input", () => {
          renderParamSummaryFor(bodyEl, summaryEl);
          if (onInputChange) {
            onInputChange();
          }
        });
        input.addEventListener("change", () => {
          renderParamSummaryFor(bodyEl, summaryEl);
          if (onInputChange) {
            onInputChange();
          }
        });
      });
      renderParamSummaryFor(bodyEl, summaryEl);
    }

    async function loadBackendState(backend, { preserveUrl = false, preserveModels = false } = {}) {
      const data = await fetchJson(`/api/backend-state?backend=${encodeURIComponent(backend)}`);
      const previousUrl = els.baseUrl.value.trim();
      const previousModels = els.modelsInput.value.trim();
      appState.backendCatalog[backend] = data;
      appState.backendState = data;
      if (!preserveUrl || !previousUrl) {
        els.baseUrl.value = data.default_url || "";
      }
      if (!preserveModels || !previousModels) {
        els.modelsInput.value = data.default_models_text || "";
      }
      renderDetectedModels(data.detected_models || []);
      renderParamRowsInto(els.paramBody, els.paramSummary, data.param_rows || []);
    }

    function collectParamPayloadFromBody(bodyEl) {
      const params = {};
      bodyEl.querySelectorAll("tr[data-param-key]").forEach((row) => {
        const key = row.dataset.paramKey;
        params[key] = {
          enabled: row.querySelector("input[type='checkbox']")?.checked || false,
          raw_value: row.querySelector("input[type='text']")?.value || "",
        };
      });
      return params;
    }

    function collectParamPayload() {
      return collectParamPayloadFromBody(els.paramBody);
    }

    function updateMaintenanceGeneralDraft() {
      if (!appState.maintenanceDraft) {
        return;
      }
      appState.maintenanceDraft.default_backend = els.defaultsBackend.value || "llama.cpp";
      appState.maintenanceDraft.default_capability = els.defaultsCapability.value || "chat";
    }

    function flushMaintenanceCapabilityDraft(targetCapability = null) {
      if (!appState.maintenanceDraft) {
        return;
      }
      const capability = targetCapability || appState.activeMaintenanceCapability || els.maintenanceCapability.value;
      if (!capability) {
        return;
      }
      appState.maintenanceDraft.capability_defaults = appState.maintenanceDraft.capability_defaults || {};
      appState.maintenanceDraft.capability_defaults[capability] = {
        prompt: String(els.maintenancePrompt.value || "").trim() || appState.capabilityDefaults[capability] || "",
        system_prompts: parseSystemPromptBlocks(els.maintenanceSystemPrompts.value),
      };
    }

    function flushMaintenanceBackendDraft(targetBackend = null) {
      if (!appState.maintenanceDraft) {
        return;
      }
      const backend = targetBackend || appState.activeMaintenanceBackend || els.maintenanceBackend.value;
      if (!backend) {
        return;
      }
      const fallbackState = appState.backendCatalog?.[backend] || {};
      appState.maintenanceDraft.backend_defaults = appState.maintenanceDraft.backend_defaults || {};
      appState.maintenanceDraft.backend_defaults[backend] = {
        url: String(els.maintenanceBaseUrl.value || "").trim() || fallbackState.default_url || "",
        models: String(els.maintenanceModels.value || "").trim(),
        params: collectParamPayloadFromBody(els.maintenanceParamBody),
      };
    }

    function renderMaintenanceCapabilityEditor() {
      if (!appState.maintenanceDraft) {
        return;
      }
      const capability = els.maintenanceCapability.value || appState.maintenanceDraft.default_capability || "chat";
      const draftValues = appState.maintenanceDraft.capability_defaults?.[capability] || {};
      appState.activeMaintenanceCapability = capability;
      els.maintenancePrompt.value = draftValues.prompt || appState.capabilityDefaults[capability] || "";
      els.maintenanceSystemPrompts.value = formatSystemPromptBlocks(draftValues.system_prompts || []);
    }

    function renderMaintenanceBackendEditor() {
      if (!appState.maintenanceDraft) {
        return;
      }
      const backend = els.maintenanceBackend.value || appState.maintenanceDraft.default_backend || "llama.cpp";
      const backendState = appState.backendCatalog?.[backend] || { param_rows: [] };
      const draftValues = appState.maintenanceDraft.backend_defaults?.[backend] || {};
      appState.activeMaintenanceBackend = backend;
      els.maintenanceBaseUrl.value = draftValues.url || backendState.default_url || "";
      els.maintenanceModels.value = draftValues.models || backendState.default_models_text || "";
      renderParamRowsInto(
        els.maintenanceParamBody,
        els.maintenanceParamSummary,
        buildRowsWithParamDefaults(backendState.param_rows || [], draftValues.params || {}),
        { onInputChange: () => flushMaintenanceBackendDraft(backend) }
      );
    }

    function renderMaintenanceEditors() {
      if (!appState.maintenanceDraft) {
        return;
      }
      els.uiDefaultsPath.textContent = `Defaults file / 預設檔: ${appState.uiDefaultsPath || ""}`;
      els.defaultsBackend.value = appState.maintenanceDraft.default_backend || "llama.cpp";
      els.defaultsCapability.value = appState.maintenanceDraft.default_capability || "chat";

      if (!els.maintenanceBackend.value || !appState.backendCatalog?.[els.maintenanceBackend.value]) {
        els.maintenanceBackend.value = appState.maintenanceDraft.default_backend || "llama.cpp";
      }
      if (!els.maintenanceCapability.value || !appState.capabilityDefaults?.[els.maintenanceCapability.value]) {
        els.maintenanceCapability.value = appState.maintenanceDraft.default_capability || "chat";
      }
      renderMaintenanceCapabilityEditor();
      renderMaintenanceBackendEditor();
    }

    function maybeApplyCapabilityDefaults(force = false) {
      const nextPromptDefault = appState.capabilityDefaults[els.capability.value] || "";
      const nextSystemPromptsDefault = appState.capabilitySystemPromptDefaults[els.capability.value] || "";
      const isBuiltInSuite = ["suite-smoke-7", "local-expert-battle-48"].includes(els.capability.value);
      const currentPromptValue = els.promptInput.value.trim();
      const currentSystemPromptsValue = els.systemPromptsInput.value.trim();

      if (force || !currentPromptValue || currentPromptValue === appState.lastPromptDefault) {
        els.promptInput.value = nextPromptDefault;
      }
      if (force || !currentSystemPromptsValue || currentSystemPromptsValue === appState.lastSystemPromptsDefault) {
        els.systemPromptsInput.value = nextSystemPromptsDefault;
      }
      appState.lastPromptDefault = nextPromptDefault;
      appState.lastSystemPromptsDefault = nextSystemPromptsDefault;
      els.promptInput.disabled = isBuiltInSuite;
      els.promptHint.textContent = isBuiltInSuite
        ? (els.capability.value === "local-expert-battle-48"
          ? "Local Expert Battle runs 48 fixed questions and loads 2–3K-character excerpts from ~/wiki; the prompt cannot be edited here. / Local Expert Battle 會執行 48 題固定題目，並自 ~/wiki 擷取 2–3K 字，提示不可修改。"
          : "Built-in suite-smoke-7 runs seven fixed questions; the prompt cannot be edited here. / 內建 suite-smoke-7 會執行七道固定題目，此處不可修改。")
        : "This prompt is sent once per configuration. / 此提示會依每組設定送出一次。";
    }

    function applyUiDefaultsPayload(data, { applyToBenchmark = false } = {}) {
      appState.uiDefaults = cloneData(data.ui_defaults || {});
      appState.maintenanceDraft = cloneData(data.ui_defaults || {});
      appState.backendCatalog = cloneData(data.backend_catalog || {});
      appState.capabilityDefaults = { ...(data.capability_defaults || {}) };
      appState.capabilitySystemPromptDefaults = { ...(data.capability_system_prompt_defaults || {}) };
      appState.uiDefaultsPath = data.ui_defaults_path || "";
      renderMaintenanceEditors();

      if (applyToBenchmark) {
        const defaultBackend = appState.uiDefaults?.default_backend || "llama.cpp";
        const defaultCapability = appState.uiDefaults?.default_capability || "chat";
        els.backend.value = defaultBackend;
        els.capability.value = defaultCapability;
        const backendState = appState.backendCatalog?.[defaultBackend];
        if (backendState) {
          appState.backendState = backendState;
          els.baseUrl.value = backendState.default_url || "";
          els.modelsInput.value = backendState.default_models_text || "";
          renderDetectedModels(backendState.detected_models || []);
          renderParamRowsInto(els.paramBody, els.paramSummary, backendState.param_rows || []);
        }
        maybeApplyCapabilityDefaults(true);
      }
    }

    function collectConfigPayload() {
      return {
        backend: els.backend.value,
        capability: els.capability.value,
        url: els.baseUrl.value.trim(),
        models: els.modelsInput.value.trim(),
        llama_cpp_auto_switch: els.backend.value === "llama.cpp" && els.llamaCppAutoSwitch.checked,
        prompt: els.promptInput.value,
        system_prompts: els.systemPromptsInput.value,
        params: collectParamPayload(),
      };
    }

    function renderArtifactLinks(container, links, extraActions = []) {
      const parts = [];
      (links || []).forEach((link) => {
        parts.push(`<a href="${escapeHtml(link.url)}" target="_blank" rel="noreferrer">${escapeHtml(link.label)}</a>`);
      });
      extraActions.forEach((action) => {
        parts.push(`<button type="button" data-action="${escapeHtml(action.action)}">${escapeHtml(action.label)}</button>`);
      });
      container.innerHTML = parts.join("");
    }

    function renderJobSummary(job) {
      if (!job || !job.result || !job.result.counts) {
        els.jobSummaryStrip.innerHTML = "";
        return;
      }
      const counts = job.result.counts;
      els.jobSummaryStrip.innerHTML = `
        <div class="summary-item"><strong>${counts.ok}</strong><span>OK / 正常</span></div>
        <div class="summary-item"><strong>${counts.warning}</strong><span>Warning / 警告</span></div>
        <div class="summary-item"><strong>${counts.error}</strong><span>Error / 錯誤</span></div>
        <div class="summary-item"><strong>${escapeHtml(job.result.config_summary?.estimated_run_count ?? job.result.config_summary?.combination_count ?? "-")}</strong><span>Runs / 執行數</span></div>
      `;
    }

    function maybePreviewLatestReport(job) {
      const latestReport = job?.result?.latest_report;
      if (!latestReport || !latestReport.html_url) {
        return;
      }
      if (appState.activeReportId === latestReport.id) {
        return;
      }
      previewReport(latestReport.id);
    }

    function renderJob(job) {
      appState.job = job;
      const status = job?.status || "idle";
      els.jobStatusChip.className = `status-chip status-${status}`;
      els.jobStatusChip.textContent = statusLabels[status] || status;
      els.jobMessage.textContent = job?.message || "Ready. / 就緒。";

      const logLines = job?.logs?.length ? job.logs.join("\\n") : "No job started yet. / 尚未開始任務。";
      els.jobLogs.textContent = logLines;

      const errorText = job?.error ? `${job.error}` : "";
      if (errorText) {
        showAlert(errorText);
      } else {
        showAlert("");
      }

      renderJobSummary(job);

      if (job?.result?.console_summary_text) {
        els.consoleSummary.textContent = job.result.console_summary_text;
      } else {
        els.consoleSummary.textContent = "Benchmark table output will appear here after a run. / 執行後會在這裡顯示 benchmark 表格摘要。";
      }

      renderArtifactLinks(els.runArtifacts, job?.result?.artifact_links || []);

      const running = status === "running";
      els.startButton.disabled = running;
      els.startButton.textContent = running
        ? "Benchmark Running... / 測試執行中..."
        : "Start Benchmark / 開始測試";

      if (status === "succeeded") {
        maybePreviewLatestReport(job);
      }
    }

    function renderReportList(reports) {
      appState.reports = reports || [];
      if (!appState.reports.length) {
        els.reportList.innerHTML = '<div class="empty-state">No saved HTML reports yet. Start a benchmark run and the latest report will appear here. / 目前還沒有已儲存的 HTML 報告。開始一次 benchmark 後，最新報告會顯示在這裡。</div>';
        return;
      }

      els.reportList.innerHTML = appState.reports.map((report) => `
        <article class="report-card ${appState.activeReportId === report.id ? "active" : ""}" data-report-id="${escapeHtml(report.id)}">
          <div>
            <h4>${escapeHtml(report.title)}</h4>
            <div class="report-meta">
              <span>${escapeHtml(report.modified_at)}</span>
              <span>${escapeHtml(report.size_kib)} KiB</span>
            </div>
          </div>
          <div class="artifact-links">
            <button type="button" data-preview-report="${escapeHtml(report.id)}">Preview / 預覽</button>
            <a href="${escapeHtml(report.html_url)}" target="_blank" rel="noreferrer">Open / 開啟</a>
          </div>
        </article>
      `).join("");

      els.reportList.querySelectorAll("[data-preview-report]").forEach((button) => {
        button.addEventListener("click", () => previewReport(button.dataset.previewReport));
      });
    }

    function previewReport(reportId) {
      const report = appState.reports.find((item) => item.id === reportId);
      if (!report) {
        return;
      }
      appState.activeReportId = report.id;
      els.reportFrame.src = report.html_url;
      els.viewerCaption.textContent = `${report.title} / ${report.modified_at}`;
      renderArtifactLinks(els.viewerActions, report.artifact_links || []);
      renderReportList(appState.reports);
    }

    async function refreshReports({ preserveSelection = true } = {}) {
      const data = await fetchJson(`/api/reports?token=${encodeURIComponent(String(Date.now()))}`);
      renderReportList(data.reports || []);
      if (!preserveSelection && data.reports?.length) {
        previewReport(data.reports[0].id);
        return;
      }
      if (preserveSelection && appState.activeReportId) {
        const stillExists = data.reports?.some((report) => report.id === appState.activeReportId);
        if (stillExists) {
          previewReport(appState.activeReportId);
          return;
        }
      }
      if (!appState.activeReportId && data.reports?.length) {
        previewReport(data.reports[0].id);
      }
    }

    async function refreshJob() {
      const data = await fetchJson(`/api/job?token=${encodeURIComponent(String(Date.now()))}`);
      const previousStatus = appState.job?.status;
      renderJob(data.job);
      if (previousStatus === "running" && data.job?.status !== "running") {
        await refreshReports({ preserveSelection: false });
      }
    }

    async function startBenchmark() {
      try {
        showAlert("");
        els.startButton.disabled = true;
        const payload = collectConfigPayload();
        const data = await fetchJson("/api/start-benchmark", {
          method: "POST",
          body: JSON.stringify(payload),
        });
        renderJob(data.job);
      } catch (error) {
        els.startButton.disabled = false;
        showAlert(error.message);
      }
    }

    async function reloadUiDefaults() {
      const data = await fetchJson("/api/ui-defaults");
      applyUiDefaultsPayload(data, { applyToBenchmark: true });
    }

    async function saveUiDefaults() {
      updateMaintenanceGeneralDraft();
      flushMaintenanceCapabilityDraft();
      flushMaintenanceBackendDraft();
      const data = await fetchJson("/api/ui-defaults", {
        method: "POST",
        body: JSON.stringify(appState.maintenanceDraft || {}),
      });
      applyUiDefaultsPayload(data, { applyToBenchmark: true });
    }

    async function resetUiDefaults() {
      const data = await fetchJson("/api/ui-defaults/reset", {
        method: "POST",
        body: JSON.stringify({}),
      });
      applyUiDefaultsPayload(data, { applyToBenchmark: true });
    }

    async function shutdownUi() {
      try {
        await fetchJson("/api/shutdown", { method: "POST", body: JSON.stringify({}) });
        els.shutdownUi.disabled = true;
        els.startButton.disabled = true;
        els.jobMessage.textContent = "Local UI is shutting down... / 本機 UI 正在關閉...";
      } catch (error) {
        showAlert(error.message);
      }
    }

    async function bootstrap() {
      try {
        const data = await fetchJson("/api/bootstrap");
        const backendOptionsHtml = (data.backends || []).map((item) => `
          <option value="${escapeHtml(item.value)}">${escapeHtml(item.label)}</option>
        `).join("");
        const capabilityOptionsHtml = (data.capabilities || []).map((item) => `
          <option value="${escapeHtml(item.value)}">${escapeHtml(item.label)}</option>
        `).join("");
        els.backend.innerHTML = backendOptionsHtml;
        els.defaultsBackend.innerHTML = backendOptionsHtml;
        els.maintenanceBackend.innerHTML = backendOptionsHtml;
        els.capability.innerHTML = capabilityOptionsHtml;
        els.defaultsCapability.innerHTML = capabilityOptionsHtml;
        els.maintenanceCapability.innerHTML = capabilityOptionsHtml;

        applyUiDefaultsPayload(data, { applyToBenchmark: true });
        renderJob(data.job);
        renderReportList(data.reports || []);
        if (data.reports?.length) {
          previewReport(data.reports[0].id);
        }

        els.backend.addEventListener("change", async () => {
          try {
            await loadBackendState(els.backend.value, { preserveUrl: false, preserveModels: false });
          } catch (error) {
            showAlert(error.message);
          }
        });

        els.capability.addEventListener("change", () => maybeApplyCapabilityDefaults(false));
        els.refreshModels.addEventListener("click", async () => {
          try {
            await loadBackendState(els.backend.value, { preserveUrl: true, preserveModels: true });
          } catch (error) {
            showAlert(error.message);
          }
        });
        els.modelsInput.addEventListener("input", () => {
          renderDetectedModels(appState.backendState?.detected_models || []);
        });
        els.refreshReports.addEventListener("click", () => refreshReports({ preserveSelection: true }).catch((error) => showAlert(error.message)));
        els.startButton.addEventListener("click", startBenchmark);
        els.shutdownUi.addEventListener("click", shutdownUi);
        els.defaultsBackend.addEventListener("change", updateMaintenanceGeneralDraft);
        els.defaultsCapability.addEventListener("change", updateMaintenanceGeneralDraft);
        els.maintenanceCapability.addEventListener("change", () => {
          flushMaintenanceCapabilityDraft(appState.activeMaintenanceCapability);
          renderMaintenanceCapabilityEditor();
        });
        els.maintenanceBackend.addEventListener("change", () => {
          flushMaintenanceBackendDraft(appState.activeMaintenanceBackend);
          renderMaintenanceBackendEditor();
        });
        els.maintenancePrompt.addEventListener("input", () => flushMaintenanceCapabilityDraft());
        els.maintenanceSystemPrompts.addEventListener("input", () => flushMaintenanceCapabilityDraft());
        els.maintenanceBaseUrl.addEventListener("input", () => flushMaintenanceBackendDraft());
        els.maintenanceModels.addEventListener("input", () => flushMaintenanceBackendDraft());
        els.saveUiDefaults.addEventListener("click", () => saveUiDefaults().catch((error) => showAlert(error.message)));
        els.reloadUiDefaults.addEventListener("click", () => reloadUiDefaults().catch((error) => showAlert(error.message)));
        els.resetUiDefaults.addEventListener("click", () => resetUiDefaults().catch((error) => showAlert(error.message)));

        setInterval(() => {
          refreshJob().catch((error) => showAlert(error.message));
        }, 1600);
      } catch (error) {
        showAlert(error.message);
      }
    }

    els.toggleMaintenance.addEventListener("click", () => {
      setMaintenancePanelVisible(els.maintenancePanel.hidden);
    });
    setMaintenancePanelVisible(false);
    bootstrap();
  </script>
</body>
</html>
"""


def run_web_benchmark_job(app_state, job_id, config):
    try:
        result = run_benchmark_workflow(
            config,
            progress_callback=lambda message: app_state.append_log(job_id, message),
        )
        for warning_text in result.get("warnings", []):
            app_state.append_log(job_id, warning_text)
        app_state.complete_job(job_id, result)
    except Exception as exc:
        traceback.print_exc()
        app_state.fail_job(job_id, f"{type(exc).__name__}: {exc}")


class BenchmarkWebUiRequestHandler(BaseHTTPRequestHandler):
    server_version = "DIYLLMBenchmarkUI/1.0"

    def log_message(self, _format, *_args):
        return

    @property
    def app_state(self):
        return self.server.app_state

    @property
    def report_root(self):
        return self.server.report_root

    def _send_json(self, payload, status=200):
        body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _send_html(self, html_text, status=200):
        body = html_text.encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "text/html; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _send_file(self, file_path):
        target_path = Path(file_path)
        body = target_path.read_bytes()
        content_type = mimetypes.guess_type(target_path.name)[0] or "application/octet-stream"
        if content_type.startswith("text/") or content_type == "application/json":
            content_type += "; charset=utf-8"

        self.send_response(200)
        self.send_header("Content-Type", content_type)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(body)

    def _read_json_body(self):
        content_length = int(self.headers.get("Content-Length", "0") or "0")
        raw_body = self.rfile.read(content_length) if content_length else b"{}"
        if not raw_body.strip():
            return {}
        return json.loads(raw_body.decode("utf-8"))

    def _not_found(self):
        self._send_json({"error": "Not found."}, status=404)

    def do_GET(self):
        parsed = urlparse(self.path)

        if parsed.path == "/":
            self._send_html(build_single_file_benchmark_ui_html())
            return

        if parsed.path == "/api/bootstrap":
            self._send_json(build_web_ui_bootstrap_payload(self.app_state))
            return

        if parsed.path == "/api/backend-state":
            query = parse_qs(parsed.query)
            backend = (query.get("backend") or [DEFAULT_BACKEND])[0]
            self._send_json(build_web_ui_backend_state(backend))
            return

        if parsed.path == "/api/ui-defaults":
            self._send_json(build_ui_defaults_api_payload())
            return

        if parsed.path == "/api/job":
            self._send_json({"job": self.app_state.snapshot()})
            return

        if parsed.path == "/api/reports":
            self._send_json({"reports": list_report_entries(self.report_root)})
            return

        if parsed.path.startswith("/report-files/"):
            file_name = unquote(parsed.path.removeprefix("/report-files/"))
            if not file_name or Path(file_name).name != file_name:
                self._not_found()
                return
            target_path = self.report_root / file_name
            if not target_path.exists() or not target_path.is_file():
                self._not_found()
                return
            self._send_file(target_path)
            return

        self._not_found()

    def do_POST(self):
        parsed = urlparse(self.path)

        if parsed.path == "/api/start-benchmark":
            try:
                payload = self._read_json_body()
                config = normalize_web_ui_config(payload)
                job_id = self.app_state.start_job(config)
                worker = threading.Thread(
                    target=run_web_benchmark_job,
                    args=(self.app_state, job_id, config),
                    daemon=True,
                )
                worker.start()
                self._send_json({"job": self.app_state.snapshot()}, status=202)
            except Exception as exc:
                self._send_json({"error": str(exc)}, status=400)
            return

        if parsed.path == "/api/ui-defaults":
            try:
                payload = self._read_json_body()
                saved_defaults = save_ui_defaults(payload)
                self._send_json(build_ui_defaults_api_payload(saved_defaults))
            except Exception as exc:
                self._send_json({"error": str(exc)}, status=400)
            return

        if parsed.path == "/api/ui-defaults/reset":
            try:
                reset_defaults = reset_ui_defaults()
                self._send_json(build_ui_defaults_api_payload(reset_defaults))
            except Exception as exc:
                self._send_json({"error": str(exc)}, status=400)
            return

        if parsed.path == "/api/shutdown":
            self._send_json({"ok": True})
            threading.Thread(target=self.server.shutdown, daemon=True).start()
            return

        self._not_found()


def launch_single_file_benchmark_ui():
    report_root = ensure_report_output_dir().resolve()
    app_state = BenchmarkWebUiState()
    server = None
    for port_offset in range(DEFAULT_UI_PORT_SCAN_LIMIT):
        candidate_port = DEFAULT_UI_PORT + port_offset
        try:
            server = ThreadingHTTPServer((DEFAULT_UI_HOST, candidate_port), BenchmarkWebUiRequestHandler)
            break
        except OSError:
            continue
    if server is None:
        server = ThreadingHTTPServer((DEFAULT_UI_HOST, 0), BenchmarkWebUiRequestHandler)
    server.daemon_threads = True
    server.app_state = app_state
    server.report_root = report_root

    url = f"http://{DEFAULT_UI_HOST}:{server.server_address[1]}/"
    launch_hint_path = persist_ui_launch_hint(url, report_root)
    print("\n" + "=" * 62)
    print("DIY LLM Benchmark | Single-file HTML UI")
    print("=" * 62)
    print(f"UI URL: {url}")
    print(f"Report directory: {report_root}")
    print(f"Launch hint file: {launch_hint_path.resolve()}")
    print("Press Ctrl-C in this terminal to stop the local UI server.")

    browser_opened, browser_error = try_open_browser(url)
    if browser_opened:
        print("Browser auto-open: OK")
    else:
        print("Browser auto-open: failed")
        print(f"Open this URL manually: {url}")
        print(f"Browser error: {browser_error}")
    if not sys.stdin.isatty():
        threading.Thread(
            target=show_windows_info_dialog,
            args=(
                "llm_expert_bench UI Ready",
                build_ui_launch_message(url, report_root, launch_hint_path, browser_opened),
            ),
            daemon=True,
        ).start()

    try:
        server.serve_forever()
    finally:
        server.server_close()


def interactive_config():
    print("\n" + "=" * 62)
    print("LLM Benchmark")
    print("=" * 62)
    state = {}
    stage_index = 0

    while True:
        if stage_index == 0:
            backend = ask_select_with_back(
                "Select backend:",
                choices=[
                    Choice("Ollama", value="ollama"),
                    Choice("llama.cpp (llama-server)", value="llama.cpp"),
                ],
                default=state.get("backend"),
            )
            if backend in (None, BACK_ACTION):
                return None
            if backend != state.get("backend"):
                state.pop("url", None)
                state.pop("models", None)
                state.pop("params", None)
            state["backend"] = backend
            stage_index = 1
            continue

        if stage_index == 1:
            capability = ask_select_with_back(
                "Select benchmark mode:",
                choices=[
                    Choice("Chat | Standard chat response benchmark", value="chat"),
                    Choice("Tools | Check whether the model emits tool_calls", value="tools"),
                    Choice(
                        "Suite Smoke 7 | Seven fixed capability questions / 七項能力固定題庫",
                        value="suite-smoke-7",
                    ),
                    Choice(
                        "Local Expert Battle 48 | PLC、工程計算、繁中與 wiki 摘要",
                        value=LOCAL_EXPERT_BATTLE_SUITE_ID,
                    ),
                ],
                default=state.get("capability"),
            )
            if capability is None:
                return None
            if capability == BACK_ACTION:
                stage_index = 0
                continue
            state["capability"] = capability
            stage_index = 2
            continue

        if stage_index == 2:
            model_result = select_models_and_url(
                state["backend"],
                previous_url=state.get("url"),
                previous_models=state.get("models"),
            )
            if model_result is None:
                return None
            if model_result == BACK_ACTION:
                stage_index = 1
                continue
            url, models = model_result
            if not models:
                print("No models available. Cancelled.")
                return None
            state["url"] = url
            state["models"] = models
            stage_index = 3
            continue

        if stage_index == 3:
            final_params = edit_param_grid(state["backend"], initial_params=state.get("params"))
            if final_params is None:
                print("Parameter grid cancelled.")
                return None
            if final_params == BACK_ACTION:
                stage_index = 2
                continue
            state["params"] = final_params
            stage_index = 4
            continue

        if stage_index == 4:
            if state["capability"] in BUILTIN_SUITES:
                state["prompt"] = CAPABILITY_DEFAULTS[state["capability"]]
                stage_index = 5
                continue
            prompt = ask_text_with_back(
                "Benchmark prompt:",
                default=state.get("prompt", CAPABILITY_DEFAULTS[state["capability"]]),
            )
            if prompt is None:
                return None
            if prompt == BACK_ACTION:
                stage_index = 3
                continue
            state["prompt"] = prompt
            stage_index = 5
            continue

        if stage_index == 5:
            system_prompts = select_system_prompt_variants(state.get("system_prompts"))
            if system_prompts is None:
                return None
            if system_prompts == BACK_ACTION:
                stage_index = 3 if state["capability"] in BUILTIN_SUITES else 4
                continue
            state["system_prompts"] = system_prompts
            stage_index = 6
            continue

        config = {
            "backend": state["backend"],
            "capability": state["capability"],
            "url": state["url"],
            "models": state["models"],
            "params": state.get("params", {}),
            "prompt": state["prompt"],
            "system_prompts": state.get("system_prompts", []),
        }

        print_config_review(config)
        confirmed = ask_confirm_with_back(
            "Start benchmark with this configuration?",
            default=True,
        )
        if confirmed is None:
            return None
        if confirmed == BACK_ACTION:
            stage_index = 5
            continue
        if not confirmed:
            print("Cancelled before benchmark run.")
            return None
        return config


def main():
    launch_single_file_benchmark_ui()


def main_cli():
    config = interactive_config()
    if not config:
        return

    result = run_benchmark_workflow(config)

    print("\n" + "=" * 62)
    print(result["console_summary_text"])
    counts = result["counts"]
    print(
        f"\nResult counts: ok={counts['ok']}, warning={counts['warning']}, error={counts['error']}"
    )
    print(f"Saved artifacts directory: {result['report_dir']}")
    for link in result.get("artifact_links", []):
        print(f"- {link['label']}: {link['name']}")
    if result.get("warnings"):
        print("\nWarnings:")
        for warning_text in result["warnings"]:
            print(f"- {warning_text}")

if __name__ == "__main__":
    exit_code = 0
    try:
        # Explorer and shortcuts may start the script from an unrelated working directory.
        os.chdir(Path(__file__).resolve().parent)
        use_cli_mode = "--cli" in sys.argv
        ensure_runtime_ready(require_questionary=use_cli_mode)
        if use_cli_mode:
            main_cli()
        else:
            main()
    except KeyboardInterrupt:
        print("\n已取消執行。")
    except Exception as exc:
        exit_code = 1
        handle_fatal_error(exc)
    finally:
        pause_before_exit()
    if exit_code:
        sys.exit(exit_code)
